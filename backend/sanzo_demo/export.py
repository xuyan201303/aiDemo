"""Export current snapshots with actual provenance; grouping never writes classification."""
from collections import Counter
from .comparison import comparison_groups, comparison_projection
from .review_state import review_facts
from .current_document_review import current_document_reviews, issue_text
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CONDITIONS = ("transport_mode", "equipment", "billing_unit", "charge_scope", "extra_terms")


def append_text(sheet, values):
    # Literal text preserves long decimals and prevents Excel formula injection.
    sheet.append([None] * len(values))
    for column, value in enumerate(values, 1):
        cell = sheet.cell(sheet.max_row, column)
        cell.value = "" if value is None else str(value)
        cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def source_position(item):
    source = item.get("source_ref") or {}
    locator = source.get("locator") or {}
    status = source.get("locator_status", "UNAVAILABLE")
    if status in {"VERIFIED", "PAGE_ONLY"}:
        page = locator.get("page_number")
        if source.get("kind") == "PDF_PAGE" and type(page) is int and page > 0:
            return f"PDF {page} ページ" + ("（ページ内範囲は未確定）" if status == "PAGE_ONLY" else "")
        sheet = locator.get("sheet_name")
        if source.get("kind") == "EXCEL_RANGE" and isinstance(sheet, str) and sheet:
            area = locator.get("A1_range")
            if status == "VERIFIED" and isinstance(area, str) and area:
                return sheet + "!" + area
            return sheet + "（セル位置は未取得）"
    return "原文位置は未取得"


def _issue_text(value):
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return "：".join(str(value[k]) for k in ("code", "reason", "limitation", "field", "resolution", "message") if value.get(k)) or json.dumps(value, ensure_ascii=False)
    return str(value)


def _document_status(document, tasks):
    matching = [t for t in tasks if t.get("document_id") == document["document_id"]]
    task = max(matching, key=lambda t: str(t.get("created_at") or t.get("updated_at") or ""), default={})
    state = task.get("status") or task.get("state") or document.get("latest_job_state") or document.get("state") or "UNKNOWN"
    code = task.get("error_code") or document.get("error_code")
    if not code and state in {"BLOCKED", "FAILED", "INTERRUPTED"}:
        code = document.get("blocked_reason")
    issues = list(task.get("issues") or document.get("issues") or [])
    reasons = ([str(code)] if code else []) + [_issue_text(issue) for issue in issues]
    return state, " / ".join(reasons)


def _item_issues(item):
    return review_facts(item)["review_reasons"]


def export_review_state(items, documents, tasks=()):
    """Current quotation blockers only; future learning notices remain audit facts."""
    reviews = current_document_reviews(items, documents, tasks)
    # File-level unresolved issues and member-level issues are each listed once.
    # The document card may summarize both, but an item's review must not create
    # a duplicate empty warning row for its otherwise available document.
    document_states = {document_id: (
        "AVAILABLE" if review["historical_state"] in {"AVAILABLE", "REVIEW_REQUIRED"}
            and not review["unresolved_issues"] else review["current_state"],
        " / ".join(issue_text(issue) for issue in review["unresolved_issues"]))
        for document_id, review in reviews.items()}
    item_issues = {item["item_id"]: _item_issues(item) for item in items}
    draft = (not items or
             any(state != "AVAILABLE" or reason for state, reason in document_states.values()) or
             any(item_issues.values()))
    return draft, document_states, item_issues


def _classification_label(item):
    if not item.get("current_fee_id") and item.get("ai_suggested_name"):
        return "AI 提案・未確認"
    return {"MEMORY_APPLIED": "過去の修正を適用",
            "HUMAN_CONFIRMED": "今回の確認を保存"}.get(item.get("saved_status"), item.get("saved_status") or "未確認")


def workbook_bytes(case: dict, items: list[dict], documents: list[dict], *,
                   tasks: list[dict] | None = None, fees: list[dict] | None = None, reader_mode: str | None = None) -> bytes:
    tasks, fees = tasks or [], fees or []
    docs = {d["document_id"]: d for d in documents}
    draft, document_states, item_issues = export_review_state(items, documents, tasks)
    marker = "DRAFT／未確認・草稿" if draft else "保存済み案件スナップショット"
    counts = Counter(item.get("input_kind") or "UNKNOWN" for item in items)
    origin_names = {"MOCK_SDK": "デモ用解析データ／SDK未接続", "SDK_OBSERVED": "SDK 読取", "MANUAL_CONFIRMED": "内部デバッグ人工入力",
                    "SYNTHETIC_STRUCTURED": "内部デバッグ合成入力", "UNKNOWN": "入力方式未取得"}
    origin_summary = " / ".join(f"{origin_names.get(kind, kind)} {count} 明細" for kind, count in sorted(counts.items()))
    if reader_mode == "MOCK_SDK" and not counts:
        origin_summary = "デモ用解析データ／SDK未接続（明細処理未完了）"
    book = Workbook()
    try:
        comparison = book.active
        comparison.title = "原価単価一覧表"
        companies = sorted({i["observation"].get("forwarder_id") or "会社未確認" for i in items})
        append_text(comparison, ["案件", case["name"], "案件版", case["case_revision"]])
        append_text(comparison, [marker, "各社の今回資料の単価・通貨"])
        projection = comparison_projection(items, fees, documents)
        companies = projection["companies"]
        append_text(comparison, ["統一費目"] + companies)
        for group in projection["rows"]:
            label = group["display_name"]
            values = [label]
            for company in companies:
                values.append("\n".join(f'{cell["price"]} {cell["currency"]}' for cell in group["cells_by_forwarder"][company]) or "—")
            append_text(comparison, values)
        details = book.create_sheet("明細と出典")
        append_text(details, ["明細ID", "会社", "原始費目", "採用時の費目名", "単価（本次）", "通貨",
                              *CONDITIONS, "入力方式", "保存状態", "初回適用学習版", "明細版",
                              "原本名", "原本SHA256", "原本ID", "位置", "原文金額", "見積日", "有効期限",
                              "AI 提案（未確認の参考情報）", "位置確認状態", "現在の確認記録状態", "資料処理状態", "読取由来", "fixture版", "fixtureSHA256", "学習由来",
                              "元の機材", "人工確認した機材", "有効な機材",
                              "元の課金単位", "人工確認した課金単位", "有効な課金単位",
                              "元の範囲", "人工確認した範囲", "有効な範囲",
                              "元の追加条件", "人工確認した追加条件", "有効な追加条件", "現在の資料確認状態"])
        current_documents = current_document_reviews(items, documents, tasks)
        for item in items:
            obs = item["observation"]
            doc = docs.get(obs["document_id"], {})
            initial = item.get("initial_classification") or {}
            context = {**obs["context"], **(item.get("effective_context") or {})}
            overrides = item.get("condition_overrides") or {}
            append_text(details, [item["item_id"], obs.get("forwarder_id"), obs["raw_label"],
                                  item.get("current_fee_name"), item["current_price"], item["currency"],
                                  *(context.get(k) for k in CONDITIONS),
                                  item["input_kind"], item["saved_status"], initial.get("memory_version"),
                                  item["version"], doc.get("original_filename"), doc.get("sha256"), obs["document_id"],
                                  source_position(item), item.get("raw_value"), item.get("quotation_date"), item.get("valid_until"),
                                  item.get("ai_suggested_name"), (item.get("source_ref") or {}).get("locator_status", "UNAVAILABLE"),
                                  item.get("current_memory_status"), _document_status(doc, tasks)[0] if doc else "UNKNOWN",
                                  (item.get("source_ref") or {}).get("origin"),
                                  (item.get("source_ref") or {}).get("fixture_version"),
                                  (item.get("source_ref") or {}).get("fixture_sha256"),
                                  item.get("current_memory_origin"),
                                  obs["context"].get("equipment"), (overrides.get("equipment") or {}).get("human_confirmed_value"), context.get("equipment"),
                                  obs["context"].get("billing_unit"), (overrides.get("billing_unit") or {}).get("human_confirmed_value"), context.get("billing_unit"),
                                  obs["context"].get("charge_scope"), (overrides.get("charge_scope") or {}).get("human_confirmed_value"), context.get("charge_scope"),
                                  obs["context"].get("extra_terms"), (overrides.get("extra_terms") or {}).get("human_confirmed_value"), context.get("extra_terms"),
                                  current_documents.get(obs["document_id"], {}).get("current_state", "UNKNOWN")])
        issues = book.create_sheet("未確認事項")
        append_text(issues, [marker, origin_summary, "保存時点の処理・確認状態。AI提案は人工確認ではありません。"])
        append_text(issues, ["ファイル／明細", "処理・確認状態", "確認事項", "SHA256"])
        for document in documents:
            status, reason = document_states[document["document_id"]]
            if status != "AVAILABLE" or reason:
                append_text(issues, [document["original_filename"], status, reason, document.get("sha256")])
        for item in items:
            if item_issues[item["item_id"]]:
                append_text(issues, [item["item_id"], item["saved_status"], " / ".join(item_issues[item["item_id"]]), ""])
        for sheet in book:
            header_row = 3 if sheet is comparison else 2 if sheet is issues else 1
            for cell in sheet[header_row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="25558A")
            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
            for col in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 30 if col < 4 else 24
        output = BytesIO()
        row_audit = book.create_sheet("確認情報")
        append_text(row_audit, ["統一費目", "統一状態", "会社", "原文費目", "AI提案", "採用時の費目名",
                                "費目確認", "学習適格", "有効な学習記録", "今回の学習適用", "要確認理由", "確認記録版", "出典", "学習上の注意",
                                "今回の条件確認", "今回の原文照合", "人工確認の根拠・操作人・日時", "資料の参考情報", "元の処理状態・警告"])
        document_reviews = current_document_reviews(items, documents, tasks)
        for row in projection["rows"]:
            for member in row["member_items"]:
                facts = member.get("review_facts") or review_facts(member)
                confirmations = member.get("quote_confirmations") or {}
                confirmation_events = list((confirmations.get("conditions") or {}).values())
                if confirmations.get("source"):
                    confirmation_events.append(confirmations["source"])
                evidence = "\n".join((event.get("field_name") or "原文") + "：" + event.get("evidence_note", "")
                    + " / " + event.get("actor_id", "") + " / " + str(event.get("created_at", ""))
                    + ("（現在有効）" if event.get("valid") else "（対象値変更により再確認）")
                    for event in confirmation_events)
                for field, overlay in (member.get("condition_overrides") or {}).items():
                    evidence += "\n" + field + "（条件修正）：" + (overlay.get("evidence_note") or "旧形式の修正記録：補足根拠なし")
                    evidence += " / " + overlay.get("actor_id", "") + " / " + str(overlay.get("confirmed_at", ""))
                doc_review = document_reviews.get(member["observation"]["document_id"], {})
                coverage = doc_review.get("coverage_confirmation")
                if coverage:
                    evidence += "\n資料全体：" + coverage.get("evidence_note", "") + " / " + coverage.get("actor_id", "")
                    evidence += " / " + str(coverage.get("created_at", "")) + ("（現在有効）" if coverage.get("valid") else "（再確認）")
                append_text(row_audit, [row["display_name"], row["confirmation_status"], member["forwarder"],
                    member["observation"].get("raw_label"), member.get("ai_suggested_name"), member.get("current_fee_name"),
                    "確認済み" if facts["classification_confirmed"] else "未確認",
                    "適格" if facts["learning_eligible"] else "対象外",
                    facts["reusable_memory"],
                    "適用あり" if facts["memory_applied"] else "適用なし",
                    " / ".join(facts["review_reasons"]),
                    member.get("current_memory_version", 0), source_position(member),
                    " / ".join(facts["learning_notices"]),
                    " / ".join(field for field, confirmed in facts["condition_confirmations"].items() if confirmed),
                    "人工照合済み（原文位置の精度は変更なし）" if facts["source_confirmed"] else "人工照合記録なし",
                    evidence, " / ".join(facts["metadata_notices"]),
                    doc_review.get("historical_state", "") + " / " + " / ".join(issue_text(issue) for issue in doc_review.get("historical_issues", []))])
        for cell in row_audit[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="25558A")
        row_audit.freeze_panes = "A2"
        row_audit.auto_filter.ref = f"A1:S{row_audit.max_row}"
        widths = (32, 18, 16, 30, 25, 25, 17, 17, 20, 19, 46, 16, 29, 46)
        for column, width in enumerate(widths, 1):
            row_audit.column_dimensions[get_column_letter(column)].width = width
        row_audit.row_dimensions[1].height = 32
        for number in range(2, row_audit.max_row + 1):
            row_audit.row_dimensions[number].height = 42
        book.save(output)
        return output.getvalue()
    finally:
        book.close()
