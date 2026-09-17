"""Bounded original storage and a preview-only Excel renderer, never SDK extraction."""
from __future__ import annotations

from hashlib import sha256
from html import escape
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_UPLOAD = 20 * 1024 * 1024
MEDIA = {".pdf": "application/pdf",
         ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def validate_original(path: Path, suffix: str) -> str:
    if suffix not in MEDIA:
        raise ValueError("PDF または .xlsx ファイルを選択してください")
    size = path.stat().st_size
    if not 0 < size <= MAX_UPLOAD:
        raise ValueError("ファイルは 1 byte ～ 20 MiB にしてください")
    if suffix == ".pdf":
        with path.open("rb") as handle:
            if not handle.read(8).startswith(b"%PDF-"):
                raise ValueError("PDF の署名を確認できません")
        return MEDIA[suffix]
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 5000 or sum(m.file_size for m in members) > 80 * 1024 * 1024:
                raise ValueError("Excel の展開サイズがプレビュー上限を超えています")
            seen = set()
            for member in members:
                name = member.filename
                rel = PurePosixPath(name)
                if (name in seen or rel.is_absolute() or ".." in rel.parts
                        or "\\" in name or member.flag_bits & 1
                        or (member.external_attr >> 16) & 0o170000 == 0o120000):
                    raise ValueError("安全でない Excel アーカイブです")
                seen.add(name)
                if "vbaproject" in name.lower():
                    raise ValueError("マクロ付きファイルには対応していません")
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= seen:
                raise ValueError("Excel ブック構造を確認できません")
            if archive.testzip() is not None:
                raise ValueError("Excel アーカイブの CRC が一致しません")
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if not book.sheetnames:
                raise ValueError("ワークシートがありません")
        finally:
            book.close()
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("読み取れる .xlsx ファイルではありません") from exc
    return MEDIA[suffix]


def original_path(storage_dir: Path, document: dict) -> Path:
    root = storage_dir.resolve()
    candidate = root / document["storage_key"]
    if candidate.is_symlink() or not candidate.is_file() or candidate.resolve().parent != root:
        raise ValueError("保存した原本を利用できません")
    if candidate.stat().st_size != document["size_bytes"]:
        raise ValueError("原本のサイズが登録時から変わっています")
    if sha256(candidate.read_bytes()).hexdigest() != document["sha256"]:
        raise ValueError("原本の SHA256 が登録時から変わっています")
    return candidate


def _a1_bounds(cell: str) -> tuple[int, int, int, int]:
    """Only a physical A1 cell/range; no formulas, named ranges or sheet prefixes."""
    import re
    from openpyxl.utils.cell import range_boundaries
    atom = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}"
    if not isinstance(cell, str) or len(cell) > 40 or not re.fullmatch(atom + r"(?::" + atom + r")?", cell):
        raise ValueError("A1 または A1:B2 形式のセル範囲を指定してください")
    bounds = range_boundaries(cell.replace("$", "").upper())
    left, top, right, bottom = bounds
    if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
        raise ValueError("Excel の有効なセル範囲を指定してください")
    return bounds


def excel_preview(path: Path, filename: str, sheet_name: str | None = None,
                  cell: str | None = None) -> str:
    """Bounded original-cell preview. Location comes from the caller, never inferred."""
    if cell is not None and sheet_name is None:
        raise ValueError("セル範囲にはワークシート名を指定してください")
    selected = _a1_bounds(cell) if cell is not None else None
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    out = ['<!doctype html><html lang="ja"><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width,initial-scale=1">',
           '<title>Excel 原本プレビュー</title>',
           '<style>body{font:14px system-ui;margin:24px;color:#243d5a}table{border-collapse:collapse;margin-bottom:32px}td,th{border:1px solid #d3deeb;padding:8px;white-space:pre-wrap;max-width:350px}th{background:#edf3fc}.source-cell{background:#fff0b9;outline:2px solid #d1a032;outline-offset:-2px}h1{font-size:20px}p{line-height:1.7}</style>',
           '<h1>' + escape(filename) + '</h1>',
           '<p>原本のセル値を表示しています。このプレビューは SDK の抽出結果ではありません。'
           '数式は実行せず文字列で表示します。表示は各シート最大 200 行・30 列、'
           '全体表示では最大 20 シートです。書式・画像の完全再現には対応していません。</p>']
    try:
        sheets = book.worksheets
        if sheet_name is not None:
            sheets = [sheet for sheet in sheets if sheet.title == sheet_name]
            if not sheets:
                raise ValueError("指定したワークシートが原本にありません")
        else:
            sheets = sheets[:20]
        for sheet in sheets:
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            max_row, max_col = sheet.max_row or 1, sheet.max_column or 1
            if selected:
                left, top, right, bottom = selected
                if right > max_col or bottom > max_row:
                    raise ValueError("指定したセル範囲が原本ワークシートの範囲外です")
                row_start, col_start = max(1, top - 3), max(1, left - 2)
                row_end = min(max_row, bottom + 3, row_start + 199)
                col_end = min(max_col, right + 2, col_start + 29)
                out.append('<p>指定範囲：<strong>' + escape(sheet.title + "!" + cell) + '</strong>。黄色のセルが指定範囲です。</p>')
                if bottom > row_end or right > col_end:
                    out.append('<p>指定範囲が表示上限を超えるため、先頭付近のみ表示しています。</p>')
            else:
                row_start, col_start = 1, 1
                row_end, col_end = min(max_row, 200), min(max_col, 30)
            out.append('<h2>' + escape(sheet.title) + '</h2>')
            out.append('<p>表示範囲：' + get_column_letter(col_start) + str(row_start) + ':' + get_column_letter(col_end) + str(row_end) + '</p>')
            out.append('<table><thead><tr><th>行</th>')
            out.extend('<th>' + get_column_letter(i) + '</th>' for i in range(col_start, col_end + 1))
            out.append('</tr></thead><tbody>')
            for row_number, row in enumerate(sheet.iter_rows(min_row=row_start, max_row=row_end,
                                                              min_col=col_start, max_col=col_end), row_start):
                out.append('<tr><th>' + str(row_number) + '</th>')
                for col_number, value_cell in enumerate(row, col_start):
                    highlight = selected and selected[0] <= col_number <= selected[2] and selected[1] <= row_number <= selected[3]
                    coordinate = get_column_letter(col_number) + str(row_number)
                    value = str(value_cell.value)[:10000] if value_cell.value is not None else ''
                    out.append('<td' + (' class="source-cell"' if highlight else '') + ' data-cell="' + coordinate + '">' + escape(value) + '</td>')
                out.append('</tr>')
            out.append('</tbody></table>')
    finally:
        book.close()
    out.append('</html>')
    return ''.join(out)
