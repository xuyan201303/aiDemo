"""HTTP tests against a real disposable PostgreSQL database; manual fixture data only."""
from __future__ import annotations
from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sanzo_demo.app import create_app
from sanzo_demo.comparison import comparison_projection
from sanzo_core.demo_store import MANUAL_ACKNOWLEDGMENT

DSN = os.environ.get("SANZO_TEST_DSN")


def xlsx_bytes(label="送货钱", price=10):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["SYNTHETIC MANUAL REHEARSAL; NOT SDK/AI"])
    sheet.append(["Original label", "Current unit price", "Currency"])
    sheet.append([label, price, "JPY"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def pdf_bytes():
    """A tiny valid synthetic PDF, used only for upload/preview tests."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 500 300] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = b"BT /F1 16 Tf 30 240 Td (SYNTHETIC MANUAL REHEARSAL - NOT SDK / AI) Tj ET"
    objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    content = b"%PDF-1.4\n"
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(content))
        content += str(index).encode() + b" 0 obj\n" + value + b"\nendobj\n"
    xref = len(content)
    content += b"xref\n0 6\n0000000000 65535 f \n"
    content += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    content += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"
    return content


def manual_body(price="10", company="A"):
    return {
        "raw_label": "送货钱", "unit_price": price, "currency": "JPY", "forwarder_id": company,
        "context": {"transport_mode": "SEA", "equipment": "20FT", "billing_unit": "CONTAINER",
                    "charge_scope": "DESTINATION", "extra_terms": "NONE"},
        "source_acknowledgment": MANUAL_ACKNOWLEDGMENT,
    }


@unittest.skipUnless(DSN, "NOT_RUN: dedicated real PostgreSQL SANZO_TEST_DSN required")
class DemoHTTPTests(unittest.TestCase):
    def setUp(self):
        self.storage = tempfile.TemporaryDirectory(prefix="sanzo_http_test_")
        self.addCleanup(self.storage.cleanup)
        self.client = TestClient(create_app(DSN, self.storage.name, tenant_id="http-" + uuid4().hex, debug_manual=True))
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)

    def case(self, name="SYNTHETIC MANUAL REHEARSAL"):
        response = self.client.post("/api/cases", json={"name": name})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["case_id"]

    def upload(self, case, content=None, filename="synthetic.xlsx"):
        response = self.client.post(f"/api/cases/{case}/documents",
                                    files={"file": (filename, content or xlsx_bytes())})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["document_id"]

    def manual(self, document, body=None):
        response = self.client.post(f"/api/documents/{document}/manual-lines", json=body or manual_body())
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def correction(self, row, **overrides):
        data = {
            "expected_case_revision": row["case_revision"],
            "expected_item_version": row["version"],
            "expected_memory_version": row["current_memory_version"],
            "idempotency_key": str(uuid4()), "new_fee_name": "配送费",
        }
        data.update(overrides)
        return data

    def test_manual_input_is_disabled_in_default_application(self):
        with TestClient(create_app(DSN, self.storage.name,
                                  tenant_id="debug-disabled-" + uuid4().hex)) as client:
            self.assertFalse(client.get("/api/status").json()["debug_manual"])
            response = client.post("/api/documents/" + str(uuid4()) + "/manual-lines", json=manual_body())
            self.assertEqual(response.status_code, 403)

    def test_status_upload_preview_and_blocked_processing_are_honest(self):
        status = self.client.get("/api/status").json()
        self.assertFalse(status["sdk"]["ready"])
        self.assertFalse(status["ai"]["ready"])
        case = self.case()
        excel = self.upload(case)
        pdf = self.upload(case, pdf_bytes(), "synthetic.pdf")
        detail = self.client.get(f"/api/cases/{case}").json()
        self.assertEqual(len(detail["documents"]), 2)
        self.assertEqual(detail["items"], [])
        self.assertTrue(all(d["state"] == "RECEIVED" for d in detail["documents"]))
        for doc in detail["documents"]:
            self.assertNotIn("storage_key", doc)
        preview = self.client.get(f"/api/documents/{excel}/preview")
        self.assertEqual(preview.status_code, 200)
        self.assertIn("送货钱", preview.text)
        self.assertIn("text/html", preview.headers["content-type"])
        pdf_preview = self.client.get(f"/api/documents/{pdf}/preview")
        self.assertEqual(pdf_preview.status_code, 200)
        self.assertIn("text/html", pdf_preview.headers["content-type"])
        self.assertIn("pdf-page-image", pdf_preview.text)
        rendered = self.client.get(f"/api/documents/{pdf}/pages/1.png")
        self.assertEqual(rendered.status_code, 200)
        self.assertTrue(rendered.content.startswith(b"\x89PNG\r\n\x1a\n"))
        from PIL import Image
        with Image.open(BytesIO(rendered.content)) as image:
            self.assertGreater(image.width, 0)
            self.assertGreater(image.height, 0)
            image.verify()
        self.assertEqual(self.client.get(f"/api/documents/{pdf}/pages/0.png").status_code, 422)
        self.assertEqual(self.client.get(f"/api/documents/{pdf}/pages/2.png").status_code, 422)
        original = self.client.get(f"/api/documents/{pdf}/original")
        self.assertTrue(original.content.startswith(b"%PDF"))
        blocked = self.client.post(f"/api/documents/{excel}/process", json={"idempotency_key": str(uuid4())})
        self.assertEqual(blocked.status_code, 202)
        self.assertEqual(blocked.json()["status"], "BLOCKED")
        task = self.client.get("/api/tasks/" + blocked.json()["task_id"]).json()
        self.assertEqual(task["error_code"], "INTEGRATION_PREREQUISITES_MISSING")
        self.assertEqual(task["events"][-1]["metadata"]["ai_calls"], 0)
        snapshot = self.client.get(f"/api/cases/{case}").json()
        self.assertEqual(snapshot["tasks"][0]["status"], "BLOCKED")
        self.assertEqual(snapshot["tasks"][0]["issues"], [])
        self.assertNotIn("tenant_id", snapshot["tasks"][0])
        self.assertEqual(self.client.get(f"/api/cases/{case}").json()["items"], [])

    def test_two_rounds_use_saved_classification_and_current_price_with_export(self):
        first = self.case("Round 1 - SYNTHETIC MANUAL")
        row = self.manual(self.upload(first))
        self.assertEqual(row["saved_status"], "AI_REQUIRED")
        self.assertEqual(row["input_kind"], "MANUAL_CONFIRMED")
        self.assertIsNone(row["ai_suggested_name"])
        body = self.correction(row)
        result = self.client.post(f"/api/items/{row['item_id']}/corrections", json=body)
        self.assertEqual(result.status_code, 200, result.text)
        fee = result.json()["fee_id"]
        second = self.case("Round 2 - SYNTHETIC MANUAL")
        second_doc = self.upload(second, xlsx_bytes(price=20), "synthetic-followup.xlsx")
        current = self.manual(second_doc, manual_body("20"))
        self.assertEqual(current["saved_status"], "MEMORY_APPLIED")
        self.assertEqual(current["current_fee_id"], fee)
        self.assertEqual(current["current_fee_name"], "配送费")
        self.assertEqual(current["current_price"], "20")
        self.assertEqual(current["source_ref"]["document_id"], second_doc)
        old = self.client.get(f"/api/cases/{first}").json()["items"][0]
        self.assertEqual(old["current_price"], "10")
        self.assertEqual(old["current_fee_name"], "配送费")
        snapshot = self.client.get(f"/api/cases/{second}").json()
        self.assertEqual(snapshot["comparison"], comparison_projection(snapshot["items"], snapshot["fees"]))
        exported = self.client.get(f"/api/cases/{second}/export.xlsx")
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=False)
        try:
            values = [str(cell.value) for sheet in workbook for row in sheet for cell in row]
            self.assertIn("配送费", values)
            self.assertIn("20", values)
            self.assertIn(second_doc, values)
            projected = snapshot["comparison"]["rows"][0]
            self.assertEqual(workbook["原価単価一覧表"]["A4"].value, projected["label"] + "（" + projected["equipment"] + "）")
            self.assertEqual(workbook["原価単価一覧表"]["B4"].value, projected["cells"]["A"][0]["price"] + " " + projected["cells"]["A"][0]["currency"])
        finally:
            workbook.close()

    def test_correction_idempotency_content_conflict_and_stale_version(self):
        row = self.manual(self.upload(self.case()))
        body = self.correction(row)
        url = f"/api/items/{row['item_id']}/corrections"
        first = self.client.post(url, json=body)
        self.assertEqual(first.status_code, 200)
        self.assertEqual(self.client.post(url, json=body).json(), first.json())
        changed = dict(body, new_fee_name="Other")
        self.assertEqual(self.client.post(url, json=changed).status_code, 409)
        stale = dict(body, idempotency_key=str(uuid4()), new_fee_name="Other")
        self.assertEqual(self.client.post(url, json=stale).status_code, 409)
        row_after = self.client.get(f"/api/cases/{row['case_id']}").json()["items"][0]
        self.assertEqual(row_after["current_fee_name"], "配送费")
        self.assertEqual(row_after["version"], row["version"] + 1)

    def test_amount_only_does_not_learn_and_manual_requires_acknowledgment(self):
        case = self.case()
        doc = self.upload(case)
        invalid = manual_body()
        del invalid["source_acknowledgment"]
        self.assertEqual(self.client.post(f"/api/documents/{doc}/manual-lines", json=invalid).status_code, 422)
        self.assertEqual(self.client.get(f"/api/cases/{case}").json()["items"], [])
        row = self.manual(doc)
        change = self.correction(row, new_fee_name=None, unit_price="99")
        response = self.client.post(f"/api/items/{row['item_id']}/corrections", json=change)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["memory_effect"], "UNCHANGED_AMOUNT_ONLY")
        second = self.manual(self.upload(self.case()), manual_body("20"))
        self.assertEqual(second["saved_status"], "AI_REQUIRED")
        self.assertEqual(second["current_price"], "20")

    def test_scope_and_browser_request_boundaries(self):
        self.assertEqual(self.client.post("/api/cases", json={"name": "evil"},
                                         headers={"Origin": "https://evil.example"}).status_code, 403)
        self.assertEqual(self.client.post("/api/cases", json={"name": "evil"},
                                         headers={"Sec-Fetch-Site": "cross-site"}).status_code, 403)
        self.assertEqual(self.client.get("/api/status", headers={"Host": "evil.example"}).status_code, 400)
        self.assertEqual(self.client.get("/api/cases/nonexistent").status_code, 404)
        self.assertEqual(self.client.get("/api/documents/nonexistent/original").status_code, 404)
        self.assertEqual(self.client.post("/api/cases", json={"name": "ok", "tenant_id": "other"}).status_code, 422)

    def test_invalid_uploads_leave_no_registered_document_or_file(self):
        case = self.case()
        for name, content in [("bad.xlsx", b"not a workbook"), ("bad.pdf", b"not a pdf"),
                              ("macro.xlsm", b"no macros")]:
            response = self.client.post(f"/api/cases/{case}/documents", files={"file": (name, content)})
            self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(self.client.get(f"/api/cases/{case}").json()["documents"], [])
        self.assertEqual(list(Path(self.storage.name).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
