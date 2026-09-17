"""Offline generated-workbook tests; no customer samples, SDK or AI calls."""
from copy import deepcopy
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest

from openpyxl import Workbook, load_workbook

from sanzo_demo.files import excel_preview
from sanzo_demo.export import comparison_groups, workbook_bytes


class CellParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.cells = []
    def handle_starttag(self, tag, attrs):
        if tag == "td":
            self.cells.append(dict(attrs))


class SourcePreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.path = Path(self.temp.name) / "generated-preview.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "費用"
        sheet["A1"] = "<script>generated marker</script>"
        sheet["D205"] = "TARGET_MARKER"
        sheet["E206"] = "=1+1"
        sheet["AH230"] = "LAST_PHYSICAL_CELL"
        book.create_sheet("別紙")["B2"] = "OTHER_SHEET"
        book.save(self.path)
        book.close()
    def tearDown(self):
        self.temp.cleanup()

    def test_late_source_window_uses_real_coordinates_and_highlights_only_source(self):
        html = excel_preview(self.path, "<generated>.xlsx", sheet_name="費用", cell="$D$205:$E$206")
        parser = CellParser()
        parser.feed(html)
        selected = {cell["data-cell"] for cell in parser.cells if cell.get("class") == "source-cell"}
        self.assertEqual(selected, {"D205", "E205", "D206", "E206"})
        self.assertIn("TARGET_MARKER", html)
        self.assertIn("=1+1", html)
        self.assertIn("&lt;generated&gt;", html)
        self.assertNotIn("OTHER_SHEET", html)
        self.assertLessEqual(len(parser.cells), 200 * 30)

    def test_missing_and_out_of_bounds_locations_fail_closed(self):
        cases = [
            {"sheet_name": "不存在", "cell": "A1"},
            {"sheet_name": "費用", "cell": "AI1"},
            {"sheet_name": "費用", "cell": "D231"},
            {"sheet_name": "費用", "cell": "E206:D205"},
            {"sheet_name": "費用", "cell": "A0"},
            {"sheet_name": "費用", "cell": "XFE1"},
            {"sheet_name": "費用", "cell": "A1048577"},
            {"sheet_name": "費用", "cell": "A1;B2"},
            {"sheet_name": "費用", "cell": "=A1"},
            {"sheet_name": "費用", "cell": "別紙!A1"},
            {"cell": "A1"},
        ]
        for kwargs in cases:
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                excel_preview(self.path, "generated.xlsx", **kwargs)

    def test_large_selection_is_explicitly_truncated_and_html_is_escaped(self):
        html = excel_preview(self.path, "generated.xlsx", sheet_name="費用", cell="A1:AH230")
        parser = CellParser()
        parser.feed(html)
        self.assertEqual(len(parser.cells), 200 * 30)
        self.assertIn("先頭付近のみ表示", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)

    def test_book_is_closed_when_a_real_sheet_lookup_fails(self):
        book = load_workbook(self.path, read_only=True, data_only=False, keep_links=False)
        with patch("sanzo_demo.files.load_workbook", return_value=book):
            with patch.object(book, "close", wraps=book.close) as close:
                with self.assertRaises(ValueError):
                    excel_preview(self.path, "generated.xlsx", sheet_name="missing")
                close.assert_called_once()


def item(identity, company, *, price="20", fee_id=None, name="TEST_FEE"):
    return {
        "item_id": identity, "current_fee_id": fee_id,
        "current_fee_name": name if fee_id else None,
        "ai_suggested_name": None if fee_id else name,
        "current_price": price, "currency": "JPY", "version": 1,
        "input_kind": "SDK_OBSERVED", "saved_status": "HUMAN_CONFIRMED" if fee_id else "AI_SUGGESTED",
        "current_memory_status": "ACTIVE" if fee_id else None,
        "initial_classification": {"memory_version": 3} if fee_id else {},
        "source_ref": {"kind": "EXCEL_RANGE", "locator_status": "VERIFIED",
                       "locator": {"sheet_name": "費用", "A1_range": "D205:E206"}},
        "observation": {
            "document_id": "d", "forwarder_id": company, "raw_label": "TEST_RAW_" + identity,
            "source_ready": True,
            "context": {"ready": True, "transport_mode": "SEA", "equipment": "20FT",
                        "billing_unit": "CONTAINER", "charge_scope": "TEST_SCOPE", "extra_terms": "TEST_NONE"}
        },
    }


class SourceExportTests(unittest.TestCase):
    def export(self, items, *, tasks=None, fees=None, documents=None):
        return load_workbook(BytesIO(workbook_bytes(
            {"name": "=CASE", "case_revision": 7}, items,
            documents or [{"document_id": "d", "original_filename": "=source.xlsx",
                           "sha256": "a" * 64, "state": "BLOCKED", "blocked_reason": "SDK_NOT_CONFIGURED"}],
            tasks=tasks, fees=fees)), data_only=False)

    def test_actual_tasks_provenance_ai_column_and_exact_price_are_exported(self):
        row = item("a", "A", price="12345678901234567890.1234567890", name="=FORMULA_LIKE_NAME")
        before = deepcopy(row)
        book = self.export([row], tasks=[{"task_id": "t", "document_id": "d", "status": "AVAILABLE", "created_at": "2026-01-01"}])
        try:
            detail = book["明細と出典"]
            self.assertEqual(detail["E2"].value, row["current_price"])
            self.assertEqual(detail["E2"].data_type, "s")
            self.assertEqual(detail["S2"].value, "費用!D205:E206")
            self.assertEqual(detail["W2"].value, "=FORMULA_LIKE_NAME")
            self.assertEqual(detail["W2"].data_type, "s")
            self.assertEqual(detail["P2"].value, "=source.xlsx")
            self.assertEqual(detail["P2"].data_type, "s")
            self.assertEqual(detail["Z2"].value, "AVAILABLE")
            self.assertEqual(book["未確認事項"]["B3"].value, "AVAILABLE")
            values = [str(cell.value or "") for sheet in book for cells in sheet for cell in cells]
            self.assertNotIn("SDK_NOT_CONFIGURED", " ".join(values))
            self.assertNotIn("SDK・AI未接続", " ".join(values))
            self.assertIn("AI 提案・未確認", " ".join(values))
        finally:
            book.close()
        self.assertEqual(row, before)

    def test_latest_task_error_and_issue_are_preserved(self):
        tasks = [
            {"document_id": "d", "created_at": "2026-01-01", "state": "AVAILABLE"},
            {"document_id": "d", "created_at": "2026-01-02", "state": "FAILED",
             "error_code": "TEST_PARSE_FAILED", "issues": [{"code": "TEST_LOCATION_MISSING"}]},
        ]
        book = self.export([item("a", "A")], tasks=tasks)
        try:
            self.assertEqual(book["未確認事項"]["B3"].value, "FAILED")
            self.assertIn("TEST_PARSE_FAILED", book["未確認事項"]["C3"].value)
            self.assertIn("TEST_LOCATION_MISSING", book["未確認事項"]["C3"].value)
            self.assertIn("DRAFT", book["原価単価一覧表"]["A2"].value)
        finally:
            book.close()

    def test_pdf_page_only_and_unavailable_positions_remain_distinct(self):
        page = item("a", "A")
        page["source_ref"] = {"kind": "PDF_PAGE", "locator_status": "PAGE_ONLY", "locator": {"page_number": 2}}
        unknown = item("b", "B")
        unknown["source_ref"] = {"kind": "PDF_PAGE", "locator_status": "UNAVAILABLE", "locator": {"page_number": 99}}
        book = self.export([page, unknown])
        try:
            self.assertEqual(book["明細と出典"]["S2"].value, "PDF 2 ページ（ページ内範囲は未確定）")
            self.assertEqual(book["明細と出典"]["S3"].value, "原文位置は未取得")
        finally:
            book.close()

    def test_provisional_ai_comparison_preserves_two_companies_and_duplicate_quotes(self):
        rows = [item("a", "A", price="10"), item("b", "B", price="20"), item("c", "B", price="30")]
        before = deepcopy(rows)
        book = self.export(rows)
        try:
            comparison = book["原価単価一覧表"]
            self.assertEqual(comparison.max_row, 4)
            self.assertEqual(comparison["A4"].value, "TEST_FEE（20FT）")
            self.assertNotIn("AI仮判定", comparison["B4"].value)
            self.assertIn("10 JPY", comparison["B4"].value)
            self.assertIn("20 JPY", comparison["C4"].value)
            self.assertIn("30 JPY", comparison["C4"].value)
        finally:
            book.close()
        self.assertEqual(rows, before)

    def test_unique_confirmed_name_aligns_ai_but_ambiguous_names_never_merge_ids(self):
        ai, first, second = item("a", "A"), item("b", "B", fee_id="f1"), item("c", "C", fee_id="f2")
        self.assertEqual(len(comparison_groups([ai, first], [{"fee_id": "f1", "display_name": "TEST_FEE"}])), 1)
        self.assertIsNone(ai["current_fee_id"])
        self.assertEqual(len(comparison_groups([ai, first, second])), 3)

    def test_unknown_candidate_is_ambiguous_between_conflicting_equipment(self):
        a, b, c, d, e = [item(str(i), "A") for i in range(5)]
        b["observation"]["context"]["ready"] = False
        c["observation"]["context"]["equipment"] = None
        c["observation"]["forwarder_id"] = "B"
        e["observation"]["forwarder_id"] = "C"
        d["observation"]["forwarder_id"] = None
        e["observation"]["context"]["equipment"] = "40FT"
        groups = comparison_groups([a, b, c, d, e])
        self.assertEqual(len(groups), 4)
        self.assertTrue(next(g for g in groups if c in g["items"])["ambiguous"])
        self.assertFalse(any(c in g["items"] and e in g["items"] for g in groups))


if __name__ == "__main__":
    unittest.main(verbosity=2)
