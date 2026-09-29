"""R3 stable fee identity: isolated PostgreSQL with synthetic facts, never AI/SDK.

All writes use a unique per-test tenant in SANZO_TEST_DSN. These tests do not
connect to the running Demo database and never clear any schema or history.
The temporal probes deliberately learn both raw aliases in an earlier case,
then rename while only one company's aliases are present in the editing case.
"""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import os
from pathlib import Path
import tempfile
from threading import Barrier, Event
import unittest
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from sanzo_core.domain import (ConflictError, Context, CorrectionCommand,
    DecisionKind, Observation, Principal, RevokeCommand, ScopeError, SyntheticLine,
    canonical_json, memory_key)
from sanzo_core.demo_store import DemoStore
from sanzo_core.ingestion_store import IngestionStore
from sanzo_core.postgres import PostgresRepository
from sanzo_core.service import ReviewService
from sanzo_demo.comparison import comparison_projection
from sanzo_demo.app import create_app
from sanzo_demo.export import workbook_bytes
from sanzo_demo.row_confirmation import (RowConfirmationCommand,
    RowConfirmationService, RowMember)
import test_postgres as core_tests

DSN = os.environ.get("SANZO_TEST_DSN")
OLD_NAME = "書類作成手数料_v1"
NEW_NAME = "書類作成手数料_v2"


class FeeRenameIntentContractTests(unittest.TestCase):
    def command(self, **values):
        data = dict(case_id="case", row_id="a" * 64, expected_case_revision=1,
                    members=(RowMember("item", 1, 0),), idempotency_key="intent-probe",
                    new_fee_name=NEW_NAME, intent="RENAME_FEE",
                    expected_fee_id="stable-fee", expected_name_version=1)
        data.update(values)
        return RowConfirmationCommand(**data)

    def test_rename_requires_server_checked_identity_and_name_version(self):
        self.command().validate()
        for changes in ({"expected_fee_id": None}, {"expected_name_version": None},
                        {"expected_name_version": 0}, {"expected_name_version": True},
                        {"target_fee_id": "another", "new_fee_name": None}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                self.command(**changes).validate()

    def test_arbitrary_intent_is_rejected_instead_of_inferred_from_name(self):
        with self.assertRaises(ValueError):
            self.command(intent="RENAME_BY_SIMILARITY").validate()


class FeeRenameScopeContractTests(unittest.TestCase):
    """Offline boundary probes; marker dictionaries are not SDK read evidence."""
    def real_item(self):
        return {"current_fee_id": "stable-fee", "input_kind": "SDK_OBSERVED",
                "source_ref": {"origin": "REAL_SDK", "extraction_origin": "REAL_SDK"}}

    def mock_item(self):
        return {"current_fee_id": "stable-fee", "input_kind": "MOCK_SDK",
                "source_ref": {"origin": "MOCK_SDK", "extraction_origin": "MOCK_SDK",
                    "fixture_version": "OFFLINE_SCOPE_TEST", "fixture_sha256": "a" * 64,
                    "fixture_integrity": "VERIFIED"}}

    def test_mixed_REAL_Mock_row_cannot_rename_shared_identity(self):
        from sanzo_core.fee_names import rename_row
        owner = SimpleNamespace(repository=object())
        command = SimpleNamespace(expected_fee_id="stable-fee")
        with self.assertRaises(ScopeError):
            rename_row(owner, None, Principal("tenant", "human"), command,
                       {"real": self.real_item(), "mock": self.mock_item()}, "digest")

    def test_REAL_row_cannot_rename_identity_also_linked_from_active_Mock(self):
        from sanzo_core.fee_names import rename_row
        class ReadOnlyOriginRows:
            def execute(self, query, parameters):
                assert "SELECT DISTINCT origin" in query
                return self
            def fetchall(self):
                return [{"origin": "REAL_SDK"}, {"origin": "MOCK_SDK"}]
        owner = SimpleNamespace(repository=object())
        command = SimpleNamespace(expected_fee_id="stable-fee")
        with self.assertRaises(ScopeError):
            rename_row(owner, ReadOnlyOriginRows(), Principal("tenant", "human"), command,
                       {"real": self.real_item()}, "digest")


@unittest.skipUnless(DSN, "NOT_RUN: dedicated isolated PostgreSQL required")
class FeeRenamePostgresR3Tests(unittest.TestCase):
    database_snapshot = core_tests.PostgresIntegrationTests.snapshot

    def setUp(self):
        self.dsn = DSN
        self.tenant = "r3-fee-identity-test-" + uuid4().hex
        self.actor = Principal(self.tenant, "r3-explicit-human")
        self.repo = PostgresRepository(DSN)
        self.repo.migrate()
        self.store = DemoStore(self.repo)
        self.service = ReviewService(self.repo)
        self.rows = RowConfirmationService(self.repo)
        self.context = Context("SEA", "N/A", "PER_DOCUMENT", "ORIGIN", "STANDARD")
        self.counter = 0
        self.new_case("history")
        self.history_a = self.import_line("history", "A", "DOC FEE", "8000")
        self.history_b = self.import_line("history", "B", "DOCUMENTATION FEE", "7500.00")
        saved = self.rows.confirm(self.actor, self.command("history", new_fee_name=OLD_NAME))
        self.fee_id = saved["fee_id"]
        self.initial_history = self.snapshot("history")
        self.initial_export = self.export_values("history")
        self.initial_export_bytes = workbook_bytes(self.initial_history["case"],
            self.initial_history["items"], self.initial_history["documents"], fees=self.initial_history["fees"])

    def new_case(self, name):
        self.service.create_case(self.actor, name, "INJECTED_TEST_CASE_R3 " + name)
        return name

    def prepared_line(self, case, company, raw_label, price, *, context=None, currency="JPY"):
        self.counter += 1
        suffix = str(self.counter)
        observation = Observation(self.tenant, "item-" + suffix, case,
            "document-" + suffix, "source-" + suffix, "company-" + company,
            raw_label, price, currency, context or self.context, True)
        source = {"synthetic": True, "document_id": observation.document_id,
                  "document_revision": 1, "document_sha256": "a" * 64,
                  "source_block_ids": [observation.source_id], "kind": "EXCEL_RANGE",
                  "locator_status": "VERIFIED",
                  "locator": {"sheet_name": "Synthetic", "A1_range": "A2:F2"}}
        return SyntheticLine(observation=observation, source_ref=source,
                             raw_value=price, ai_suggested_name=OLD_NAME)

    def commit_line(self, line):
        obs = line.observation
        self.store.register_document(self.actor, obs.case_id, document_id=obs.document_id,
            original_filename=obs.document_id + ".xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=1, sha256="a" * 64, storage_key="synthetic/" + obs.document_id)
        return self.store.add_synthetic_lines(self.actor, obs.document_id, [line])[0]

    def import_line(self, case, company, raw_label, price, **values):
        return self.commit_line(self.prepared_line(case, company, raw_label, price, **values))

    def snapshot(self, case):
        return self.store.case_snapshot(self.actor, case)

    def projection(self, case):
        snap = self.snapshot(case)
        return comparison_projection(snap["items"], snap["fees"], snap["documents"])

    def current(self, item):
        return next(i for i in self.store.list_items(self.actor, item["case_id"])
                    if i["item_id"] == item["item_id"])

    def fee(self, fee_id=None):
        with self.repo.transaction() as conn:
            return self.repo.fee(conn, self.tenant, fee_id or self.fee_id)

    def name_versions(self, fee_id=None):
        with self.repo.transaction() as conn:
            return conn.execute("SELECT * FROM sanzo_core.fee_name_versions WHERE tenant_id=%s AND fee_id=%s ORDER BY name_version",
                                (self.tenant, fee_id or self.fee_id)).fetchall()

    def slot(self, item):
        item = self.current(item)
        obs = Observation.from_json(canonical_json(item["observation"]))
        with self.repo.transaction() as conn:
            return self.repo.lookup_memory(conn, obs, item["source_ref"], item["input_kind"])

    def command(self, case, *, selected_item=None, **changes):
        snap = self.snapshot(case)
        rows = comparison_projection(snap["items"], snap["fees"], snap["documents"])["rows"]
        if selected_item:
            row = next(row for row in rows if selected_item["item_id"] in row["item_ids"])
        else:
            self.assertEqual(len(rows), 1, rows)
            row = rows[0]
        values = dict(case_id=case, row_id=row["row_id"],
            expected_case_revision=snap["case"]["case_revision"],
            members=tuple(RowMember(i["item_id"], i["version"], i["current_memory_version"])
                          for i in row["member_items"]), idempotency_key=uuid4().hex,
            new_fee_name=OLD_NAME)
        values.update(changes)
        if values.get("target_fee_id") is not None:
            selected = next(fee for fee in snap["fees"] if fee["fee_id"] == values["target_fee_id"])
            values.setdefault("expected_fee_id", selected["fee_id"])
            values.setdefault("expected_name_version", selected["name_version"])
        return RowConfirmationCommand(**values)

    def rename_command(self, case, name=NEW_NAME, **changes):
        values = dict(new_fee_name=name, intent="RENAME_FEE", expected_fee_id=self.fee_id,
                      expected_name_version=self.fee()["name_version"])
        values.update(changes)
        return self.command(case, **values)

    def rename(self, case, name=NEW_NAME, **changes):
        return self.rows.confirm(self.actor, self.rename_command(case, name, **changes))

    def export_values(self, case):
        snap = self.snapshot(case)
        data = workbook_bytes(snap["case"], snap["items"], snap["documents"], fees=snap["fees"])
        book = load_workbook(BytesIO(data), read_only=True)
        try:
            return {name: list(book[name].values) for name in ("原価単価一覧表", "明細と出典")}
        finally:
            book.close()

    def assert_old_case_preserved(self):
        current = self.snapshot("history")
        for old, new in zip(self.initial_history["items"], current["items"]):
            for field in ("current_fee_id", "current_fee_name", "current_fee_name_snapshot",
                          "current_fee_name_version", "current_price", "currency", "source_ref",
                          "observation", "raw_value", "version"):
                self.assertEqual(old.get(field), new.get(field), field)
        # An existing exported artifact stays frozen. A newly requested workbook
        # uses the current catalog for its comparison sheet, but its item/source
        # detail sheet must retain each originally adopted name and source.
        old_book = load_workbook(BytesIO(self.initial_export_bytes), read_only=True)
        try:
            self.assertEqual({name: list(old_book[name].values) for name in self.initial_export},
                             self.initial_export)
        finally:
            old_book.close()
        exported = self.export_values("history")
        self.assertEqual(exported["明細と出典"], self.initial_export["明細と出典"])
        self.assertEqual(exported["原価単価一覧表"][3][0], self.fee()["display_name"])
        self.assertEqual(exported["原価単価一覧表"][3][1:], self.initial_export["原価単価一覧表"][3][1:])

    def assert_new_import(self, item, price):
        self.assertEqual(item["current_fee_id"], self.fee_id)
        self.assertEqual(item["current_fee_name"], NEW_NAME)
        self.assertEqual(item["current_fee_name_version"], 2)
        self.assertEqual((item["current_price"], item["currency"]), (price, "JPY"))
        self.assertEqual(item["saved_status"], "MEMORY_APPLIED")
        self.assertEqual(item["initial_classification"]["kind"], "MEMORY_APPLIED")

    def test_R3_01_A_only_rename_B_later_uses_same_identity_new_name(self):
        self.new_case("editing")
        a = self.import_line("editing", "A", "DOC FEE", "8500")
        before_b = deepcopy(self.slot(self.history_b))
        result = self.rename("editing")
        self.assertEqual(result["fee_id"], self.fee_id)
        self.assertEqual(self.slot(self.history_b), before_b, "Rename must not rewrite absent raw alias memory")
        b = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.assert_new_import(b, "7500.00")
        self.assertEqual(self.current(a)["current_fee_name"], NEW_NAME)
        self.assertEqual(len(self.projection("editing")["rows"]), 1)
        self.assert_old_case_preserved()

    def test_R3_02_reverse_B_only_rename_A_later(self):
        self.new_case("editing")
        self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.rename("editing")
        a = self.import_line("editing", "A", "DOC FEE", "8500")
        self.assert_new_import(a, "8500")
        self.assertEqual(len(self.projection("editing")["rows"]), 1)
        self.assert_old_case_preserved()

    def test_R3_03_two_A_same_alias_one_logical_rename_B_afterwards(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        self.import_line("editing", "A", "DOC FEE", "8500")
        before = self.name_versions()
        self.rename("editing")
        after = self.name_versions()
        self.assertEqual(len(after), len(before) + 1)
        self.assertEqual([r["name_version"] for r in after], [1, 2])
        b = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.assert_new_import(b, "7500.00")
        rows = self.projection("editing")["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(len(rows[0]["member_items"]), 3)
        self.assertEqual(sorted(i["current_price"] for i in rows[0]["member_items"]), ["7500.00", "8000", "8500"])

    def test_R3_04_B_prelookup_old_then_rename_commit_then_import_commit(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        prepared = self.prepared_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        with self.repo.transaction() as conn:
            cached_slot = self.repo.lookup_memory(conn, prepared.observation, prepared.source_ref, prepared.input_kind)
            cached_fee = self.repo.fee(conn, self.tenant, cached_slot["fee_id"])
        self.assertEqual(cached_fee["display_name"], OLD_NAME)
        self.rename("editing")
        late = self.commit_line(prepared)
        self.assertEqual(cached_fee["display_name"], OLD_NAME, "Probe really retained an old prelookup")
        self.assert_new_import(late, "7500.00")
        self.assertEqual(len(self.projection("editing")["rows"]), 1)

    def test_R3_05_both_present_rename_once_without_new_identity(self):
        self.new_case("editing")
        a = self.import_line("editing", "A", "DOC FEE", "8000")
        b = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        original_slots = [deepcopy(self.slot(item)) for item in (a, b)]
        result = self.rename("editing")
        self.assertEqual(result["fee_id"], self.fee_id)
        self.assertEqual([self.slot(item) for item in (a, b)], original_slots)
        self.assertEqual(len(self.name_versions()), 2)
        self.assertEqual({i["current_fee_name"] for i in self.snapshot("editing")["items"]}, {NEW_NAME})
        self.assert_old_case_preserved()

    def test_R3_06_reclassify_A_only_does_not_rename_or_reassign_B(self):
        self.new_case("editing")
        a = self.import_line("editing", "A", "DOC FEE", "8500")
        result = self.rows.confirm(self.actor, self.command("editing", new_fee_name="Different fee concept",
                                                         intent="RECLASSIFY_MEMBERS"))
        self.assertNotEqual(result["fee_id"], self.fee_id)
        self.assertEqual(self.fee()["display_name"], OLD_NAME)
        b = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.assertEqual((b["current_fee_id"], b["current_fee_name"]), (self.fee_id, OLD_NAME))
        self.assertEqual(self.slot(a)["fee_id"], result["fee_id"])
        self.assertEqual(len(self.projection("editing")["rows"]), 2)

    def test_R3_07_existing_other_fee_name_collision_never_silent_merge(self):
        self.new_case("other-fee")
        self.import_line("other-fee", "C", "ANOTHER RAW", "99")
        other = self.rows.confirm(self.actor, self.command("other-fee", new_fee_name=NEW_NAME))
        self.assertNotEqual(other["fee_id"], self.fee_id)
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        before = self.database_snapshot()
        with self.assertRaises((ConflictError, ValueError)):
            self.rename("editing")
        self.assertEqual(self.database_snapshot(), before)

    def test_R3_08_stale_name_version_rejected_even_if_item_case_memory_are_current(self):
        self.new_case("one")
        self.new_case("two")
        self.import_line("one", "A", "DOC FEE", "8000")
        self.import_line("two", "B", "DOCUMENTATION FEE", "7500.00")
        stale = self.rename_command("two", "Stale overwrite")
        self.rename("one")
        before = self.database_snapshot()
        with self.assertRaises(ConflictError):
            self.rows.confirm(self.actor, stale)
        self.assertEqual(self.database_snapshot(), before)

    def test_R3_09_replayed_request_is_identical_changed_payload_rejected(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        command = self.rename_command("editing")
        result = self.rows.confirm(self.actor, command)
        before = self.database_snapshot()
        self.assertEqual(self.rows.confirm(self.actor, command), result)
        self.assertEqual(self.database_snapshot(), before)
        with self.assertRaises(ConflictError):
            self.rows.confirm(self.actor, replace(command, new_fee_name="Different payload"))
        self.assertEqual(self.database_snapshot(), before)

    def test_R3_10_concurrent_name_updates_have_one_winner(self):
        self.new_case("one")
        self.new_case("two")
        self.import_line("one", "A", "DOC FEE", "8000")
        self.import_line("two", "B", "DOCUMENTATION FEE", "7500.00")
        commands = [self.rename_command("one", "Winner one"), self.rename_command("two", "Winner two")]
        barrier = Barrier(2)
        def save(command):
            barrier.wait(timeout=10)
            try:
                return RowConfirmationService(PostgresRepository(DSN)).confirm(self.actor, command)
            except ConflictError:
                return "CONFLICT"
        with ThreadPoolExecutor(2) as executor:
            results = list(executor.map(save, commands))
        self.assertEqual(sum(isinstance(result, dict) for result in results), 1, results)
        self.assertEqual(results.count("CONFLICT"), 1, results)
        self.assertEqual(len(self.name_versions()), 2)
        self.assert_old_case_preserved()

    def test_R3_11_failure_after_name_write_rolls_back_all_changes(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        before = self.database_snapshot()
        seen = []
        def fail(stage):
            seen.append(stage)
            if stage == "after_fee_rename":
                raise RuntimeError("INJECTED_TEST_CASE_R3_RENAME_ROLLBACK")
        with self.assertRaisesRegex(RuntimeError, "R3_RENAME_ROLLBACK"):
            RowConfirmationService(self.repo, fault_hook=fail).confirm(self.actor, self.rename_command("editing"))
        self.assertIn("after_fee_rename", seen)
        self.assertEqual(self.database_snapshot(), before)

    def test_R3_12_revoked_absent_B_alias_is_not_restored_by_rename(self):
        b = self.current(self.history_b)
        self.service.revoke_memory(self.actor, RevokeCommand(b["item_id"], b["case_revision"], b["version"],
            b["current_memory_version"], uuid4().hex, "INJECTED_TEST_CASE_R3 revoke B"))
        revoked = deepcopy(self.slot(self.history_b))
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        self.rename("editing")
        later = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.assertIsNone(later["current_fee_id"])
        self.assertEqual(later["current_memory_status"], "REVOKED")
        self.assertEqual(self.slot(self.history_b), revoked)
        self.assertEqual(self.service.classify(self.actor, later["item_id"]).kind, DecisionKind.REVIEW_REQUIRED)

    def test_R3_13_conflicted_absent_B_alias_is_not_restored_by_rename(self):
        b = self.current(self.history_b)
        result = self.service.correct(self.actor, CorrectionCommand(item_id=b["item_id"], expected_item_version=b["version"],
            expected_case_revision=b["case_revision"], expected_memory_version=b["current_memory_version"],
            idempotency_key=uuid4().hex, new_fee_name="Conflicting old API target"))
        self.assertEqual(result["memory_effect"], "CONFLICTED")
        conflicted = deepcopy(self.slot(self.history_b))
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        self.rename("editing")
        later = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.assertIsNone(later["current_fee_id"])
        self.assertEqual(later["current_memory_status"], "CONFLICTED")
        self.assertEqual(self.slot(self.history_b), conflicted)

    def test_R3_14_previously_reclassified_B_alias_stays_other_identity(self):
        self.new_case("reclassify-b")
        b = self.import_line("reclassify-b", "B", "DOCUMENTATION FEE", "7500.00")
        changed = self.rows.confirm(self.actor, self.command("reclassify-b", new_fee_name="Different B meaning",
                                                          intent="RECLASSIFY_MEMBERS"))
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        self.rename("editing")
        later = self.import_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        self.assertEqual(later["current_fee_id"], changed["fee_id"])
        self.assertEqual(later["current_fee_name"], "Different B meaning")
        self.assertNotEqual(later["current_fee_id"], self.fee_id)

    def test_R3_15_new_price_source_restart_and_old_export_stay_separate(self):
        self.new_case("editing")
        a = self.import_line("editing", "A", "DOC FEE", "8000")
        self.rename("editing")
        self.repo = PostgresRepository(DSN)
        self.store = DemoStore(self.repo)
        self.service = ReviewService(self.repo)
        self.new_case("after-restart")
        later = self.import_line("after-restart", "A", "DOC FEE", "8500")
        self.assert_new_import(later, "8500")
        self.assertNotEqual(later["source_ref"]["document_id"], a["source_ref"]["document_id"])
        self.assertEqual(self.current(a)["current_price"], "8000")
        self.assert_old_case_preserved()
        labels = [row[0] for row in self.export_values("after-restart")["原価単価一覧表"]]
        self.assertIn(NEW_NAME, labels)
        self.assertNotIn(OLD_NAME, labels)

    def test_R3_16_changed_reliable_conditions_do_not_become_wildcards(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        self.rename("editing")
        for number, context in enumerate((replace(self.context, equipment="20FT"),
                                          replace(self.context, equipment="40FT"),
                                          replace(self.context, billing_unit="PER_CONTAINER"),
                                          replace(self.context, equipment=None, ready=False))):
            case = self.new_case("scope-" + str(number))
            item = self.import_line(case, "A", "DOC FEE", "8500", context=context)
            self.assertIsNone(item["current_fee_id"], context)
            self.assertNotEqual(self.service.classify(self.actor, item["item_id"]).kind, DecisionKind.MEMORY_APPLIED)

    def test_R3_17_wrong_identity_and_cross_tenant_rename_are_rejected(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        before = self.database_snapshot()
        with self.assertRaises((ConflictError, ScopeError, ValueError)):
            self.rows.confirm(self.actor, self.rename_command("editing", expected_fee_id="not-this-fee"))
        self.assertEqual(self.database_snapshot(), before)
        outsider = Principal("r3-outsider-" + uuid4().hex, "other")
        with self.assertRaises(ScopeError):
            self.rows.confirm(outsider, self.rename_command("editing"))
        self.assertEqual(self.database_snapshot(), before)


    def test_R3_18_mixed_adopted_versions_share_current_row_without_rewriting_history(self):
        self.new_case("historical-view")
        old_a = self.import_line("historical-view", "A", "DOC FEE", "8000")
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8500")
        self.rename("editing")
        new_b = self.import_line("historical-view", "B", "DOCUMENTATION FEE", "7500.00")
        retained_a = self.current(old_a)
        self.assertEqual(retained_a["current_fee_id"], new_b["current_fee_id"])
        self.assertEqual((retained_a["current_fee_name"], retained_a["current_fee_name_version"]), (OLD_NAME, 1))
        self.assertEqual((new_b["current_fee_name"], new_b["current_fee_name_version"]), (NEW_NAME, 2))
        self.assertEqual(retained_a["effective_fee_name"], NEW_NAME)
        self.assertEqual(retained_a["effective_fee_name_version"], 2)
        before = self.database_snapshot()
        rows = self.projection("historical-view")["rows"]
        self.assertEqual(len(rows), 1, "Adopted name versions are audit facts, not current business row identities")
        self.assertEqual({row["display_name"] for row in rows}, {NEW_NAME})
        self.assertEqual({row["fee_identity"] for row in rows}, {self.fee_id})
        self.assertEqual({item["current_fee_name_version"] for item in rows[0]["member_items"]}, {1, 2})
        values = self.export_values("historical-view")["原価単価一覧表"]
        self.assertEqual([row[0] for row in values[3:]], [NEW_NAME])
        self.assertEqual(values[3][1:3], ("8000 JPY", "7500.00 JPY"))
        self.assertEqual(self.current(old_a), retained_a)
        self.assertEqual(self.database_snapshot(), before, "Projection/export must not write snapshots")

    def test_R3_21_current_http_and_xlsx_merge_40ft_versions_without_history_writes(self):
        self.new_case("active-freight")
        twenty = Context("SEA", "20FT", "PER_CONTAINER", "ORIGIN", "STANDARD")
        forty = replace(twenty, equipment="40FT")
        a20 = self.import_line("active-freight", "A", "OCEAN 20", "1200", context=twenty, currency="USD")
        self.import_line("active-freight", "B", "FREIGHT 20", "1150", context=twenty, currency="USD")
        a40 = self.import_line("active-freight", "A", "OCEAN 40", "1900", context=forty, currency="USD")
        b40 = self.import_line("active-freight", "B", "FREIGHT 40", "1850", context=forty, currency="USD")
        first = self.rows.confirm(self.actor, self.command("active-freight", selected_item=a20,
                                                         new_fee_name="海上運賃"))
        fee_id = first["fee_id"]
        self.rows.confirm(self.actor, self.command("active-freight", selected_item=a40,
            new_fee_name=None, target_fee_id=fee_id, intent="CONFIRM_MEMBERS"))
        adopted_forty = deepcopy([self.current(a40), self.current(b40)])
        old_snap = self.snapshot("active-freight")
        old_xlsx = workbook_bytes(old_snap["case"], old_snap["items"], old_snap["documents"], fees=old_snap["fees"])
        self.rows.confirm(self.actor, self.command("active-freight", selected_item=a20,
            new_fee_name="海上運賃_v1", intent="RENAME_FEE", expected_fee_id=fee_id, expected_name_version=1))
        # RENAME_FEE remains scoped to explicitly confirmed row members. The
        # current read must not require bulk adoption by the other 40FT row.
        for old in adopted_forty:
            current = self.current(old)
            for field in ("current_fee_name", "current_fee_name_version", "version", "observation",
                          "source_ref", "raw_value", "current_price", "currency"):
                self.assertEqual(current[field], old[field], field)
            self.assertEqual((current["effective_fee_name"], current["effective_fee_name_version"]),
                             ("海上運賃_v1", 2))
        later20 = self.import_line("active-freight", "A", "OCEAN 20", "1350", context=twenty, currency="USD")
        later40 = self.import_line("active-freight", "A", "OCEAN 40", "2050", context=forty, currency="USD")
        self.assertEqual((later40["current_fee_id"], later40["current_fee_name_version"]), (fee_id, 2))
        self.assertEqual((later20["current_fee_id"], later20["current_fee_name_version"]), (fee_id, 2))
        before = self.database_snapshot()
        with tempfile.TemporaryDirectory(prefix="sanzo-projection-") as storage:
            historical_file = Path(storage) / "already-generated.xlsx"
            historical_file.write_bytes(old_xlsx)
            # No real SDK/AI pipeline is constructed or invoked by this GET-only
            # HTTP probe. It uses the actual store and the isolated test tenant.
            runtime = SimpleNamespace(store=IngestionStore(self.repo), reader_mode="SYNTHETIC_STRUCTURED")
            with TestClient(create_app(DSN, storage, tenant_id=self.tenant, integrations=runtime)) as client:
                response = client.get("/api/cases/active-freight")
                self.assertEqual(response.status_code, 200)
                payload = response.json()
                projection = payload["comparison"]
                exported = client.get("/api/cases/active-freight/export.xlsx")
                self.assertEqual(exported.status_code, 200)
            self.assertEqual(historical_file.read_bytes(), old_xlsx)
        self.assertEqual(len(projection["rows"]), 2)
        by_equipment = {row["equipment"]: row for row in projection["rows"]}
        for equipment, prices in (("20FT", ["1200", "1150", "1350"]),
                                  ("40FT", ["1900", "1850", "2050"])):
            row = by_equipment[equipment]
            self.assertEqual(row["display_name"], "海上運賃_v1（" + equipment + "）")
            self.assertEqual(row["fee_identity"], fee_id)
            self.assertCountEqual([cell["price"] for cells in row["cells"].values() for cell in cells], prices)
        book = load_workbook(BytesIO(exported.content), read_only=True)
        try:
            sheet = book["原価単価一覧表"]
            self.assertEqual(sheet.max_row - 3, len(projection["rows"]))
            for number, row in enumerate(projection["rows"], 4):
                self.assertEqual(sheet.cell(number, 1).value, row["display_name"])
                for column, company in enumerate(projection["companies"], 2):
                    expected = "\n".join(cell["price"] + " " + cell["currency"] for cell in row["cells"][company]) or "—"
                    self.assertEqual(sheet.cell(number, column).value, expected)
            adopted = {row[0]: row[3] for row in list(book["明細と出典"].values)[1:]}
            self.assertEqual(adopted[a40["item_id"]], "海上運賃")
            self.assertEqual(adopted[later40["item_id"]], "海上運賃_v1")
        finally:
            book.close()
        self.assertEqual(self.database_snapshot(), before,
            "GET/projection/export must preserve all items, correction events, memory, and fee-name history")

    def test_R3_19_legacy_amount_only_correction_keeps_old_adopted_name(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8500")
        self.rename("editing")
        old = self.current(self.history_a)
        old_slot = deepcopy(self.slot(old))
        history = deepcopy(self.name_versions())
        self.assertEqual((old["current_fee_name"], old["effective_fee_name"]), (OLD_NAME, NEW_NAME))
        # Existing low-level price API is tested only in this isolated tenant;
        # the ordinary customer UI still has no price editing control.
        result = self.service.correct(self.actor, CorrectionCommand(item_id=old["item_id"],
            expected_case_revision=old["case_revision"], expected_item_version=old["version"],
            expected_memory_version=old["current_memory_version"], idempotency_key=uuid4().hex,
            unit_price="8001"))
        current = self.current(old)
        self.assertEqual(result["memory_effect"], "UNCHANGED_AMOUNT_ONLY")
        self.assertEqual((current["current_fee_name"], current["current_fee_name_version"]), (OLD_NAME, 1))
        self.assertEqual(current["current_price"], "8001")
        self.assertEqual(current["source_ref"], old["source_ref"])
        self.assertEqual(current["observation"], old["observation"])
        self.assertEqual(current["raw_value"], old["raw_value"])
        self.assertEqual(self.slot(current), old_slot)
        self.assertEqual(self.name_versions(), history)

    def test_R3_20_dual_connection_old_prelookup_then_rename_then_late_commit(self):
        self.new_case("editing")
        self.import_line("editing", "A", "DOC FEE", "8000")
        prepared = self.prepared_line("editing", "B", "DOCUMENTATION FEE", "7500.00")
        obs = prepared.observation
        document = self.store.register_document(self.actor, obs.case_id,
            document_id=obs.document_id, original_filename=obs.document_id + ".xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=1, sha256="a" * 64, storage_key="synthetic/" + obs.document_id)
        old_read = Event()
        rename_committed = Event()
        observed = {}
        def delayed_insert():
            repository = PostgresRepository(DSN)
            store = DemoStore(repository)
            # Hold this real PostgreSQL connection open across another
            # connection's committed rename. READ COMMITTED must reread the
            # active name immediately before persisting the imported item.
            with repository.transaction() as conn:
                observed["reader_pid"] = conn.execute("SELECT pg_backend_pid() AS pid").fetchone()["pid"]
                slot = repository.lookup_memory(conn, obs, prepared.source_ref, prepared.input_kind)
                observed["old_fee"] = repository.fee(conn, self.tenant, slot["fee_id"])
                old_read.set()
                if not rename_committed.wait(timeout=15):
                    raise RuntimeError("INJECTED_TEST_CASE_R3 missing rename completion")
                repository.case(conn, self.tenant, obs.case_id, lock=True)
                item_id = store._insert_line(conn, self.actor, document, prepared)
                repository.advance_case(conn, self.tenant, obs.case_id)
                return store._item(conn, self.actor, item_id)
        with ThreadPoolExecutor(1) as executor:
            future = executor.submit(delayed_insert)
            try:
                self.assertTrue(old_read.wait(timeout=15), "Worker must read the pre-rename version")
                self.assertEqual(observed["old_fee"]["display_name"], OLD_NAME)
                with self.repo.transaction() as independent:
                    other_pid = independent.execute("SELECT pg_backend_pid() AS pid").fetchone()["pid"]
                self.assertNotEqual(observed["reader_pid"], other_pid)
                self.rename("editing")
            finally:
                rename_committed.set()
            late = future.result(timeout=20)
        self.assert_new_import(late, "7500.00")
        self.assertEqual(self.current(late)["current_fee_name_version"], 2)
        self.assertEqual(len(self.projection("editing")["rows"]), 1)
        self.assert_old_case_preserved()
