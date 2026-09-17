"""Two real browser pages, real FastAPI row endpoints, isolated PostgreSQL.

Only the explicitly selected localhost test port is allowed. No response is
fabricated. The lost-response case commits through route.fetch(), then discards
that response before the exact original request is replayed by the actual UI.
"""
import argparse
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import traceback
from urllib.parse import urlsplit

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


def run(output, base):
    parsed = urlsplit(base)
    if parsed.scheme != 'http' or parsed.hostname != '127.0.0.1' or parsed.port in {None, 8765} or parsed.path:
        raise ValueError('An explicit isolated localhost test port is required; production8765 is forbidden')
    output.mkdir(parents=True, exist_ok=False)
    result = {'scope': 'TWO_BROWSER_PAGES_REAL_APP_ISOLATED_POSTGRES', 'status': 'RUNNING',
              'checks': [], 'writes': [], 'http_trace': [], 'ai_calls': 0, 'sdk_calls': 0}
    pages = []

    def dump(name, value):
        (output / (name + '.json')).write_text(json.dumps(value, ensure_ascii=False, indent=2))

    def check(label, condition=True):
        assert condition, label
        result['checks'].append(label)
        print('PASS ' + label, flush=True)
        dump('progress', result)

    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome', headless=True)
            context = browser.new_context(viewport={'width': 1366, 'height': 768})
            errors = []
            a, b = context.new_page(), context.new_page()
            pages = [a, b]
            for page, label in [(a, 'A'), (b, 'B')]:
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.on('request', lambda request, tab=label: result['http_trace'].append({'tab': tab, 'method': request.method, 'path': urlsplit(request.url).path}))

            def guard(route):
                request = route.request
                if not request.url.startswith(base + '/'):
                    raise AssertionError('Non-test origin prohibited')
                path = urlsplit(request.url).path
                if request.method != 'GET':
                    assert request.method == 'POST' and '/comparison/' in path and path.endswith('/confirmation'), path
                assert not path.endswith('/process') and not path.endswith('/field-sources'), path
                route.continue_()

            context.route('**/*', guard)

            def get(path):
                response = a.request.get(base + path)
                assert response.ok, (path, response.status)
                return response.json()

            def audit(tag):
                data = get('/__test/audit')
                dump(tag, data)
                return data

            def open_case(page, case_id):
                page.goto(base, wait_until='networkidle')
                page.locator('#case-select').select_option(case_id)
                page.locator('.row-confirm-button').first.click()
                expect(page.locator('#edit-dialog')).to_be_visible()
                return get('/api/cases/' + case_id)

            def submit(page, tag):
                start = len(result['http_trace'])
                with page.expect_response(lambda response: response.request.method == 'POST' and response.url.endswith('/confirmation')) as awaited:
                    page.locator('#save-correction').click()
                response = awaited.value
                record = {'tag': tag, 'status': response.status, 'request': response.request.post_data_json, 'response': response.json()}
                result['writes'].append(record)
                own_tab = 'A' if page == a else 'B'
                trace = [entry for entry in result['http_trace'][start:] if entry['tab'] == own_tab]
                first_post = next(index for index, entry in enumerate(trace) if entry['method'] == 'POST')
                record['before_save_trace'] = trace[:first_post + 1]
                check(tag + ' sends POST without save-time GET', all(entry['method'] != 'GET' for entry in trace[:first_post]))
                dump(tag + '_response', record)
                return record

            def rename(case_id, name, tag):
                snapshot = open_case(a, case_id)
                a.locator('#edit-fee-name').fill(name)
                record = submit(a, tag)
                check(tag + ' real rename commits', record['status'] == 200 and record['request']['intent'] == 'RENAME_FEE')
                expect(a.locator('#edit-dialog')).not_to_be_visible()
                return record, snapshot

            def stale_b(tag, before, expected_id, expected_version):
                record = submit(b, tag)
                after = audit(tag + '_audit_after')
                # Capture actual persisted state before asserting, so RED records a
                # real200/write if the vulnerable implementation accepts the form.
                check(tag + ' stale name rejected409', record['status'] == 409)
                check(tag + ' rejected save writes no table', after == before)
                check(tag + ' payload carries originally seen target identity/version', record['request'].get('expected_fee_id') == expected_id and record['request'].get('expected_name_version') == expected_version)
                expect(b.locator('#edit-dialog')).to_be_visible()
                expect(b.locator('#edit-error')).to_contain_text('他の更新')
                b.screenshot(path=str(output / (tag + '.png')))
                return record

            def refresh_and_choose(expected_draft, expected_name, fee_id, tag):
                b.locator('#edit-details > summary').click()
                b.locator('#refresh-item').click()
                expect(b.locator('#edit-fee-name')).to_have_value(expected_draft)
                expect(b.locator('#fee-candidates')).to_be_visible()
                option = b.locator('[data-fee-id="' + fee_id + '"]')
                expect(option).to_contain_text(expected_name)
                expect(b.locator('#edit-intent-help')).to_contain_text(expected_name)
                check(tag + ' refresh displays latest candidate and keeps draft')
                previous_writes = len(result['writes'])
                request_count = len([entry for entry in result['http_trace'] if entry['tab'] == 'B' and entry['method'] == 'POST'])
                b.locator('#save-correction').click()
                expect(b.locator('#edit-error')).to_contain_text('候補')
                b.wait_for_timeout(100)
                check(tag + ' unchanged stale draft cannot submit with refreshed version', len(result['writes']) == previous_writes and len([entry for entry in result['http_trace'] if entry['tab'] == 'B' and entry['method'] == 'POST']) == request_count)
                option.click()
                expect(b.locator('#edit-fee-name')).to_have_value(expected_name)
                return submit(b, tag + '_reviewed_save')

            baseline = audit('baseline')
            check('No AI tasks or ledger in isolated synthetic scenario', not baseline.get('ai_call_attempts') and not baseline.get('ingestion_tasks'))
            old_b = open_case(b, 'confirm-b')
            target = old_b['items'][0]['current_fee_id']
            seen_version = old_b['items'][0]['effective_fee_name_version']
            expect(b.locator('#edit-fee-name')).to_have_value('確認対象 X')
            check('Existing business dialog prices remain readonly', b.locator('#edit-dialog input:visible').count() == 1 and b.locator('#edit-price,#edit-currency,#amount-only').count() == 0)
            a.screenshot(path=str(output / 'before_a_navigation.png'))
            rename('confirm-a', '確認対象 Y', 'confirm_other_tab_rename')
            before = audit('confirm_stale_audit_before')
            stale_b('confirm_stale', before, target, seen_version)
            expect(b.locator('#edit-fee-name')).to_have_value('確認対象 X')
            reviewed = refresh_and_choose('確認対象 X', '確認対象 Y', target, 'confirm_refresh')
            check('Reviewed confirm uses latest name version', reviewed['status'] == 200 and reviewed['request']['intent'] == 'CONFIRM_MEMBERS' and reviewed['request']['expected_name_version'] == seen_version + 1)
            new_b = get('/api/cases/confirm-b')
            check('Fresh confirmation keeps original amount currency source', all(old_b['items'][0][key] == new_b['items'][0][key] for key in ('current_price', 'currency', 'observation', 'source_ref')) and new_b['items'][0]['current_fee_name'] == '確認対象 Y')
            b.screenshot(path=str(output / 'confirm_saved.png'))

            # A different existing classification target requires its own viewed version.
            source_b = open_case(b, 'reclass-b')
            target_info = next(fee for fee in source_b['fees'] if fee['display_name'] == '分類先 X')
            b.locator('#edit-fee-name').fill('分類先 X')
            b.locator('[data-fee-id="' + target_info['fee_id'] + '"]').click()
            rename('reclass-a', '分類先 Y', 'reclass_other_tab_rename')
            before = audit('reclass_stale_audit_before')
            rejected = stale_b('reclass_stale', before, target_info['fee_id'], target_info['name_version'])
            check('Reclassification retains explicit member-only intent', rejected['request']['intent'] == 'RECLASSIFY_MEMBERS')
            reviewed = refresh_and_choose('分類先 X', '分類先 Y', target_info['fee_id'], 'reclass_refresh')
            check('Reviewed reclassification adopts current target only', reviewed['status'] == 200 and reviewed['request']['intent'] == 'RECLASSIFY_MEMBERS' and reviewed['request']['expected_name_version'] == target_info['name_version'] + 1)
            new_source = get('/api/cases/reclass-b')
            check('Reclassification keeps price and original evidence', all(source_b['items'][0][key] == new_source['items'][0][key] for key in ('current_price', 'currency', 'observation', 'source_ref')) and new_source['items'][0]['current_fee_id'] == target_info['fee_id'])

            # X -> Y -> X must still reject the old version even if the name matches.
            old_cycle = open_case(b, 'cycle-b')
            cycle_id = old_cycle['items'][0]['current_fee_id']
            rename('cycle-a', '循環対象 Y', 'cycle_to_y')
            rename('cycle-a', '循環対象 X', 'cycle_back_x')
            # Exercise a real background snapshot read while the old dialog remains open.
            b.evaluate('loadSnapshot()')
            expect(b.locator('#edit-fee-name')).to_have_value('循環対象 X')
            before = audit('cycle_stale_audit_before')
            stale_b('cycle_stale', before, cycle_id, 1)
            reviewed = refresh_and_choose('循環対象 X', '循環対象 X', cycle_id, 'cycle_refresh')
            check('Same name reviewed again uses version3', reviewed['status'] == 200 and reviewed['request']['expected_name_version'] == 3)

            # Commit a real request, lose only its HTTP response, rename elsewhere,
            # then prove exact stored response replay is allowed before current-version checks.
            old_replay = open_case(b, 'replay-b')
            captured = []
            def lose_once(route):
                actual = route.fetch()
                captured.append({'request': route.request.post_data_json, 'status': actual.status, 'response': actual.json()})
                route.abort('connectionreset')
            b.route('**/api/cases/replay-b/comparison/*/confirmation', lose_once, times=1)
            b.locator('#save-correction').click()
            expect(b.locator('#edit-error')).to_contain_text('保存結果を確認できません')
            check('Lost response came from actual committed200', len(captured) == 1 and captured[0]['status'] == 200)
            expect(b.locator('#edit-fee-name')).to_be_disabled()
            rename('replay-a', '再送対象 Y', 'replay_other_tab_rename')
            before = audit('replay_after_rename_before_retry')
            replayed = submit(b, 'exact_replay_after_rename')
            check('Exact stale-version request replays saved success', replayed['status'] == 200 and replayed['request'] == captured[0]['request'] and replayed['response'] == captured[0]['response'])
            check('Replay creates no new event or memory', audit('replay_after_retry') == before)
            replay_after = get('/api/cases/replay-b')
            check('Replay keeps originally adopted name and amount', replay_after['items'][0]['current_fee_name'] == '再送対象 X' and replay_after['items'][0]['current_price'] == old_replay['items'][0]['current_price'])
            dump('lost_response_actual_commit', captured)

            # Existing export and preview continue to read stored facts.
            open_case(b, 'confirm-b')
            b.locator('.member-preview').first.click()
            expect(b.frame_locator('#preview-frame').locator('[data-cell="A2"]')).to_contain_text('SYNTHETIC ORIGINAL confirm-b')
            b.screenshot(path=str(output / 'source_preview.png'))
            b.locator('#preview-dialog .close-dialog').click()
            expect(b.locator('#edit-fee-name')).to_have_value('確認対象 Y')
            exported = a.request.get(base + '/api/cases/confirm-b/export.xlsx')
            check('Existing export succeeds', exported.ok)
            book = load_workbook(BytesIO(exported.body()), data_only=True)
            values = [book['原価単価一覧表'].cell(4, column).value for column in (1, 2)]
            check('Export shows saved name and source price', values == ['確認対象 Y（TRUCK）', '9500 JPY'])
            dump('export_check', {'sha256': sha256(exported.body()).hexdigest(), 'row': values})
            book.close()
            check('No browser JS errors', not errors)
            result['status'] = 'PASS'
            browser.close()
    except Exception:
        result['status'] = 'FAIL'
        result['error'] = traceback.format_exc()
        print(result['error'], flush=True)
        for index, page in enumerate(pages):
            try:
                page.screenshot(path=str(output / ('failed_tab_' + str(index) + '.png')))
            except Exception:
                pass
    finally:
        dump('result', result)
    return 0 if result['status'] == 'PASS' else 1


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--evidence', type=Path, required=True)
    parser.add_argument('--base-url', default='http://127.0.0.1:28765')
    args = parser.parse_args()
    raise SystemExit(run(args.evidence, args.base_url.rstrip('/')))
