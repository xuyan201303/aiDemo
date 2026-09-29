"""Offline comparison projection probes; no sample answers, database or AI calls."""
from copy import deepcopy
from io import BytesIO
from itertools import permutations
import unittest

from openpyxl import load_workbook

from sanzo_demo.comparison import comparison_groups, comparison_projection
from sanzo_demo.export import workbook_bytes


def quote(identity, company, *, name="Synthetic transport", fee_id=None,
          equipment="20FT", billing_unit="CONTAINER", transport_mode="SEA",
          ready=True, price="100"):
    return {
        "item_id": identity,
        "current_fee_id": fee_id,
        "current_fee_name": name if fee_id else None,
        "ai_suggested_name": None if fee_id else name,
        "current_price": price, "currency": "JPY", "version": 1,
        "input_kind": "MOCK_SDK",
        "saved_status": "HUMAN_CONFIRMED" if fee_id else ("AI_SUGGESTED" if ready else "REVIEW_REQUIRED"),
        "current_memory_status": "ACTIVE" if fee_id else None,
        "initial_classification": {},
        "source_ref": {"kind": "EXCEL_RANGE", "locator_status": "VERIFIED",
                       "origin": "MOCK_SDK", "locator": {"sheet_name": "Test", "A1_range": "A2:D2"}},
        "observation": {
            "document_id": "doc-" + company if company else "doc-unknown",
            "forwarder_id": company, "raw_label": "Synthetic original " + identity,
            "source_ready": True,
            "context": {"transport_mode": transport_mode, "equipment": equipment,
                        "billing_unit": billing_unit, "charge_scope": "ORIGIN",
                        "extra_terms": "STANDARD", "ready": ready},
        },
    }


def partition(items):
    return {frozenset(i["item_id"] for i in group["items"]) for group in comparison_groups(items)}


class ComparisonProjectionTests(unittest.TestCase):
    def test_legacy_confirmed_name_does_not_repeat_equipment_and_source_order_is_stable(self):
        later = quote("later", "A", name="Customer fee (40FT)", fee_id="f", equipment="40FT")
        earlier = quote("earlier", "A", name="Earlier fee")
        later["source_ref"]["locator"]["A1_range"] = "B13"
        earlier["source_ref"]["locator"]["A1_range"] = "B12"
        rows = comparison_projection([later, earlier])["rows"]
        self.assertEqual(rows[0]["item_ids"], ["earlier"])
        self.assertEqual(rows[1]["display_name"], "Customer fee (40FT)")
        self.assertEqual(rows[1]["distinguishing_conditions"], {"equipment": "40FT"})
        self.assertEqual(later["current_fee_name"], "Customer fee (40FT)")

    def test_missing_transport_allows_unique_provisional_pair_without_promotion(self):
        left = quote("a", "A", fee_id="fee-one")
        right = quote("b", "B", transport_mode=None, ready=False)
        items, before = [left, right], deepcopy([left, right])
        groups = comparison_groups(items)
        self.assertEqual(len(groups), 1)
        self.assertTrue(groups[0]["has_confirmed"])
        self.assertTrue(groups[0]["has_suggestions"])
        self.assertFalse(groups[0]["ambiguous"])
        projection = comparison_projection(items)
        self.assertEqual(projection["rows"][0]["cells"]["B"][0]["status"], "● AI仮判定")
        self.assertTrue(projection["rows"][0]["cells"]["B"][0]["needs_review"])
        self.assertEqual(projection["rows"][0]["cells"]["B"][0]["price"], right["current_price"])
        self.assertEqual(items, before)
        self.assertIsNone(right["current_fee_id"])
        self.assertIsNone(right["current_memory_status"])

    def test_two_ai_suggestions_with_unique_known_conditions_can_align(self):
        rows = [quote("a", "A"), quote("b", "B", transport_mode=None, ready=False)]
        self.assertEqual(partition(rows), {frozenset(("a", "b"))})
        self.assertTrue(all(item["current_fee_id"] is None for item in rows))
        self.assertEqual(comparison_projection(rows)["rows"][0]["cells"]["A"][0]["status"], "● AI仮判定")

    def test_known_equipment_and_billing_conflicts_never_align(self):
        for change in ({"equipment": "40FT"}, {"billing_unit": "SHIPMENT"},
                       {"transport_mode": "AIR"}):
            with self.subTest(change=change):
                self.assertEqual(partition([quote("a", "A"), quote("b", "B", **change)]),
                                 {frozenset(("a",)), frozenset(("b",))})

    def test_known_scope_and_extra_terms_conflicts_never_align(self):
        for field in ("charge_scope", "extra_terms"):
            with self.subTest(field=field):
                a, b = quote("a", "A"), quote("b", "B")
                b["observation"]["context"][field] = "CONFLICTING"
                self.assertEqual(len(comparison_groups([a, b])), 2)

    def test_unknown_markers_allow_only_unique_provisional_candidate(self):
        for value in (None, "", "  ", "未確認", "unknown"):
            with self.subTest(value=value):
                a, b = quote("a", "A"), quote("b", "B", transport_mode=value, ready=False)
                self.assertEqual(partition([a, b]), {frozenset(("a", "b"))})
                self.assertEqual(b["observation"]["context"]["transport_mode"], value)

    def test_unknown_equipment_cannot_choose_between_twenty_and_forty(self):
        rows = [quote("twenty", "A"), quote("forty", "A", equipment="40FT"),
                quote("unknown", "B", equipment=None, ready=False)]
        expected = {frozenset((r["item_id"],)) for r in rows}
        for order in permutations(rows):
            with self.subTest(order=[r["item_id"] for r in order]):
                self.assertEqual(partition(order), expected)
                projection = comparison_projection(order)
                candidate = next(r for r in projection["rows"] if "unknown" in r["item_ids"])
                self.assertTrue(candidate["ambiguous"])
                self.assertEqual(candidate["label"], "費目未確認")
                self.assertIn("確認必要", candidate["cells"]["B"][0]["status"])

    def test_different_confirmed_ids_never_merge_and_ai_stays_ambiguous(self):
        rows = [quote("one", "A", fee_id="fee-one"), quote("two", "B", fee_id="fee-two"),
                quote("candidate", "C", transport_mode=None, ready=False)]
        for order in permutations(rows):
            with self.subTest(order=[r["item_id"] for r in order]):
                self.assertEqual(len(comparison_groups(order)), 3)
                candidate = next(r for r in comparison_projection(order)["rows"]
                                 if r["item_ids"] == ["candidate"])
                self.assertTrue(candidate["ambiguous"])
                self.assertEqual(candidate["label"], "費目未確認")

    def test_known_condition_can_disambiguate_distinct_confirmed_ids(self):
        rows = [quote("one", "A", fee_id="fee-one"),
                quote("two", "A", fee_id="fee-two", equipment="40FT"),
                quote("candidate", "B", transport_mode=None, ready=False)]
        self.assertEqual(partition(rows), {frozenset(("one", "candidate")), frozenset(("two",))})

    def test_confirmed_identity_priority_precedes_tentative_candidates(self):
        rows = [quote("confirmed-20", "A", fee_id="fee-one"),
                quote("confirmed-unknown", "B", fee_id="fee-one", equipment=None, ready=False),
                quote("suggested-40", "C", equipment="40FT")]
        expected = {frozenset(("confirmed-20", "confirmed-unknown")), frozenset(("suggested-40",))}
        for order in permutations(rows):
            with self.subTest(order=[r["item_id"] for r in order]):
                self.assertEqual(partition(order), expected)

    def test_same_confirmed_id_different_names_still_aligns_without_known_conflict(self):
        rows = [quote("a", "A", name="Old display", fee_id="fee-one"),
                quote("b", "B", name="Current display", fee_id="fee-one", transport_mode=None, ready=False)]
        self.assertEqual(partition(rows), {frozenset(("a", "b"))})

    def test_same_confirmed_id_with_twenty_forty_conflict_stays_separate(self):
        rows = [quote("a", "A", fee_id="fee-one"),
                quote("b", "B", fee_id="fee-one", equipment="40FT")]
        self.assertEqual(len(comparison_groups(rows)), 2)

    def test_current_projection_merges_adopted_versions_and_preserves_three_quotes(self):
        fee_id = "ac4edd3c-cea6-4319-ada3-219d952c2549"
        items = []
        for equipment, prices in (("20FT", ("1200", "1150", "1350")),
                                  ("40FT", ("1900", "1850", "2050"))):
            for index, (company, price) in enumerate(zip(("A初回", "B初回", "A次回"), prices)):
                old = equipment == "40FT" and index < 2
                item = quote(equipment + "-" + company, company,
                    name="海上運賃" if old else "海上運賃_v1", fee_id=fee_id,
                    equipment=equipment, price=price)
                item.update(current_fee_name_version=1 if old else 2,
                            current_fee_name_snapshot=item["current_fee_name"],
                            effective_fee_name="海上運賃_v1", effective_fee_name_version=2,
                            currency="USD")
                items.append(item)
        fees = [{"fee_id": fee_id, "display_name": "海上運賃_v1", "name_version": 2}]
        before = deepcopy((items, fees))
        # Input order cannot pick an old adopted name as the current row label.
        for ordered in (items, list(reversed(items)), items[3:] + items[:3]):
            with self.subTest(order=[item["item_id"] for item in ordered]):
                projection = comparison_projection(ordered, fees)
                self.assertEqual(len(projection["rows"]), 2)
                by_equipment = {row["equipment"]: row for row in projection["rows"]}
                for equipment, prices in (("20FT", ("1200", "1150", "1350")),
                                          ("40FT", ("1900", "1850", "2050"))):
                    row = by_equipment[equipment]
                    self.assertEqual(row["display_name"], "海上運賃_v1（" + equipment + "）")
                    self.assertEqual(row["fee_identity"], fee_id)
                    self.assertEqual(row["rename_name_version"], 2)
                    self.assertEqual(len(row["item_ids"]), 3)
                    for company, price in zip(("A初回", "B初回", "A次回"), prices):
                        self.assertEqual([(c["price"], c["currency"]) for c in row["cells"][company]],
                                         [(price, "USD")])
        documents = [{"document_id": "doc-" + company, "original_filename": company + ".xlsx",
                      "sha256": "a" * 64, "state": "AVAILABLE"}
                     for company in ("A初回", "B初回", "A次回")]
        projection = comparison_projection(items, fees, documents)
        book = load_workbook(BytesIO(workbook_bytes({"name": "Synthetic rename", "case_revision": 1},
                             items, documents, fees=fees)), data_only=False)
        try:
            sheet = book["原価単価一覧表"]
            self.assertEqual(sheet.max_row - 3, 2)
            self.assertEqual([sheet.cell(3, n).value for n in range(2, 5)], projection["companies"])
            for number, row in enumerate(projection["rows"], 4):
                self.assertEqual(sheet.cell(number, 1).value, row["display_name"])
                for column, company in enumerate(projection["companies"], 2):
                    self.assertEqual(sheet.cell(number, column).value,
                        "\n".join(c["price"] + " " + c["currency"] for c in row["cells"][company]))
            detail_names = {row[0]: row[3] for row in list(book["明細と出典"].values)[1:]}
            self.assertEqual(detail_names, {item["item_id"]: item["current_fee_name"] for item in items})
        finally:
            book.close()
        self.assertEqual((items, fees), before, "Read-only projection/export must preserve adopted history and sources")

    def test_same_identity_with_mixed_versions_respects_every_known_condition_conflict(self):
        for field, conflict in (("transport_mode", "AIR"), ("equipment", "40FT"),
                                ("billing_unit", "SHIPMENT"), ("charge_scope", "DESTINATION"),
                                ("extra_terms", "SPECIAL")):
            with self.subTest(field=field):
                a = quote("a", "A", fee_id="stable", name="Old adopted")
                b = quote("b", "B", fee_id="stable", name="Latest")
                a.update(current_fee_name_version=1, effective_fee_name="Latest", effective_fee_name_version=2)
                b.update(current_fee_name_version=2, effective_fee_name="Latest", effective_fee_name_version=2)
                b["observation"]["context"][field] = conflict
                self.assertEqual(partition([a, b]), {frozenset(("a",)), frozenset(("b",))})

    def test_same_current_label_never_merges_distinct_fee_identities(self):
        items = [quote("a", "A", name="Shared label", fee_id="identity-a"),
                 quote("b", "B", name="Shared label", fee_id="identity-b")]
        for item in items:
            item.update(current_fee_name_version=1, effective_fee_name="Shared label", effective_fee_name_version=2)
        fees = [{"fee_id": identity, "display_name": "Shared label", "name_version": 2}
                for identity in ("identity-a", "identity-b")]
        rows = comparison_projection(items, fees)["rows"]
        self.assertEqual(len(rows), 2)
        self.assertEqual({row["fee_identity"] for row in rows}, {"identity-a", "identity-b"})

    def test_current_name_comes_from_catalog_or_effective_fields_not_suffix_order(self):
        item = quote("old", "A", name="海上運賃_v999", fee_id="stable")
        item.update(current_fee_name_version=1, effective_fee_name="海上運賃", effective_fee_name_version=2)
        before = deepcopy(item)
        for fees in ([], [{"fee_id": "stable", "display_name": "海上運賃", "name_version": 2}]):
            with self.subTest(catalog=bool(fees)):
                row = comparison_projection([item], fees)["rows"][0]
                self.assertEqual(row["label"], "海上運賃")
                self.assertEqual(row["rename_name_version"], 2)
        # A catalog supplied for this read is authoritative even for a legacy
        # item dictionary without joined effective-name fields.
        legacy = {k: v for k, v in item.items() if not k.startswith("effective_")}
        row = comparison_projection([legacy], [{"fee_id": "stable", "display_name": "最新統一名", "name_version": 3}])["rows"][0]
        self.assertEqual(row["label"], "最新統一名")
        self.assertEqual(row["rename_name_version"], 3)
        self.assertEqual(item, before)

    def test_catalog_name_and_version_override_stale_effective_metadata_together(self):
        item = quote("a", "A", name="Adopted", fee_id="stable")
        item.update(current_fee_name_version=1, effective_fee_name="Earlier", effective_fee_name_version=2)
        row = comparison_projection([item], [{"fee_id": "stable", "display_name": "Current", "name_version": 3}])["rows"][0]
        self.assertEqual((row["label"], row["rename_name_version"]), ("Current", 3))
        # A partial legacy catalog must not pair its name with a different
        # effective name's concurrency version.
        row = comparison_projection([item], [{"fee_id": "stable", "display_name": "Current"}])["rows"][0]
        self.assertEqual((row["label"], row["rename_name_version"]), ("Current", None))

    def test_effective_metadata_fallback_uses_explicit_version_independent_of_order(self):
        old = quote("a", "A", name="Adopted", fee_id="stable")
        new = quote("b", "B", name="Adopted", fee_id="stable")
        old.update(effective_fee_name="Earlier_v999", effective_fee_name_version=2)
        new.update(effective_fee_name="Current", effective_fee_name_version=3)
        before = deepcopy([old, new])
        for items in permutations((old, new)):
            row = comparison_projection(items)["rows"][0]
            self.assertEqual((row["label"], row["rename_name_version"]), ("Current", 3))
            self.assertEqual(set(row["item_ids"]), {"a", "b"})
        self.assertEqual([old, new], before)

    def test_unknown_bridge_cannot_transitively_merge_incompatible_contexts(self):
        a, b, c = quote("a", "A"), quote("b", "B", equipment=None, ready=False), quote("c", "C", equipment="40FT")
        for order in permutations((a, b, c)):
            with self.subTest(order=[r["item_id"] for r in order]):
                self.assertEqual(len(comparison_groups(order)), 3)

    def test_exact_duplicate_quotes_are_preserved_in_same_company_cell(self):
        rows = [quote("a", "A", price="100"), quote("b", "B", price="90"),
                quote("b-second", "B", price="95")]
        projection = comparison_projection(rows)
        self.assertEqual(len(projection["rows"]), 1)
        self.assertEqual([c["item_id"] for c in projection["rows"][0]["cells"]["B"]], ["b", "b-second"])
        self.assertEqual(partition(rows), {frozenset(("a", "b", "b-second"))})

    def test_partial_same_company_candidates_do_not_supply_cross_company_match(self):
        rows = [quote("one", "A"), quote("two", "A", transport_mode=None, ready=False)]
        self.assertEqual(len(comparison_groups(rows)), 2)

    def test_missing_company_or_missing_suggestion_is_not_guessed(self):
        for b in (quote("b", None), quote("b", "B", name=None)):
            with self.subTest(item=b["item_id"], company=b["observation"]["forwarder_id"], name=b["ai_suggested_name"]):
                self.assertEqual(len(comparison_groups([quote("a", "A"), b])), 2)

    def test_classification_names_require_exact_agreement(self):
        self.assertEqual(len(comparison_groups([quote("a", "A"), quote("b", "B", name="Synthetic OTHER")])), 2)

    def test_catalog_display_name_can_identify_unique_confirmed_candidate(self):
        a, b = quote("a", "A", name=None, fee_id="fee-one"), quote("b", "B")
        groups = comparison_groups([a, b], [{"fee_id": "fee-one", "display_name": "Synthetic transport"}])
        self.assertEqual(len(groups), 1)
        self.assertIsNone(b["current_fee_id"])

    def test_incomplete_or_unverified_item_retains_review_even_if_legacy_ready_true(self):
        missing, unverified = quote("missing", "A", transport_mode=None), quote("unverified", "B")
        unverified["observation"]["source_ready"] = False
        for item in (missing, unverified):
            with self.subTest(identity=item["item_id"]):
                row = comparison_projection([item])["rows"][0]
                self.assertTrue(row["cells"][item["observation"]["forwarder_id"]][0]["needs_review"])

    def test_projection_covers_every_item_once_and_does_not_mutate_inputs(self):
        items = [quote("a", "A", fee_id="fee-one"), quote("b", "B", transport_mode=None, ready=False),
                 quote("c", "B", name="Synthetic other", equipment="40FT")]
        fees = [{"fee_id": "fee-one", "display_name": "Synthetic transport"}]
        before = deepcopy((items, fees))
        result = comparison_projection(items, fees)
        self.assertEqual(result["version"], 1)
        self.assertEqual(result["companies"], ["A", "B"])
        self.assertCountEqual([identity for row in result["rows"] for identity in row["item_ids"]], ["a", "b", "c"])
        self.assertCountEqual([cell["item_id"] for row in result["rows"] for cells in row["cells"].values() for cell in cells], ["a", "b", "c"])
        self.assertEqual((items, fees), before)

    def test_shared_export_has_same_five_pairs_and_preserves_review_and_prices(self):
        # Purely synthetic names and prices test mechanics; not a fixture/AI substitute.
        rows = []
        for n, (name, equipment) in enumerate((("Transport", "20FT"), ("Transport", "40FT"),
                                               ("Docs", "COMMON"), ("Clearance", "COMMON"), ("Local", "COMMON"))):
            rows += [quote("a" + str(n), "A", name=name, equipment=equipment,
                           price=str(100 + n), fee_id="fee-local" if n == 4 else None),
                     quote("b" + str(n), "B", name=name, equipment=equipment,
                           price=str(90 + n), transport_mode=None, ready=False)]
        before = deepcopy(rows)
        projection = comparison_projection(rows)
        self.assertEqual(len(projection["rows"]), 5)
        self.assertTrue(all(len(row["cells"]["A"]) == len(row["cells"]["B"]) == 1 for row in projection["rows"]))
        document_rows = [{"document_id": "doc-" + company, "original_filename": "synthetic-" + company + ".xlsx",
                          "sha256": "a" * 64, "state": "AVAILABLE"} for company in ("A", "B")]
        book = load_workbook(BytesIO(workbook_bytes(
            {"name": "SYNTHETIC PROJECTION TEST", "case_revision": 1},
            rows, document_rows, reader_mode="MOCK_SDK")), data_only=False)
        try:
            sheet = book["原価単価一覧表"]
            self.assertEqual(sheet.max_row - 3, len(projection["rows"]))
            self.assertEqual([sheet.cell(3, i).value for i in range(2, 4)], projection["companies"])
            by_id = {i["item_id"]: i for i in rows}
            for number, row in enumerate(projection["rows"], 4):
                self.assertIn(row["label"], sheet.cell(number, 1).value)
                for col, company in enumerate(projection["companies"], 2):
                    text = sheet.cell(number, col).value
                    for cell in row["cells"][company]:
                        item = by_id[cell["item_id"]]
                        self.assertIn(item["current_price"] + " " + item["currency"], text)
                        self.assertNotIn(cell["status"], text)
                    if company == "B":
                        self.assertNotIn("AI仮判定", text)
                        self.assertTrue(row["cells"]["B"][0]["needs_review"])
            self.assertIn("DRAFT", sheet["A2"].value)
        finally:
            book.close()
        self.assertEqual(rows, before)
        self.assertTrue(all(i["current_fee_id"] is None for i in rows if i["observation"]["forwarder_id"] == "B"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
