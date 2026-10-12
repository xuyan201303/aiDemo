"""Synthetic current-quotation confirmation contracts, never SDK/AI results.

Writes are restricted to SANZO_TEST_DSN. Original files are tiny synthetic
fixtures under the caller's TMPDIR; no formal upload or database is accessed.
"""
from __future__ import annotations

from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Barrier
import unittest
from uuid import uuid4

from openpyxl import load_workbook
import psycopg
from psycopg import sql
from psycopg.types.json import Jsonb

from sanzo_core.condition_corrections import ConditionCorrectionCommand, ConditionCorrectionService
from sanzo_core.demo_store import DemoStore
from sanzo_core.domain import ConflictError, Context, CorrectionCommand, Observation, Principal, RevokeCommand, ScopeError, SyntheticLine
from sanzo_core.postgres import PostgresRepository
from sanzo_core.ingestion_store import IngestionStore
from sanzo_core.quote_confirmations import (CONDITION_FIELDS, DocumentReviewCommand,
    DocumentReviewService, QuoteConfirmationCommand, QuoteConfirmationService, project_confirmations)
from sanzo_core.service import ReviewService
from sanzo_demo.case_overview import summarize_case
from sanzo_demo.comparison import comparison_projection
from sanzo_demo.current_document_review import current_document_reviews
from sanzo_demo.export import workbook_bytes
from sanzo_demo.review_state import review_facts
from sanzo_demo.row_confirmation import RowConfirmationCommand, RowConfirmationService, RowMember

DSN = os.environ.get("SANZO_TEST_DSN")


class QuoteConfirmationCommandTests(unittest.TestCase):
    def command(self, **changes):
        values = dict(case_id="synthetic-case", row_id="a" * 64,
            item_id="synthetic-item", expected_case_revision=1,
            expected_item_version=1, idempotency_key="synthetic-save",
            kind="CONDITION_VALUE", field="equipment", value="20FT",
            evidence_note="Synthetic original B14 explicitly says 20FT")
        return QuoteConfirmationCommand(**{**values, **changes})

    def test_confirmation_command_is_bounded_and_cannot_claim_arbitrary_fields(self):
        self.command().validate()
        self.command(kind="SOURCE_REVIEW", field=None, value=None).validate()
        for changes in ({"evidence_note": " "}, {"evidence_note": "x" * 2001},
                        {"kind": "ALL_CONFIRMED"}, {"field": "forwarder_id"},
                        {"field": "transport_mode"}, {"value": None},
                        {"expected_case_revision": True}, {"expected_item_version": 0},
                        {"row_id": "bad"}, {"kind": "SOURCE_REVIEW"},
                        {"value": "x" * 1025}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                self.command(**changes).validate()


class QuoteConfirmationProjectionTests(unittest.TestCase):
    def synthetic_saved_reader_item(self):
        """Injected saved-result structure; the REAL marker tests its policy only."""
        record = ["EXCEL_RANGE", "SYNTHETIC ONLY", "injected-record"]
        values = {"raw_label": "SYNTHETIC OCEAN", "raw_value": "1200", "currency": "USD",
                  "equipment": "20FT", "billing_unit": "CONTAINER",
                  "charge_scope": "ORIGIN", "extra_terms": "STANDARD"}
        fields = {name: {"value": value, "original_value": value,
            "source_block_id": "synthetic-" + name, "record_key": record,
            "evidence_unit": {"kind": "EXCEL_CELL", "field_name": name,
                              "record_id": record[-1], "scope": "LINE"}}
            for name, value in values.items()}
        return {"item_id": "synthetic-item", "current_fee_id": "synthetic-fee",
            "saved_status": "HUMAN_CONFIRMED", "current_memory_status": None,
            "current_price": "1200", "raw_value": "1200", "currency": "USD",
            "input_kind": "SDK_OBSERVED", "effective_context": {
                key: values[key] for key in CONDITION_FIELDS},
            "observation": {"tenant_id": "synthetic-tenant", "item_id": "synthetic-item",
                "case_id": "synthetic-case", "document_id": "synthetic-document",
                "source_id": "synthetic-raw_label", "forwarder_id": None,
                "raw_label": values["raw_label"], "unit_price": "1200", "currency": "USD",
                "source_ready": False, "context": {"ready": False, "transport_mode": None,
                    **{key: values[key] for key in CONDITION_FIELDS}}},
            "source_ref": {"synthetic_test_injected": True, "origin": "SDK_OBSERVED",
                "extraction_origin": "REAL_SDK", "document_id": "synthetic-document",
                "document_sha256": "a" * 64, "source_block_ids": ["synthetic-" + key for key in values],
                "kind": "EXCEL_RANGE", "locator_status": "VERIFIED", "text_verification": "VERIFIED",
                "locator": {"sheet_name": "SYNTHETIC ONLY", "A1_range": "B14:H14"},
                "evidence_unit": fields["raw_label"]["evidence_unit"], "field_sources": fields}}

    def test_document_company_and_transport_metadata_do_not_create_new_required_fields(self):
        item = self.synthetic_saved_reader_item()
        facts = review_facts(item)
        self.assertFalse(facts["review_required"])
        self.assertEqual(len(facts["metadata_notices"]), 2)
        self.assertFalse(item["observation"]["source_ready"])
        self.assertFalse(item["observation"]["context"]["ready"])

    def test_complete_field_values_do_not_fabricate_precise_page_location(self):
        item = self.synthetic_saved_reader_item()
        item["source_ref"].update(kind="PDF_PAGE", locator_status="PAGE_ONLY", locator={"page_number": 1})
        facts = review_facts(item)
        self.assertTrue(facts["review_required"])
        self.assertIn("ページ内の明細範囲は未確定", facts["review_reasons"])

    def test_mismatched_amount_evidence_remains_pending_despite_grounded_label(self):
        item = self.synthetic_saved_reader_item()
        item["source_ref"]["field_sources"]["raw_value"]["original_value"] = "9000"
        facts = review_facts(item)
        self.assertTrue(facts["review_required"])
        self.assertIn("原文未確認", facts["review_reasons"])


@unittest.skipUnless(DSN, "NOT_RUN: dedicated isolated SANZO_TEST_DSN required")
class QuoteConfirmationPostgresTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        PostgresRepository(DSN).migrate()

    def setUp(self):
        self.tenant = "synthetic-quote-confirmation-" + uuid4().hex
        self.actor = Principal(self.tenant, "synthetic-human-reviewer")
        self.repo = PostgresRepository(DSN)
        self.store = DemoStore(self.repo)
        self.review = ReviewService(self.repo)
        self.temp = TemporaryDirectory(prefix="synthetic-quote-confirmation-")
        self.addCleanup(self.temp.cleanup)
        self.storage = Path(self.temp.name)
        self.confirmations = QuoteConfirmationService(self.repo, self.storage)
        self.conditions = ConditionCorrectionService(self.repo)

    def case(self):
        case_id = "synthetic-case-" + uuid4().hex
        self.review.create_case(self.actor, case_id, "SYNTHETIC ONLY / quotation confirmation")
        return case_id

    def line(self, case_id, *, label="SYNTHETIC OCEAN FREIGHT", company="A",
             equipment="20FT", price="1200", ready=True, source_ready=True,
             locator_status="VERIFIED", document_id=None):
        if document_id is None:
            document_id = "synthetic-document-" + uuid4().hex
            original = b"SYNTHETIC ONLY: OCEAN FREIGHT 20FT / 1200 USD\n"
            key = document_id + ".xlsx"
            (self.storage / key).write_bytes(original)
            self.store.register_document(self.actor, case_id, document_id=document_id,
                original_filename="SYNTHETIC_ONLY.xlsx",
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                size_bytes=len(original), sha256=sha256(original).hexdigest(), storage_key=key)
        document = self.store.get_document(self.actor, document_id)
        item_id, source_id = "synthetic-item-" + uuid4().hex, "synthetic-source-" + uuid4().hex
        source = {"synthetic": True, "document_id": document_id,
            "document_revision": 1, "document_sha256": document["sha256"],
            "source_block_ids": [source_id], "kind": "EXCEL_RANGE",
            "locator_status": locator_status,
            "locator": {"sheet_name": "SYNTHETIC ONLY", "A1_range": "B14"}}
        if locator_status == "PAGE_ONLY":
            source.update(kind="PDF_PAGE", locator={"page_number": 1})
        observation = Observation(self.tenant, item_id, case_id, document_id, source_id,
            company, label, price, "USD",
            Context("SEA", equipment, "CONTAINER", "ORIGIN", "STANDARD", ready=ready), source_ready)
        return self.store.add_synthetic_lines(self.actor, document_id,
            [SyntheticLine(observation, source, raw_value=price)])[0]

    def current(self, item):
        return next(member for member in self.store.list_items(self.actor, item["case_id"])
                    if member["item_id"] == item["item_id"])

    def classify(self, item, *, name="海上運賃", fee_id=None):
        current = self.current(item)
        return self.review.correct(self.actor, CorrectionCommand(
            item_id=current["item_id"], expected_case_revision=current["case_revision"],
            expected_item_version=current["version"],
            expected_memory_version=current["current_memory_version"],
            idempotency_key=uuid4().hex, new_fee_name=name if fee_id is None else None,
            target_fee_id=fee_id))

    def command(self, item, *, kind="CONDITION_VALUE", field="equipment", value=None, **changes):
        snapshot = self.store.case_snapshot(self.actor, item["case_id"])
        row = next(row for row in comparison_projection(snapshot["items"], snapshot["fees"],
            snapshot["documents"])["rows"] if item["item_id"] in row["item_ids"])
        member = next(member for member in row["member_items"] if member["item_id"] == item["item_id"])
        if kind == "SOURCE_REVIEW":
            field, value = None, None
        elif value is None:
            value = member["effective_context"].get(field)
        return QuoteConfirmationCommand(**{**dict(case_id=item["case_id"], row_id=row["row_id"],
            item_id=item["item_id"], expected_case_revision=snapshot["case"]["case_revision"],
            expected_item_version=member["version"], idempotency_key=uuid4().hex,
            kind=kind, field=field, value=value,
            evidence_note="SYNTHETIC original: matched label, condition and amount on page 1"), **changes})

    def database_snapshot(self):
        with self.repo.transaction() as conn:
            tables = conn.execute("SELECT table_name FROM information_schema.columns "
                "WHERE table_schema='sanzo_core' AND column_name='tenant_id' ORDER BY table_name").fetchall()
            return {row["table_name"]: conn.execute(sql.SQL(
                "SELECT row_to_json(t)::text AS row FROM {}.{} t "
                "WHERE tenant_id=%s ORDER BY row_to_json(t)::text").format(
                    sql.Identifier("sanzo_core"), sql.Identifier(row["table_name"])),
                (self.tenant,)).fetchall() for row in tables}

    def events(self, item):
        with self.repo.transaction() as conn:
            return conn.execute("SELECT * FROM sanzo_core.quote_confirmation_events "
                "WHERE tenant_id=%s AND item_id=%s ORDER BY item_version,event_id",
                (self.tenant, item["item_id"])).fetchall()

    def inject_historical_task(self, item, issues):
        """Explicit synthetic test injection of stored reader state, with no SDK."""
        document_id = item["observation"]["document_id"]
        document = self.store.get_document(self.actor, document_id)
        with self.repo.transaction() as conn:
            conn.execute("""INSERT INTO sanzo_core.ingestion_tasks
                (tenant_id,task_id,case_id,document_id,document_sha256,actor_id,state,result_json)
                VALUES(%s,%s,%s,%s,%s,%s,'REVIEW_REQUIRED',%s)""",
                (self.tenant, "synthetic-task-" + uuid4().hex, item["case_id"], document_id,
                 document["sha256"], self.actor.actor_id,
                 Jsonb({"synthetic_test_injected": True, "issues": issues, "item_ids": [item["item_id"]]})))

    def document_command(self, item, **changes):
        snapshot = IngestionStore(self.repo).case_snapshot(self.actor, item["case_id"])
        document_id = item["observation"]["document_id"]
        review = current_document_reviews(snapshot["items"], snapshot["documents"])[document_id]
        return DocumentReviewCommand(**{**dict(case_id=item["case_id"], document_id=document_id,
            expected_case_revision=snapshot["case"]["case_revision"],
            expected_review_sha256=review["review_sha256"], confirmed_item_count=review["item_count"],
            evidence_note="SYNTHETIC original page 1: every visible quotation line counted and matched",
            idempotency_key=uuid4().hex), **changes})

    def explicitly_check_all(self, item):
        for field in CONDITION_FIELDS:
            self.confirmations.confirm(self.actor, self.command(item, field=field))
        self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))

    def processed_snapshot(self, case_id, issues=None):
        """Explicit injected historical processing state, never an SDK outcome."""
        snapshot = deepcopy(self.store.case_snapshot(self.actor, case_id))
        snapshot["documents"] = [{**doc, "state": "REVIEW_REQUIRED", "blocked_reason": None,
            "issues": issues if issues is not None else [
                {"code": "APPLICABILITY_INCOMPLETE", "item_id": item["item_id"]}
                for item in snapshot["items"] if item["observation"]["document_id"] == doc["document_id"]]}
            for doc in snapshot["documents"]]
        return snapshot

    def assert_surfaces(self, snapshot, *, complete):
        self.assertEqual(summarize_case(snapshot)["status"], "確認済み" if complete else "確認必要")
        rows = comparison_projection(snapshot["items"], snapshot["fees"], snapshot["documents"])["rows"]
        if complete:
            self.assertTrue(all(row["confirmation_status"] == "✓ 確認済み" for row in rows))
            self.assertTrue(all(not review_facts(item)["review_required"] for item in snapshot["items"]))
        book = load_workbook(BytesIO(workbook_bytes(snapshot["case"], snapshot["items"],
            snapshot["documents"], fees=snapshot["fees"])), read_only=True)
        try:
            self.assertEqual("DRAFT" not in book["原価単価一覧表"]["A2"].value, complete)
            self.assertEqual(len(book.sheetnames), 4)
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        finally:
            book.close()

    def test_same_value_confirmation_survives_refresh_without_false_change_or_learning_edits(self):
        item = self.line(self.case())
        self.classify(item)
        before, database_before = deepcopy(self.current(item)), self.database_snapshot()
        result = self.confirmations.confirm(self.actor, self.command(item))
        refreshed = self.current(item)
        self.assertEqual(result["kind"], "ITEM_QUOTE_CONFIRMATION")
        self.assertEqual(result["review_facts"], review_facts(refreshed))
        self.assertTrue(refreshed["quote_confirmations"]["conditions"]["equipment"]["valid"])
        for field in ("observation", "source_ref", "current_price", "currency",
                      "effective_context", "condition_overrides", "current_memory_version",
                      "current_memory_status", "future_reuse_eligible"):
            self.assertEqual(refreshed[field], before[field], field)
        after = self.database_snapshot()
        for table in ("condition_correction_events", "item_condition_overlays", "memory_versions",
                      "memory_events", "memory_slots", "correction_events", "fee_name_versions"):
            self.assertEqual(after[table], database_before[table], table)
        event = self.events(item)[0]
        self.assertEqual(event["actor_id"], self.actor.actor_id)
        self.assertEqual(event["value_snapshot"]["value"], "20FT")
        self.assertTrue(event["created_at"])

    def test_each_needed_condition_is_confirmed_individually_without_changing_reader_ready(self):
        item = self.line(self.case(), ready=False)
        self.classify(item)
        self.assertIn("適用条件未確認", review_facts(self.current(item))["review_reasons"])
        for field in CONDITION_FIELDS:
            self.confirmations.confirm(self.actor, self.command(item, field=field))
            facts = review_facts(self.current(item))
            self.assertEqual("適用条件未確認" in facts["review_reasons"], field != CONDITION_FIELDS[-1])
        self.assertFalse(self.current(item)["observation"]["context"]["ready"])
        self.assertEqual(self.current(item)["condition_overrides"], {})
        self.assert_surfaces(self.processed_snapshot(item["case_id"]), complete=True)

    def test_page_only_source_review_preserves_locator_and_records_human_evidence(self):
        item = self.line(self.case(), source_ready=False, locator_status="PAGE_ONLY")
        self.classify(item)
        before = deepcopy(self.current(item))
        result = self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))
        refreshed = self.current(item)
        self.assertFalse(result["review_required"])
        self.assertTrue(refreshed["quote_confirmations"]["source"]["valid"])
        self.assertEqual(refreshed["observation"], before["observation"])
        self.assertEqual(refreshed["source_ref"], before["source_ref"])
        self.assertEqual(refreshed["source_ref"]["locator_status"], "PAGE_ONLY")
        event = self.events(item)[0]
        self.assertEqual(event["document_sha256"], before["source_ref"]["document_sha256"])
        self.assertIn("page 1", event["evidence_note"])
        self.assertEqual(len(event["anchor_sha256"]), 64)
        self.assert_surfaces(self.processed_snapshot(item["case_id"]), complete=True)
        old_events = deepcopy(self.events(item))
        for mutation in (lambda value: value.update(current_price="1300"),
                         lambda value: value.update(currency="JPY"),
                         lambda value: value["source_ref"]["locator"].update(page_number=2)):
            changed = deepcopy(refreshed)
            mutation(changed)
            projection = project_confirmations(changed, old_events, event["document_sha256"])
            self.assertFalse(projection["source"]["valid"])
        self.assertFalse(project_confirmations(refreshed, old_events, "f" * 64)["source"]["valid"])
        self.assertEqual(self.events(item), old_events)

    def test_source_review_rejects_changed_or_missing_original_without_writes(self):
        item = self.line(self.case(), source_ready=False, locator_status="PAGE_ONLY")
        self.classify(item)
        document = self.store.get_document(self.actor, item["observation"]["document_id"])
        path = self.storage / document["storage_key"]
        original = path.read_bytes()
        for content in (b"X" * len(original), None):
            path.write_bytes(content) if content is not None else path.unlink()
            before = self.database_snapshot()
            with self.assertRaises(ValueError):
                self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))
            self.assertEqual(self.database_snapshot(), before)
            self.assertTrue(review_facts(self.current(item))["review_required"])

    def test_value_confirmation_does_not_perform_a_field_change(self):
        item = self.line(self.case())
        self.classify(item)
        before = self.database_snapshot()
        with self.assertRaises(ConflictError):
            self.confirmations.confirm(self.actor, self.command(item, value="40FT"))
        self.assertEqual(self.database_snapshot(), before)

    def test_confirmation_history_is_append_only_and_evidence_exports_as_literal_text(self):
        item = self.line(self.case(), source_ready=False, locator_status="PAGE_ONLY")
        self.classify(item)
        note = '=HYPERLINK("https://synthetic.invalid", "SYNTHETIC page 1 reviewed")'
        self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW", evidence_note=note))
        event = self.events(item)[0]
        for statement in ("UPDATE sanzo_core.quote_confirmation_events SET evidence_note='changed' "
                          "WHERE tenant_id=%s AND event_id=%s",
                          "DELETE FROM sanzo_core.quote_confirmation_events WHERE tenant_id=%s AND event_id=%s"):
            with self.assertRaises(psycopg.Error):
                with self.repo.transaction() as conn:
                    conn.execute(statement, (self.tenant, event["event_id"]))
        self.assertEqual(self.events(item)[0], event)
        snapshot = self.processed_snapshot(item["case_id"])
        book = load_workbook(BytesIO(workbook_bytes(snapshot["case"], snapshot["items"],
            snapshot["documents"], fees=snapshot["fees"])), read_only=True)
        try:
            cells = [cell for sheet in book for row in sheet for cell in row]
            self.assertTrue(any(note in str(cell.value) for cell in cells))
            self.assertFalse(any(cell.data_type == "f" for cell in cells))
        finally:
            book.close()

    def test_missing_condition_cannot_be_confirmed_away(self):
        item = self.line(self.case(), equipment=None, ready=False)
        self.classify(item)
        before = self.database_snapshot()
        with self.assertRaises(ValueError):
            self.confirmations.confirm(self.actor, self.command(item, value="20FT"))
        self.assertEqual(self.database_snapshot(), before)
        self.assertIn("適用条件未確認", review_facts(self.current(item))["review_reasons"])
        self.assert_surfaces(self.processed_snapshot(item["case_id"]), complete=False)

    def test_confirming_one_member_does_not_confirm_sibling_or_whole_case(self):
        case_id = self.case()
        first = self.line(case_id, source_ready=False, locator_status="PAGE_ONLY")
        second = self.line(case_id, label="SYNTHETIC O/F", company="B", price="1150",
                           source_ready=False, locator_status="PAGE_ONLY",
                           document_id=first["observation"]["document_id"])
        fee = self.classify(first)["fee_id"]
        self.classify(second, fee_id=fee)
        sibling_before = deepcopy(self.current(second))
        self.confirmations.confirm(self.actor, self.command(first, kind="SOURCE_REVIEW"))
        # case_revision is a shared read projection; the sibling's persisted
        # values and own version remain unchanged when one member is confirmed.
        self.assertEqual({key: value for key, value in self.current(second).items() if key != "case_revision"},
                         {key: value for key, value in sibling_before.items() if key != "case_revision"})
        self.assertFalse(review_facts(self.current(first))["review_required"])
        self.assertTrue(review_facts(self.current(second))["review_required"])
        partial = self.processed_snapshot(case_id)
        self.assert_surfaces(partial, complete=False)
        row = comparison_projection(partial["items"], partial["fees"], partial["documents"])["rows"][0]
        self.assertEqual(row["confirmation_status"], "⚠ 確認必要")
        self.confirmations.confirm(self.actor, self.command(second, kind="SOURCE_REVIEW"))
        self.assert_surfaces(self.processed_snapshot(case_id), complete=True)

    def test_idempotency_stale_item_case_and_scope_are_enforced(self):
        item = self.line(self.case())
        self.classify(item)
        command = self.command(item)
        result = self.confirmations.confirm(self.actor, command)
        before = self.database_snapshot()
        self.assertEqual(self.confirmations.confirm(self.actor, command), result)
        self.assertEqual(self.database_snapshot(), before)
        for invalid in (replace(command, evidence_note="different payload"),
                        replace(command, idempotency_key=uuid4().hex),
                        replace(self.command(item), expected_item_version=1),
                        replace(self.command(item), row_id="f" * 64)):
            with self.subTest(command=invalid), self.assertRaises(ConflictError):
                self.confirmations.confirm(self.actor, invalid)
            self.assertEqual(self.database_snapshot(), before)
        forbidden = Principal(self.tenant, self.actor.actor_id, allowed_case_ids=frozenset())
        with self.assertRaises(ScopeError):
            self.confirmations.confirm(forbidden, self.command(item))
        outsider = Principal("synthetic-other-tenant", self.actor.actor_id)
        with self.assertRaises(ScopeError):
            self.confirmations.confirm(outsider, self.command(item))
        self.assertEqual(self.database_snapshot(), before)

    def test_transaction_failure_rolls_back_event_versions_and_command(self):
        item = self.line(self.case())
        self.classify(item)
        for point in ("after_quote_confirmation", "before_quote_confirmation_command"):
            def fail(stage):
                if stage == point:
                    raise RuntimeError("SYNTHETIC_TRANSACTION_FAILURE")
            service = QuoteConfirmationService(self.repo, self.storage, fault_hook=fail)
            before = self.database_snapshot()
            with self.subTest(point=point), self.assertRaisesRegex(RuntimeError, "SYNTHETIC_TRANSACTION_FAILURE"):
                service.confirm(self.actor, self.command(item))
            self.assertEqual(self.database_snapshot(), before)

    def test_concurrent_confirmations_accept_exactly_one_current_version(self):
        item = self.line(self.case())
        self.classify(item)
        first = self.command(item)
        second = replace(first, idempotency_key=uuid4().hex)
        barrier = Barrier(2)
        def submit(command):
            barrier.wait(timeout=10)
            try:
                QuoteConfirmationService(PostgresRepository(DSN), self.storage).confirm(self.actor, command)
                return "SAVED"
            except ConflictError:
                return "STALE"
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(submit, (first, second)))
        self.assertCountEqual(results, ["SAVED", "STALE"])
        self.assertEqual(len(self.events(item)), 1)

    def test_changed_condition_invalidates_old_confirmation_without_erasing_history(self):
        item = self.line(self.case())
        self.classify(item)
        self.confirmations.confirm(self.actor, self.command(item))
        old_events = deepcopy(self.events(item))
        current = self.current(item)
        quoted = self.command(item)
        self.conditions.correct(self.actor, ConditionCorrectionCommand(
            case_id=item["case_id"], row_id=quoted.row_id, item_id=item["item_id"],
            expected_case_revision=quoted.expected_case_revision,
            expected_item_version=quoted.expected_item_version,
            expected_memory_version=current["current_memory_version"],
            idempotency_key=uuid4().hex, field="equipment", value="40FT"))
        self.assertFalse(self.current(item)["quote_confirmations"]["conditions"]["equipment"]["valid"])
        self.assertEqual(self.events(item), old_events)
        self.assertIn("適用条件変更後の費目未確認", review_facts(self.current(item))["review_reasons"])

    def test_amount_changed_then_restored_does_not_resurrect_old_source_confirmation(self):
        item = self.line(self.case())
        self.classify(item)
        self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))
        original = deepcopy(self.current(item))
        old_events = deepcopy(self.events(item))
        for price in ("1300", "1200"):
            current = self.current(item)
            self.review.correct(self.actor, CorrectionCommand(item_id=item["item_id"],
                expected_case_revision=current["case_revision"], expected_item_version=current["version"],
                expected_memory_version=current["current_memory_version"], idempotency_key=uuid4().hex,
                unit_price=price))
            self.assertFalse(self.current(item)["quote_confirmations"]["source"]["valid"])
            self.assertTrue(review_facts(self.current(item))["review_required"])
        self.assertEqual(self.current(item)["current_price"], original["current_price"])
        self.assertEqual(self.current(item)["observation"], original["observation"])
        self.assertEqual(self.current(item)["source_ref"], original["source_ref"])
        self.assertEqual(self.events(item), old_events)

    def test_fee_catalog_rename_preserves_unrelated_source_and_condition_attestations(self):
        item = self.line(self.case())
        fee_id = self.classify(item)["fee_id"]
        self.explicitly_check_all(item)
        old_events = deepcopy(self.events(item))
        snapshot = self.store.case_snapshot(self.actor, item["case_id"])
        row = comparison_projection(snapshot["items"], snapshot["fees"], snapshot["documents"])["rows"][0]
        RowConfirmationService(self.repo).confirm(self.actor, RowConfirmationCommand(
            case_id=item["case_id"], row_id=row["row_id"],
            expected_case_revision=snapshot["case"]["case_revision"],
            members=tuple(RowMember(member["item_id"], member["version"], member["current_memory_version"])
                          for member in row["member_items"]),
            idempotency_key=uuid4().hex, intent="RENAME_FEE", expected_fee_id=fee_id,
            expected_name_version=row["rename_name_version"], new_fee_name="改名した合成海上運賃"))
        current = self.current(item)
        self.assertTrue(current["quote_confirmations"]["source"]["valid"])
        self.assertTrue(all(event["valid"] for event in current["quote_confirmations"]["conditions"].values()))
        self.assertEqual(self.events(item), old_events)

    def test_quote_attestation_does_not_mask_revocation_or_shared_key_conflict(self):
        for state in ("REVOKED", "CONFLICTED"):
            with self.subTest(state=state):
                item = self.line(self.case(), label="SYNTHETIC " + state)
                self.classify(item, name="海上運賃 " + state)
                if state == "REVOKED":
                    current = self.current(item)
                    self.review.revoke_memory(self.actor, RevokeCommand(item["item_id"],
                        current["case_revision"], current["version"], current["current_memory_version"],
                        uuid4().hex, "SYNTHETIC explicit revocation"))
                else:
                    self.classify(item, name="異なる合成費目")
                memory_before = deepcopy(self.database_snapshot()["memory_versions"])
                self.confirmations.confirm(self.actor, self.command(item))
                self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))
                current = self.current(item)
                self.assertEqual(current["current_memory_status"], state)
                expected = "確認記録が撤回されています" if state == "REVOKED" else "確認記録が競合しています"
                self.assertIn(expected, review_facts(current)["review_reasons"])
                self.assertEqual(self.database_snapshot()["memory_versions"], memory_before)
                self.assert_surfaces(self.processed_snapshot(item["case_id"]), complete=False)

    def test_unknown_document_error_is_not_cleared_by_confirmed_members(self):
        item = self.line(self.case())
        self.classify(item)
        self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))
        self.assertFalse(review_facts(self.current(item))["review_required"])
        snapshot = self.processed_snapshot(item["case_id"],
            [{"code": "SYNTHETIC_UNKNOWN_READER_FAILURE", "message": "unexplained source mismatch"}])
        self.assert_surfaces(snapshot, complete=False)

    def test_item_bound_document_issue_resolves_only_for_its_member_preserving_history(self):
        item = self.line(self.case(), source_ready=False)
        self.classify(item)
        snapshot = self.processed_snapshot(item["case_id"])
        document_id = item["observation"]["document_id"]
        historical = deepcopy(snapshot["documents"])
        before = current_document_reviews(snapshot["items"], snapshot["documents"])[document_id]
        self.assertTrue(before["review_required"])
        self.confirmations.confirm(self.actor, self.command(item, kind="SOURCE_REVIEW"))
        snapshot["items"] = self.store.list_items(self.actor, item["case_id"])
        after = current_document_reviews(snapshot["items"], snapshot["documents"])[document_id]
        self.assertEqual(after["historical_state"], "REVIEW_REQUIRED")
        self.assertEqual(after["current_state"], "AVAILABLE")
        self.assertFalse(after["review_required"])
        self.assertEqual(len(after["resolved_issues"]), 1)
        self.assertEqual(snapshot["documents"], historical)
        wrong_target = deepcopy(snapshot["documents"])
        wrong_target[0]["issues"][0]["item_id"] = "synthetic-nonexistent-item"
        self.assertTrue(current_document_reviews(snapshot["items"], wrong_target)[document_id]["review_required"])

    def test_document_coverage_completes_supported_history_without_rewriting_task(self):
        item = self.line(self.case(), source_ready=False, locator_status="PAGE_ONLY")
        self.classify(item)
        self.inject_historical_task(item, [{"code": "SDK_LIMITATION", "limitation": "READING_ORDER_UNRESOLVED"}])
        self.explicitly_check_all(item)
        before = self.database_snapshot()
        service = DocumentReviewService(self.repo, self.storage)
        command = self.document_command(item)
        result = service.confirm(self.actor, command)
        self.assertFalse(result["review_required"])
        self.assertEqual(result, service.confirm(self.actor, command))
        after = self.database_snapshot()
        for table in ("documents", "ingestion_tasks", "ingestion_events", "items", "memory_versions",
                      "memory_events", "correction_events", "quote_confirmation_events"):
            self.assertEqual(after[table], before[table], table)
        self.assertEqual(len(after["document_review_events"]), 1)
        snapshot = IngestionStore(self.repo).case_snapshot(self.actor, item["case_id"])
        self.assert_surfaces(snapshot, complete=True)
        self.assertEqual(snapshot["documents"][0]["state"], "REVIEW_REQUIRED")
        self.assertTrue(snapshot["documents"][0]["issues"])
        # Relevant value changes make the current document attestation stale;
        # the append-only event stays available as audit history.
        changed = deepcopy(snapshot)
        changed["items"][0]["current_price"] = "1300"
        document_id = item["observation"]["document_id"]
        projection = current_document_reviews(changed["items"], changed["documents"])[document_id]
        self.assertFalse(projection["coverage_confirmation"]["valid"])
        self.assertTrue(projection["review_required"])

    def test_document_coverage_requires_member_checks_and_never_hides_unknown_issues(self):
        service = DocumentReviewService(self.repo, self.storage)
        for issue in ({"code": "SDK_LIMITATION", "limitation": "READING_ORDER_UNRESOLVED"},
                      {"code": "SOURCE_LINE_NOT_EXTRACTED", "source_block_id": "synthetic-missing-line"},
                      {"code": "SYNTHETIC_UNKNOWN_FAILURE"}):
            with self.subTest(issue=issue):
                item = self.line(self.case())
                self.classify(item, name="合成費目 " + uuid4().hex)
                self.inject_historical_task(item, [issue])
                before = self.database_snapshot()
                with self.assertRaises(ValueError):
                    service.confirm(self.actor, self.document_command(item))
                self.assertEqual(self.database_snapshot(), before)
                if issue["code"] != "SDK_LIMITATION":
                    self.explicitly_check_all(item)
                    before = self.database_snapshot()
                    with self.assertRaises(ValueError):
                        service.confirm(self.actor, self.document_command(item))
                    self.assertEqual(self.database_snapshot(), before)
                    self.assert_surfaces(IngestionStore(self.repo).case_snapshot(self.actor, item["case_id"]), complete=False)

    def test_document_coverage_count_fingerprint_scope_and_rollback_are_enforced(self):
        item = self.line(self.case())
        self.classify(item)
        self.inject_historical_task(item, [{"code": "SDK_LIMITATION", "limitation": "role_requires_review"}])
        self.explicitly_check_all(item)
        service = DocumentReviewService(self.repo, self.storage)
        command = self.document_command(item)
        before = self.database_snapshot()
        for changes, error in (({"confirmed_item_count": 2}, ValueError),
                               ({"expected_review_sha256": "f" * 64}, ConflictError),
                               ({"expected_case_revision": 1}, ConflictError)):
            with self.subTest(changes=changes), self.assertRaises(error):
                service.confirm(self.actor, replace(command, **changes))
            self.assertEqual(self.database_snapshot(), before)
        for point in ("after_document_review", "before_document_review_command"):
            def fail(stage):
                if stage == point:
                    raise RuntimeError("SYNTHETIC_DOCUMENT_ROLLBACK")
            with self.assertRaisesRegex(RuntimeError, "SYNTHETIC_DOCUMENT_ROLLBACK"):
                DocumentReviewService(self.repo, self.storage, fault_hook=fail).confirm(self.actor, command)
            self.assertEqual(self.database_snapshot(), before)
        forbidden = Principal(self.tenant, "synthetic-forbidden", allowed_case_ids=frozenset())
        with self.assertRaises(ScopeError):
            service.confirm(forbidden, command)
        service.confirm(self.actor, command)
        after = self.database_snapshot()
        with self.assertRaises(ConflictError):
            service.confirm(self.actor, replace(command, evidence_note="different evidence"))
        self.assertEqual(self.database_snapshot(), after)


if __name__ == "__main__":
    unittest.main()
