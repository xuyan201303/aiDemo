"""Synthetic Web/XLSX review facts and conservative case-list status contracts."""
from dataclasses import asdict
from io import BytesIO
import unittest
from types import SimpleNamespace

from openpyxl import load_workbook

from sanzo_core.domain import Context, Observation
from sanzo_demo.case_overview import case_summaries, summarize_case
from sanzo_demo.comparison import comparison_projection
from sanzo_demo.export import export_review_state, workbook_bytes
from sanzo_demo.review_state import review_facts


def sample(name, *, fee_id="stable-fee", status="HUMAN_CONFIRMED", memory="ACTIVE",
           initial=None, company="Forwarder A", ready=True, source_ready=True,
           locator="VERIFIED", corrected=False):
    observation = Observation("isolated-review", name, "review-case", "doc", "source-" + name,
        company, "RAW " + name, "1200", "USD",
        Context("SEA", "20FT", "PER_CONTAINER", "ORIGIN", "STANDARD", ready), source_ready)
    row = {
        "item_id": name, "observation": asdict(observation),
        "source_ref": {"synthetic": True, "kind": "EXCEL_RANGE", "document_id": "doc",
                       "source_block_ids": [observation.source_id], "locator_status": locator,
                       "locator": {"sheet_name": "Sheet1", "A1_range": "B14"}},
        "input_kind": "SYNTHETIC_STRUCTURED", "current_fee_id": fee_id,
        "current_fee_name": "海上運賃" if fee_id else None,
        "effective_fee_name": "海上運賃" if fee_id else None,
        "current_price": "1200", "currency": "USD", "raw_value": "1200",
        "saved_status": status, "current_memory_status": memory,
        "current_memory_version": 1 if memory != "NONE" else 0,
        "initial_classification": initial or {}, "version": 1,
        "ai_suggested_name": "海上運賃" if fee_id is None else None,
    }
    if corrected:
        row["effective_context"] = {**row["observation"]["context"], "equipment": "40FT"}
        row["condition_overrides"] = {"equipment": {"original_value": "20FT",
            "human_confirmed_value": "40FT", "effective_value": "40FT"}}
        row["future_reuse_eligible"] = False
        row["classification_confirmed_after_condition"] = True
    return row


class SharedReviewStateTests(unittest.TestCase):
    def test_export_labels_current_catalog_name_and_adopted_history_separately(self):
        member = sample("historical")
        book = load_workbook(BytesIO(workbook_bytes(
            {"name": "Synthetic rename", "case_revision": 2}, [member],
            [{"document_id": "doc", "original_filename": "synthetic.xlsx",
              "state": "AVAILABLE", "sha256": "a" * 64}],
            fees=[{"fee_id": "stable-fee", "display_name": "海上輸送費", "name_version": 2}])),
            read_only=True)
        try:
            self.assertEqual(book["原価単価一覧表"]["A4"].value, "海上輸送費（20FT）")
            self.assertEqual(book["明細と出典"]["D1"].value, "採用時の費目名")
            self.assertEqual(book["明細と出典"]["D2"].value, "海上運賃")
            self.assertEqual(book["確認情報"]["A2"].value, "海上輸送費（20FT）")
            self.assertEqual(book["確認情報"]["F1"].value, "採用時の費目名")
            self.assertEqual(book["確認情報"]["F2"].value, "海上運賃")
        finally:
            book.close()

    def test_web_cells_and_xlsx_share_five_facts_and_only_real_issues(self):
        rows = [
            sample("human"),
            sample("applied", status="MEMORY_APPLIED", initial={"kind": "MEMORY_APPLIED"}),
            sample("eligible-unlearned", fee_id=None, status="AI_SUGGESTED", memory="NONE"),
            sample("conflicted", memory="CONFLICTED"),
            sample("revoked", memory="REVOKED"),
            sample("corrected", corrected=True),
            sample("incomplete", fee_id=None, status="AI_SUGGESTED", memory="NONE",
                   ready=False, source_ready=False, locator="UNAVAILABLE"),
        ]
        expected = {
            "human": (True, True, "ACTIVE", False, False),
            "applied": (True, True, "ACTIVE", True, False),
            "eligible-unlearned": (False, True, "NONE", False, True),
            "conflicted": (True, True, "CONFLICTED", False, True),
            "revoked": (True, True, "REVOKED", False, True),
            "corrected": (True, False, "ACTIVE", False, False),
            "incomplete": (False, False, "NONE", False, True),
        }
        documents = [{"document_id": "doc", "original_filename": "synthetic.xlsx",
                      "state": "AVAILABLE", "sha256": "a" * 64}]
        projection = comparison_projection(rows, [{"fee_id": "stable-fee", "display_name": "海上運賃",
                                                   "name_version": 1}], documents)
        web = {member["item_id"]: member["review_facts"]
               for group in projection["rows"] for member in group["member_items"]}
        for group in projection["rows"]:
            for cells in group["cells"].values():
                for cell in cells:
                    self.assertEqual(cell["review_facts"], web[cell["item_id"]])
        for name, facts in web.items():
            self.assertEqual(tuple(facts[key] for key in (
                "classification_confirmed", "learning_eligible", "reusable_memory",
                "memory_applied", "review_required")), expected[name])
            self.assertEqual(facts, review_facts(next(row for row in rows if row["item_id"] == name)))
        self.assertIn("次回自動適用は未確認", " ".join(web["corrected"]["learning_notices"]))
        self.assertEqual(web["corrected"]["review_reasons"], [])
        self.assertEqual(web["corrected"]["reusable_memory"], "ACTIVE")

        binary = workbook_bytes({"name": "Synthetic review", "case_revision": 1}, rows,
                                documents, fees=[{"fee_id": "stable-fee", "display_name": "海上運賃",
                                                  "name_version": 1}])
        book = load_workbook(BytesIO(binary), read_only=True)
        try:
            audit = list(book["確認情報"].values)
            header = {name: column for column, name in enumerate(audit[0])}
            self.assertTrue({"費目確認", "学習適格", "有効な学習記録", "今回の学習適用", "要確認理由"} <= set(header))
            self.assertIn("学習上の注意", header)
            by_raw = {row[header["原文費目"]]: row for row in audit[1:]}
            for name, facts in web.items():
                cells = by_raw["RAW " + name]
                self.assertEqual(cells[header["費目確認"]], "確認済み" if facts["classification_confirmed"] else "未確認")
                self.assertEqual(cells[header["学習適格"]], "適格" if facts["learning_eligible"] else "対象外")
                self.assertEqual(cells[header["有効な学習記録"]], facts["reusable_memory"])
                self.assertEqual(cells[header["今回の学習適用"]], "適用あり" if facts["memory_applied"] else "適用なし")
                self.assertEqual(cells[header["要確認理由"]] or "", " / ".join(facts["review_reasons"]))
                self.assertEqual(cells[header["学習上の注意"]] or "", " / ".join(facts["learning_notices"]))
            issues = {line[0] for line in list(book["未確認事項"].values)[2:]}
            self.assertEqual(issues, {name for name, facts in web.items() if facts["review_required"]})
        finally:
            book.close()

    def test_future_learning_notice_does_not_draft_a_fully_confirmed_quote(self):
        item = sample("future-only", corrected=True, memory="NONE")
        document = {"document_id": "doc", "original_filename": "synthetic.xlsx",
                    "state": "AVAILABLE", "sha256": "a" * 64, "issues": []}
        facts = review_facts(item)
        self.assertTrue(facts["classification_confirmed"])
        self.assertFalse(facts["learning_eligible"])
        self.assertFalse(facts["review_required"])
        self.assertEqual(facts["review_reasons"], [])
        self.assertIn("次回自動適用は未確認", " ".join(facts["learning_notices"]))
        snapshot = {"case": {"case_id": "review-case", "name": "Synthetic", "case_revision": 1},
                    "documents": [document], "items": [item]}
        self.assertEqual(summarize_case(snapshot)["status"], "確認済み")
        self.assertFalse(export_review_state([item], [document])[0])
        projection = comparison_projection([item], [], [document])["rows"][0]
        self.assertEqual(projection["confirmation_status"], "✓ 確認済み")
        self.assertFalse(projection["cells"]["Forwarder A"][0]["needs_review"])
        book = load_workbook(BytesIO(workbook_bytes(snapshot["case"], [item], [document])), read_only=True)
        try:
            self.assertEqual(book["原価単価一覧表"]["A2"].value, "保存済み案件スナップショット")
            self.assertEqual(book["未確認事項"].max_row, 2)
            self.assertEqual(book["確認情報"]["K2"].value, None)
            self.assertIn("次回自動適用は未確認", book["確認情報"]["N2"].value)
        finally:
            book.close()

    def test_unconfirmed_after_overlay_and_real_current_issues_still_block(self):
        item = sample("before-confirm", corrected=True, memory="NONE")
        item["classification_confirmed_after_condition"] = False
        facts = review_facts(item)
        self.assertFalse(facts["classification_confirmed"])
        self.assertTrue(facts["review_required"])
        self.assertIn("適用条件変更後の費目未確認", facts["review_reasons"])
        item["classification_confirmed_after_condition"] = True
        for change, reason in [
            ({"effective_context": {**item["effective_context"], "equipment": None}}, "適用条件未確認"),
            ({"current_price": None}, "単価未確認"),
            ({"currency": None}, "通貨未確認"),
            ({"observation": {**item["observation"], "source_ready": False}}, "原文未確認"),
            ({"source_ref": {**item["source_ref"], "locator_status": "UNAVAILABLE"}}, "原文位置未取得"),
            ({"current_memory_status": "REVOKED"}, "確認記録が撤回されています"),
        ]:
            with self.subTest(reason=reason):
                self.assertIn(reason, review_facts({**item, **change})["review_reasons"])

    def test_guarded_future_rule_is_notice_only_after_explicit_current_confirmation(self):
        item = sample("guarded", corrected=True, memory="CONFLICTED")
        item["memory_resolution"] = {"kind": "CONDITION_CORRECTION_GUARD", "version": 2,
                                     "guarded_direct_status": "CONFLICTED",
                                     "guarded_independent_conflict": False,
                                     "guarded_independent_revocation": False}
        facts = review_facts(item)
        self.assertFalse(facts["review_required"])
        self.assertEqual(facts["reusable_memory"], "CONFLICTED")
        self.assertIn("競合", " ".join(facts["learning_notices"]))
        item["classification_confirmed_after_condition"] = False
        self.assertTrue(review_facts(item)["review_required"])
        item["classification_confirmed_after_condition"] = True
        item["memory_resolution"] = {"kind": "DIRECT", "version": 2}
        self.assertIn("確認記録が競合しています", review_facts(item)["review_reasons"])

    def test_true_current_conflict_or_failed_document_still_marks_list_and_xlsx_pending(self):
        case = {"case_id": "review-case", "name": "Synthetic", "case_revision": 1}
        available = {"document_id": "doc", "original_filename": "synthetic.xlsx",
                     "state": "AVAILABLE", "issues": [], "sha256": "a" * 64}
        item = sample("conflicted", memory="CONFLICTED")
        self.assertTrue(review_facts(item)["review_required"])
        self.assertTrue(export_review_state([item], [available])[0])
        self.assertEqual(summarize_case({"case": case, "documents": [available],
                                         "items": [item]})["status"], "確認必要")
        book = load_workbook(BytesIO(workbook_bytes(case, [item], [available])), read_only=True)
        try:
            self.assertIn("DRAFT", book["原価単価一覧表"]["A2"].value)
            self.assertIn("確認記録が競合しています", book["未確認事項"]["C3"].value)
        finally:
            book.close()

        confirmed = sample("complete", corrected=True, memory="NONE")
        failed = {**available, "state": "FAILED", "issues": [{"code": "TEST_READ_FAILED"}]}
        self.assertFalse(review_facts(confirmed)["review_required"])
        self.assertTrue(export_review_state([confirmed], [failed])[0])
        self.assertEqual(summarize_case({"case": case, "documents": [failed],
                                         "items": [confirmed]})["status"], "確認必要")
        book = load_workbook(BytesIO(workbook_bytes(case, [confirmed], [failed])), read_only=True)
        try:
            self.assertIn("DRAFT", book["原価単価一覧表"]["A2"].value)
            self.assertIn("TEST_READ_FAILED", book["未確認事項"]["C3"].value)
        finally:
            book.close()

    def test_case_summary_uses_conservative_saved_state(self):
        case = {"case_id": "review-case", "name": "Synthetic", "updated_at": "2026-09-29"}
        document = {"document_id": "doc", "state": "AVAILABLE"}
        snapshots = [
            ({"case": case, "documents": [], "items": []}, "資料なし"),
            ({"case": case, "documents": [{**document, "state": "READING"}], "items": []}, "解析中"),
            ({"case": case, "documents": [document], "items": [sample("pending", fee_id=None)]}, "確認必要"),
            ({"case": case, "documents": [document], "items": [sample("confirmed")]}, "確認済み"),
        ]
        for snapshot, status in snapshots:
            with self.subTest(status=status):
                summary = summarize_case(snapshot)
                self.assertEqual(summary["status"], status)
                self.assertEqual(summary["document_count"], len(snapshot["documents"]))
                self.assertEqual(summary["case_id"], case["case_id"])

    def test_case_list_retains_persisted_timestamps(self):
        listed = {"case_id": "review-case", "name": "Synthetic", "case_revision": 1,
                  "created_at": "2026-09-28", "updated_at": "2026-09-29"}
        store = SimpleNamespace(list_cases=lambda actor: [listed])
        pipeline = SimpleNamespace(store=SimpleNamespace(case_snapshot=lambda actor, case_id: {
            "case": {"case_id": case_id, "name": "Synthetic", "case_revision": 1},
            "documents": [], "items": []}))
        summary = case_summaries(object(), store, pipeline)[0]
        self.assertEqual(summary["updated_at"], listed["updated_at"])
        self.assertEqual(summary["created_at"], listed["created_at"])
