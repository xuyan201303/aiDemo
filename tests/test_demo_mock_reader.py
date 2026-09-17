"""Raw fixture and file binding checks. No AI, SDK or business processing."""
from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
import pypdfium2 as pdfium

from sanzo_demo.extraction import complete_field_unit, record_locations
from sanzo_demo.mock_reader import MockSDKReader

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "demo_materials" / "fixtures"
SAMPLES = ROOT / "demo_materials" / "samples"


class MockSDKReaderTests(unittest.TestCase):
    def setUp(self):
        self.folder = TemporaryDirectory(prefix="sanzo-mock-reader-")
        self.addCleanup(self.folder.cleanup)
        self.directory = Path(self.folder.name)
        self.reader = MockSDKReader(FIXTURES)
        self.manifest = json.loads((FIXTURES / "manifest.json").read_text())

    def document(self, sample, document_id="current-upload-id"):
        return {"document_id": document_id, "case_id": "fixture-contract-only-no-business-case",
                "sha256": sha256(sample.read_bytes()).hexdigest()}

    def copied_reader(self):
        destination = self.directory / "fixtures"
        shutil.copytree(FIXTURES, destination)
        return destination, MockSDKReader(destination)

    def rewrite_fixture(self, folder, filename, change):
        path = folder / filename
        payload = json.loads(path.read_text())
        change(payload)
        data = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode()
        path.write_bytes(data)
        manifest_path = folder / "manifest.json"
        manifest = json.loads(manifest_path.read_text())
        for entry in manifest["fixtures"]:
            if entry["fixture_file"] == filename:
                entry["fixture_sha256"] = sha256(data).hexdigest()
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")

    def test_all_three_original_hashes_load_only_raw_mock_fields(self):
        for entry in self.manifest["fixtures"]:
            with self.subTest(entry["sample_filename"]):
                path = SAMPLES / entry["sample_filename"]
                reading = self.reader.read(path, self.document(path))
                self.assertEqual(reading.origin, "MOCK_SDK")
                self.assertEqual(reading.sdk_identity["real_sdk_execution"], "NOT_RUN")
                self.assertEqual(reading.sdk_identity["fixture_integrity"], "VERIFIED")
                self.assertEqual(reading.sdk_identity["fixture_sha256"], entry["fixture_sha256"])
                self.assertEqual(len([b for b in reading.blocks
                    if b["source_ref"]["evidence_unit"]["field_name"] == "raw_label"]), 5)
                for block in reading.blocks:
                    self.assertNotIn("fee_id", block)
                    self.assertEqual(block["source_ref"]["origin"], "MOCK_SDK")
                    complete_field_unit(block, block["source_ref"]["evidence_unit"]["field_name"])
                _, ambiguous = record_locations(reading)
                self.assertFalse(ambiguous)

    def test_renaming_uploaded_file_does_not_change_hash_selection(self):
        source = SAMPLES / self.manifest["fixtures"][0]["sample_filename"]
        renamed = self.directory / "different-name.xlsx"
        shutil.copyfile(source, renamed)
        document = self.document(renamed)
        document["original_filename"] = "arbitrary-display-only.xlsx"
        self.assertEqual(self.reader.read(source, document), self.reader.read(renamed, document))

    def test_modified_bytes_reject_stale_declared_hash(self):
        source = SAMPLES / self.manifest["fixtures"][0]["sample_filename"]
        changed = self.directory / source.name
        changed.write_bytes(source.read_bytes() + b"changed")
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_UPLOAD_HASH_MISMATCH"):
            self.reader.read(changed, self.document(source))

    def test_modified_bytes_with_updated_declared_hash_still_unsupported(self):
        source = SAMPLES / self.manifest["fixtures"][0]["sample_filename"]
        changed = self.directory / source.name
        changed.write_bytes(source.read_bytes() + b"changed")
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_UNSUPPORTED_FILE"):
            self.reader.read(changed, self.document(changed))

    def test_unlisted_filename_cannot_select_nearest_fixture(self):
        unknown = self.directory / "01_A社_初回見積.xlsx"
        unknown.write_bytes(b"this is explicitly unsupported, not a workbook fixture")
        with self.assertRaisesRegex(ValueError, "当前模拟模式仅支持三份指定样例"):
            self.reader.read(unknown, self.document(unknown))

    def test_each_read_binds_new_document_id_without_mutating_fixture(self):
        entry = self.manifest["fixtures"][0]
        sample = SAMPLES / entry["sample_filename"]
        before = (FIXTURES / entry["fixture_file"]).read_bytes()
        one = self.reader.read(sample, self.document(sample, "upload-one"))
        two = self.reader.read(sample, self.document(sample, "upload-two"))
        self.assertEqual(one.document_sha256, two.document_sha256)
        self.assertTrue(all(b["source_ref"]["document_id"] == "upload-one" for b in one.blocks))
        self.assertTrue(all(b["source_ref"]["document_id"] == "upload-two" for b in two.blocks))
        self.assertEqual(before, (FIXTURES / entry["fixture_file"]).read_bytes())

    def test_fixture_tampering_fails_digest_before_use(self):
        folder, reader = self.copied_reader()
        entry = self.manifest["fixtures"][0]
        sample = SAMPLES / entry["sample_filename"]
        path = folder / entry["fixture_file"]
        path.write_bytes(path.read_bytes() + b" ")
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_FIXTURE_HASH_MISMATCH"):
            reader.read(sample, self.document(sample))

    def test_manifest_cannot_escape_fixture_directory(self):
        folder, _ = self.copied_reader()
        path = folder / "manifest.json"
        manifest = json.loads(path.read_text())
        manifest["fixtures"][0]["fixture_file"] = "../outside.raw.json"
        path.write_text(json.dumps(manifest))
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_INVALID_MANIFEST_ENTRY"):
            MockSDKReader(folder)

    def test_fixture_symlink_is_rejected(self):
        folder, reader = self.copied_reader()
        entry = self.manifest["fixtures"][0]
        fixture = folder / entry["fixture_file"]
        fixture.unlink()
        fixture.symlink_to(FIXTURES / entry["fixture_file"])
        sample = SAMPLES / entry["sample_filename"]
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_FIXTURE_PATH_UNSAFE"):
            reader.read(sample, self.document(sample))

    def test_answer_data_forbidden_even_if_manifest_digest_matches(self):
        folder, _ = self.copied_reader()
        entry = self.manifest["fixtures"][0]
        self.rewrite_fixture(folder, entry["fixture_file"],
                             lambda payload: payload["blocks"][0].update(fee_id="forbidden"))
        reader = MockSDKReader(folder)
        sample = SAMPLES / entry["sample_filename"]
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_ANSWER_DATA_FORBIDDEN"):
            reader.read(sample, self.document(sample))

    def test_static_document_binding_in_fixture_forbidden(self):
        folder, _ = self.copied_reader()
        entry = self.manifest["fixtures"][0]
        self.rewrite_fixture(folder, entry["fixture_file"],
            lambda payload: payload["blocks"][0]["source_ref"].update(document_id="preloaded"))
        sample = SAMPLES / entry["sample_filename"]
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_STATIC_DOCUMENT_ID_FORBIDDEN"):
            MockSDKReader(folder).read(sample, self.document(sample))

    def test_manifest_duplicate_hash_rejected(self):
        folder, _ = self.copied_reader()
        path = folder / "manifest.json"
        manifest = json.loads(path.read_text())
        manifest["fixtures"][1]["sample_sha256"] = manifest["fixtures"][0]["sample_sha256"]
        path.write_text(json.dumps(manifest))
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_INVALID_MANIFEST_ENTRY"):
            MockSDKReader(folder)

    def test_manifest_duplicate_json_key_rejected(self):
        folder, _ = self.copied_reader()
        path = folder / "manifest.json"
        raw = path.read_text().replace('"schema_version": 1,', '"schema_version": 1, "schema_version": 1,')
        path.write_text(raw)
        with self.assertRaisesRegex(ValueError, "MOCK_SDK_DUPLICATE_JSON_KEY"):
            MockSDKReader(folder)

    def test_constructor_reads_manifest_only_and_no_sample_business_processing(self):
        reads = []
        original = MockSDKReader._read_file
        def tracked(reader, filename):
            reads.append(filename)
            return original(reader, filename)
        with patch.object(MockSDKReader, "_read_file", tracked):
            reader = MockSDKReader(FIXTURES)
        self.assertEqual(reads, ["manifest.json"])
        self.assertEqual(len(reader.sample_hashes), 3)

    def test_excel_fixture_has_exact_full_original_cells_and_current_locations(self):
        for entry in self.manifest["fixtures"]:
            if not entry["sample_filename"].endswith(".xlsx"):
                continue
            sample = SAMPLES / entry["sample_filename"]
            reading = self.reader.read(sample, self.document(sample))
            with sample.open("rb") as stream:
                workbook = load_workbook(stream, read_only=True, data_only=False)
                try:
                    for block in reading.blocks:
                        source = block["source_ref"]
                        locator = source["locator"]
                        cell = workbook[locator["sheet_name"]][locator["A1_range"]]
                        self.assertEqual(block["raw_text"], str(cell.value))
                        self.assertEqual(source["fixture_boundary"]["original_number_format"], cell.number_format)
                finally:
                    workbook.close()

    def test_pdf_fixture_rows_match_complete_original_printed_rows(self):
        entry = next(v for v in self.manifest["fixtures"] if v["sample_filename"].endswith(".pdf"))
        sample = SAMPLES / entry["sample_filename"]
        reading = self.reader.read(sample, self.document(sample))
        document = pdfium.PdfDocument(sample)
        try:
            self.assertEqual(len(document), 1)
            page = document[0]
            try:
                text_page = page.get_textpage()
                try:
                    original_lines = text_page.get_text_range().splitlines()
                finally:
                    text_page.close()
            finally:
                page.close()
        finally:
            document.close()
        records = {}
        for block in reading.blocks:
            source = block["source_ref"]
            self.assertEqual(source["locator_status"], "PAGE_ONLY")
            self.assertEqual(source["locator"], {"page_number": 1})
            unit = source["evidence_unit"]
            self.assertEqual(unit["kind"], "FIXTURE_FIELD")
            if unit["scope"] == "LINE":
                records.setdefault(unit["record_id"], {})[unit["field_name"]] = block["raw_text"]
        self.assertEqual(len(records), 5)
        for record in records.values():
            self.assertNotIn("transport_mode", record)
            complete_line = " ".join(record[key] for key in (
                "raw_label", "equipment", "currency", "raw_value", "billing_unit"))
            complete_line += " " + record["charge_scope"] + "／" + record["extra_terms"]
            self.assertIn(complete_line, original_lines)

    def test_mock_reader_available_when_real_sdk_import_is_forbidden(self):
        script = (
            "import builtins, pathlib, sys\n"
            "sys.path.insert(0, 'backend')\n"
            "original = builtins.__import__\n"
            "def guarded(name, *args, **kwargs):\n"
            "    if name == 'document_sdk' or name.startswith('document_sdk.'):\n"
            "        raise AssertionError('real SDK import forbidden')\n"
            "    return original(name, *args, **kwargs)\n"
            "builtins.__import__ = guarded\n"
            "from sanzo_demo.mock_reader import MockSDKReader\n"
            "reader = MockSDKReader()\n"
            "assert reader.reader_mode == 'MOCK_SDK' and not reader.is_real_sdk\n"
            "assert len(reader.sample_hashes) == 3\n"
        )
        result = subprocess.run([sys.executable, "-B", "-c", script], cwd=ROOT,
                                capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)

    def mock_payload(self, reading):
        """MECHANISM ONLY: copied raw fields, no real AI or business insertion."""
        records, shared = {}, {}
        for block in reading.blocks:
            unit = block["source_ref"]["evidence_unit"]
            field = {"value": block["raw_text"], "source_block_id": block["source_id"]}
            if unit["scope"] == "DOCUMENT":
                for record_id in unit["applies_to_record_ids"]:
                    shared.setdefault(record_id, {})[unit["field_name"]] = field
            else:
                records.setdefault(unit["record_id"], {})[unit["field_name"]] = field
        entries = []
        for record_id, fields in records.items():
            fields = {**shared.get(record_id, {}), **fields}
            unknown = {"value": None, "source_block_id": None}
            entries.append({**{name: deepcopy(fields.get(name, unknown))
                for name in ("raw_label", "raw_value", "currency", "company")},
                "context": {name: deepcopy(fields.get(name, unknown))
                for name in ("transport_mode", "equipment", "billing_unit", "charge_scope", "extra_terms")}})
        return {"lines": entries, "document_issues": []}

    def test_f1_first_round_fixture_fields_materialize_without_business_write(self):
        from sanzo_demo.extraction import materialize
        for entry in self.manifest["fixtures"]:
            # Third-file checks stop at ReadResult/raw cells. Its business
            # Observations are reserved for the actual second round.
            if entry["fixture_file"] == "03.raw.json":
                continue
            sample = SAMPLES / entry["sample_filename"]
            document = self.document(sample)
            reading = self.reader.read(sample, document)
            lines, issues = materialize(document, "fixture-unit-test-no-database", reading, self.mock_payload(reading))
            self.assertEqual(len(lines), 5)
            if sample.suffix == ".pdf":
                self.assertTrue(all(not line.observation.context.complete() for line in lines))
                self.assertEqual({issue["code"] for issue in issues}, {"APPLICABILITY_INCOMPLETE"})
            else:
                self.assertEqual(issues, [])
                self.assertTrue(all(line.observation.context.complete() for line in lines))
            self.assertTrue(all(line.source_ref["origin"] == "MOCK_SDK" for line in lines))

    def test_f1_shortened_amount_label_and_condition_are_rejected(self):
        from sanzo_demo.extraction import materialize
        sample = SAMPLES / self.manifest["fixtures"][0]["sample_filename"]
        document = self.document(sample)
        reading = self.reader.read(sample, document)
        original = self.mock_payload(reading)
        mutations = [
            (3, "raw_value", "1200"),
            (4, "raw_label", "配送"),
            (4, "equipment", "通"),
        ]
        for index, field, shorter in mutations:
            with self.subTest(field):
                payload = deepcopy(original)
                target = payload["lines"][index]
                (target["context"] if field == "equipment" else target)[field]["value"] = shorter
                lines, issues = materialize(document, "fixture-unit-test-no-database", reading, payload)
                self.assertEqual(len(lines), 4)
                self.assertTrue(any(issue.get("line_index") == index for issue in issues))

    def test_f1_same_amount_from_another_fixture_row_is_rejected(self):
        from sanzo_demo.extraction import materialize
        sample = SAMPLES / self.manifest["fixtures"][0]["sample_filename"]
        document = self.document(sample)
        reading = self.reader.read(sample, document)
        payload = self.mock_payload(reading)
        self.assertEqual(payload["lines"][2]["raw_value"]["value"], payload["lines"][4]["raw_value"]["value"])
        payload["lines"][4]["raw_value"] = deepcopy(payload["lines"][2]["raw_value"])
        lines, issues = materialize(document, "fixture-unit-test-no-database", reading, payload)
        self.assertEqual(len(lines), 4)
        self.assertTrue(any(issue.get("line_index") == 4 for issue in issues))


if __name__ == "__main__":
    unittest.main()
