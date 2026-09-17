"""Five-row browser/XLSX projection check using synthetic GET stubs only.

This is not real AI output, sample-file processing, persistence or Demo acceptance.
"""
import argparse
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import traceback
from urllib.parse import urlsplit

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from sanzo_demo.comparison import comparison_projection
from sanzo_demo.export import workbook_bytes
from test_demo_comparison_projection import quote

ROOT = Path(__file__).resolve().parents[1]


def run(out):
    out.mkdir(parents=True, exist_ok=False)
    requests, errors, checks = [], [], []
    result = {
        "scope": "OFFLINE_SYNTHETIC_BROWSER_AND_GENERATED_XLSX",
        "real_ai_calls": 0, "real_sdk_calls": 0, "database_writes": 0,
        "sample_files_read": 0,
        "limitations": ["All quote values and AI suggestions are synthetic test data.",
                       "HTTP is replaced with GET-only local stand-ins.",
                       "No real AI, sample ingestion, learning or database acceptance is claimed."],
        "checks": checks, "requests": requests,
    }
    try:
        items = []
        for n, (name, equipment) in enumerate((("合成・輸送費", "20FT"), ("合成・輸送費", "40FT"),
                                               ("合成・書類費", "共通"), ("合成・通関費", "共通"), ("合成・配送費", "共通"))):
            items += [quote("a" + str(n), "テストA社", name=name, equipment=equipment,
                            price=str(1010 + n), fee_id="synthetic-confirmed-id" if n == 4 else None),
                      quote("b" + str(n), "テストB社", name=name, equipment=equipment,
                            price=str(910 + n), transport_mode=None, ready=False)]
        documents = [{"document_id": "doc-" + company, "original_filename": "synthetic-" + company + ".xlsx",
                      "sha256": "a" * 64, "state": "REVIEW_REQUIRED"} for company in ("テストA社", "テストB社")]
        case = {"case_id": "synthetic-case", "name": "オフライン合成テスト・実AI結果ではありません", "case_revision": 1}
        fees = [{"fee_id": "synthetic-confirmed-id", "display_name": "合成・配送費"}]
        projection = comparison_projection(items, fees)
        snapshot = {"case": case, "documents": documents, "items": items, "tasks": [], "fees": fees,
                    "comparison": projection}
        before = deepcopy(snapshot)
        generated_xlsx = workbook_bytes(case, items, documents, fees=fees, reader_mode="MOCK_SDK")
        serve_projection = True
        with sync_playwright() as pw:
            browser = pw.chromium.launch(executable_path="/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", headless=True)
            try:
                page = browser.new_page(viewport={"width": 1500, "height": 1150}, accept_downloads=True)
                page.on("pageerror", lambda e: errors.append(str(e)))
                def route(r):
                    parsed = urlsplit(r.request.url)
                    path = parsed.path
                    requests.append({"method": r.request.method, "path": path, "query": parsed.query})
                    assert parsed.hostname == "127.0.0.1", "External request prohibited"
                    assert r.request.method == "GET", "Mutation request prohibited"
                    if path in ("/", "/index.html"):
                        return r.fulfill(path=str(ROOT / "backend/sanzo_demo/static/index.html"))
                    if path.startswith("/static/"):
                        filename = path.rsplit("/", 1)[-1]
                        assert filename in {"app.js", "workflow.js", "style.css"}
                        return r.fulfill(path=str(ROOT / "backend/sanzo_demo/static" / filename))
                    if path == "/api/status":
                        data = {"reader_mode": "MOCK_SDK", "reader": {"ready": True}, "sdk": {"ready": False},
                                "ai": {"ready": False}, "missing_conditions": []}
                    elif path == "/api/cases":
                        data = [case]
                    elif path == "/api/cases/synthetic-case":
                        data = snapshot if serve_projection else {k: v for k, v in snapshot.items() if k != "comparison"}
                    elif path == "/api/cases/synthetic-case/export.xlsx":
                        return r.fulfill(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                         headers={"Content-Disposition": 'attachment; filename="synthetic-projection.xlsx"'},
                                         body=generated_xlsx)
                    else:
                        raise AssertionError("Unexpected request: " + path)
                    r.fulfill(content_type="application/json", body=json.dumps(data, ensure_ascii=False))

                page.route("**/*", route)
                page.goto("http://127.0.0.1:8765", wait_until="networkidle")
                page.locator("#case-select").select_option(case["case_id"])
                expect(page.locator(".compare-table tbody tr")).to_have_count(5)
                expect(page.locator(".quote-card")).to_have_count(10)
                browser_rows = []
                for number, row in enumerate(projection["rows"]):
                    browser_row = page.locator(".compare-table tbody tr").nth(number)
                    expect(browser_row).to_contain_text(row["label"])
                    observed = {"item_ids": [], "cells": {}}
                    for col, company in enumerate(projection["companies"], 1):
                        cell = browser_row.locator("td").nth(col)
                        expect(cell.locator(".quote-card")).to_have_count(1)
                        quote_result = row["cells"][company][0]
                        original = next(i for i in items if i["item_id"] == quote_result["item_id"])
                        actual_id = cell.locator(".quote-card").get_attribute("data-item-id")
                        assert actual_id == original["item_id"]
                        expect(cell.locator("button,input,.quote-status,.row-status")).to_have_count(0)
                        expect(browser_row.locator("td").first.locator(".row-status")).to_have_text(row["confirmation_status"])
                        expect(cell.locator(".quote-price")).to_contain_text(format(int(original["current_price"]), ","))
                        expect(cell.locator(".quote-currency")).to_have_text(original["currency"])
                        observed["item_ids"].append(actual_id)
                        observed["cells"][company] = {"item_id": actual_id, "price": original["current_price"],
                                                     "currency": original["currency"],
                                                     "status": quote_result["status"]}
                        if company == "テストB社":
                            assert observed["cells"][company]["status"] == "● AI仮判定"
                            assert quote_result["needs_review"]
                    assert sorted(observed["item_ids"]) == sorted(row["item_ids"])
                    browser_rows.append(observed)
                checks.append("Browser consumes Python projection: five rows, ten quote IDs, two companies")
                checks.append("Every B item retains AI provisional / review-required status")
                table_text = page.locator("#comparison-table").inner_text()
                assert "費目未確認" not in table_text
                assert "transport_mode" not in table_text
                assert "synthetic-" not in table_text
                page.screenshot(path=str(out / "browser_five_pairs.png"), full_page=True)
                page.locator("#comparison-table").screenshot(path=str(out / "browser_table.png"))

                with page.expect_download() as pending:
                    page.locator("#export-button").click()
                download = pending.value
                export_path = out / "generated_projection.xlsx"
                download.save_as(export_path)
                exported = export_path.read_bytes()
                assert sha256(exported).digest() == sha256(generated_xlsx).digest()
                book = load_workbook(BytesIO(exported), data_only=False)
                export_rows = []
                try:
                    sheet = book["原価単価一覧表"]
                    assert sheet.max_row == 8
                    assert [sheet.cell(3, c).value for c in (2, 3)] == projection["companies"]
                    for number, row in enumerate(projection["rows"], 4):
                        assert row["label"] in sheet.cell(number, 1).value
                        rendered = {"label": sheet.cell(number, 1).value, "item_ids": browser_rows[number - 4]["item_ids"], "cells": {}}
                        for col, company in enumerate(projection["companies"], 2):
                            text = sheet.cell(number, col).value
                            observed = browser_rows[number - 4]["cells"][company]
                            assert observed["price"] + " " + observed["currency"] in text
                            assert observed["status"] not in text
                            rendered["cells"][company] = text
                        export_rows.append(rendered)
                    detail = book["明細と出典"]
                    assert sorted(detail.cell(r, 1).value for r in range(2, detail.max_row + 1)) == sorted(i["item_id"] for i in items)
                    assert "DRAFT" in sheet["A2"].value
                    assert "SDK未接続" not in sheet["B2"].value
                    assert "確認情報" in book.sheetnames
                finally:
                    book.close()
                checks.append("Actual browser download equals generated workbook bytes and shared projection values/statuses")
                checks.append("Workbook detail item IDs cover all ten browser quote IDs; draft and SDK Mock markers retained")
                assert snapshot == before
                assert page.evaluate("state.snapshot.items.filter(i => i.observation.forwarder_id === 'テストB社').every(i => i.current_fee_id === null && i.current_memory_status === null)")
                checks.append("Browser display leaves all B classifications and memory fields unconfirmed")
                result["projection"] = projection
                result["browser_rows"] = browser_rows
                result["export_rows"] = export_rows
                result["export_sha256"] = sha256(exported).hexdigest()

                serve_projection = False
                page.reload(wait_until="networkidle")
                page.locator("#case-select").select_option(case["case_id"])
                expect(page.locator("#comparison-table")).to_contain_text("比較情報を取得できません")
                expect(page.locator(".quote-card")).to_have_count(0)
                page.locator("#comparison-table").screenshot(path=str(out / "browser_missing_projection.png"))
                checks.append("Missing server projection reports unavailable; no client grouping fallback")
                assert snapshot == before
                assert not errors, errors
                assert all(call["method"] == "GET" for call in requests)
                checks.append("GET-only requests, no external request, mutation or browser exception")
            finally:
                browser.close()
        result.update(status="PASS", exit_code=0)
    except Exception as error:
        result.update(status="FAIL", exit_code=1, error=str(error), traceback=traceback.format_exc(), browser_errors=errors)
        raise
    finally:
        (out / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
        print(json.dumps({"status": result.get("status"), "exit_code": result.get("exit_code"),
                          "checks": len(checks), "output": str(out)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence", type=Path, required=True)
    run(parser.parse_args().evidence)
