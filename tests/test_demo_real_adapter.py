"""Public fixed-wheel adapter tests. Generated inputs are synthetic, with no AI.

Live SDK tests are separately labelled and require the approved installed wheel;
synthetic table-boundary tests exercise application mapping only.
"""
from copy import deepcopy
from dataclasses import replace
from hashlib import sha256
from importlib import metadata
import os
import json
from pathlib import Path
import sys
import tempfile
import unittest
from uuid import uuid4

from sanzo_core.domain import memory_key
from sanzo_demo.document_sdk_adapter import DocumentSdkAdapter, _header, WHEEL_SHA256
from sanzo_demo.extraction import ReadResult, materialize
from sanzo_demo.facts import structure_facts


def document(path):
    return {'document_id': str(uuid4()), 'case_id': str(uuid4()),
            'sha256': sha256(path.read_bytes()).hexdigest()}


def complete_pdf_table():
    headings = ['費目表記', '通貨', '単価', '適用条件', '課金単位',
                'サービス範囲', '付帯条件', '会社名', '輸送モード']
    values = ['Unfamiliar full charge', 'JPY', '9,500.00', '20FT', 'container',
              'port-to-port', 'no setup', 'Synthetic Carrier', 'sea']
    cells = [{'row_index': r, 'column_index': c, 'row_span': 1, 'column_span': 1,
              'text': value, 'availability': 'AVAILABLE'}
             for r, row in enumerate((headings, values)) for c, value in enumerate(row)]
    return {'pages': [{'page_number': 2, 'table_ids': ['synthetic-table']}], 'tables': [{'table_id': 'synthetic-table',
        'page_number': 2, 'source_identity': {'opaque_sdk_authority': True}, 'availability': 'AVAILABLE', 'row_count': 2,
        'column_count': len(headings), 'cells': cells}]}


class PdfPublicBoundaryTests(unittest.TestCase):
    def test_limited_table_preserves_readable_fields_and_limitations(self):
        data = complete_pdf_table()
        data['tables'][0].update(availability='LIMITED_EVIDENCE', limitations=['UNSPECIFIED_SDK_LIMIT'])
        doc, reading = self.map(data)
        lines, _ = structure_facts(doc, 'tenant', reading)
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines[0].source_ref['text_verification'], 'VERIFIED')
        self.assertEqual(lines[0].source_ref['sdk_table']['table_availability'], 'LIMITED_EVIDENCE')
        self.assertEqual(lines[0].source_ref['sdk_table']['limitations'], ['UNSPECIFIED_SDK_LIMIT'])
        self.assertFalse(lines[0].observation.source_ready)
        self.assertIsNone(memory_key(lines[0].observation, lines[0].source_ref, lines[0].input_kind))

    def test_missing_or_conflicting_page_source_retains_exception(self):
        for mode in ('missing', 'conflict'):
            data = complete_pdf_table()
            if mode == 'missing':
                data['tables'][0]['source_identity'] = None
            else:
                data['pages'][0]['table_ids'] = []
            _, reading = self.map(data)
            self.assertFalse(reading.blocks)
            self.assertEqual(reading.issues[0]['reason'],
                             'PDF_TABLE_SOURCE_MISSING' if mode == 'missing' else 'PDF_TABLE_SOURCE_CONFLICT')

    def map(self, data):
        adapter = object.__new__(DocumentSdkAdapter)
        doc = {'document_id': 'synthetic-document', 'case_id': 'synthetic-case', 'sha256': 'a' * 64}
        blocks, issues = [], []
        adapter._pdf_tables(data, doc, blocks, issues)
        reading = ReadResult(doc['sha256'], tuple(blocks), {'adapter': 'mechanism-test'},
                             origin='MECHANISM_TEST', issues=tuple(issues))
        return doc, reading

    def test_full_pdf_fields_have_same_row_and_page_only(self):
        doc, reading = self.map(complete_pdf_table())
        lines, issues = structure_facts(doc, 'synthetic-tenant', reading)
        self.assertEqual(issues, [])
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines[0].observation.unit_price, '9500.00')
        for block in reading.blocks:
            self.assertEqual(block['source_ref']['locator'], {'page_number': 2})
            self.assertEqual(block['source_ref']['locator_status'], 'PAGE_ONLY')
            self.assertNotIn('bounding_box', block['source_ref'])

    def test_merged_price_boundary_does_not_certify_amount(self):
        data = complete_pdf_table()
        next(c for c in data['tables'][0]['cells'] if c['row_index'] == 1 and c['column_index'] == 2)['column_span'] = 2
        doc, reading = self.map(data)
        lines, issues = structure_facts(doc, 'synthetic-tenant', reading)
        self.assertFalse(lines)
        self.assertTrue(any(i.get('reason') == 'PDF_FIELD_BOUNDARY_UNVERIFIED' for i in reading.issues))

    def test_duplicate_cell_is_not_arbitrarily_selected(self):
        data = complete_pdf_table(); data['tables'][0]['cells'].append(deepcopy(data['tables'][0]['cells'][-1]))
        doc, reading = self.map(data)
        self.assertFalse(reading.blocks)
        self.assertEqual(reading.issues[0]['reason'], 'PDF_TABLE_MEMBERSHIP_AMBIGUOUS')

    def test_page_text_cannot_supply_price_for_other_row(self):
        data = complete_pdf_table()
        data['pages'][0]['text'] = 'Other unrelated charge 9,500.00'
        data['tables'][0]['cells'] = [c for c in data['tables'][0]['cells']
                                     if not (c['row_index'] == 1 and c['column_index'] == 2)]
        doc, reading = self.map(data)
        lines, _ = structure_facts(doc, 'synthetic-tenant', reading)
        self.assertFalse(lines)

    def test_cross_row_same_numeric_token_cannot_supply_missing_price(self):
        data = complete_pdf_table()
        amount = next(c for c in data['tables'][0]['cells'] if c['row_index'] == 1 and c['column_index'] == 2)
        amount['row_index'] = 2; data['tables'][0]['row_count'] = 3
        doc, reading = self.map(data)
        lines, _ = structure_facts(doc, 'synthetic-tenant', reading)
        self.assertFalse(lines)

    def test_shortened_header_is_not_a_header_alias(self):
        rows = {9: {1: {'value': '費目'}, 2: {'value': '通貨'}, 3: {'value': '見積単価'}}}
        self.assertEqual(_header(rows), (None, 'HEADER_MISSING'))

    def test_multiple_complete_headers_are_ambiguous(self):
        row = {1: {'value': '費目表記'}, 2: {'value': '通貨'}, 3: {'value': '見積単価'}}
        self.assertEqual(_header({1: row, 20: deepcopy(row)}), (None, 'MULTIPLE_HEADER_ROWS'))

    def test_truncated_label_and_amount_fail_existing_f1(self):
        doc, reading = self.map(complete_pdf_table())
        fields = {b['source_ref']['evidence_unit']['field_name']:
                  {'value': b['raw_text'], 'source_block_id': b['source_id']} for b in reading.blocks}
        for field, bad in [('raw_label', 'Unfamiliar'), ('raw_value', '950'), ('equipment', '20')]:
            selected = deepcopy(fields); selected[field]['value'] = bad
            entry = {name: selected[name] for name in ('raw_label', 'raw_value', 'currency', 'company')}
            entry['context'] = {name: selected[name] for name in ('equipment', 'transport_mode', 'billing_unit', 'charge_scope', 'extra_terms')}
            lines, issues = materialize(doc, 'synthetic-tenant', reading, {'lines': [entry]})
            self.assertFalse(lines, field)
            self.assertEqual(issues[0]['code'], 'AI_VALUE_NOT_IN_SDK_ORIGINAL')


try:
    SDK_READY = sys.version_info[:2] == (3, 12) and metadata.version('document-processing-sdk') == '1.8.0'
except metadata.PackageNotFoundError:
    SDK_READY = False


@unittest.skipUnless(SDK_READY, 'fixed-wheel Python 3.12 SDK environment required; no Mock substitution')
class RealFixedWheelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.reader = DocumentSdkAdapter()

    def workbook(self, path, *, offset=0, amount=9321, formula=False, missing_company=False, duplicate_header=False):
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active; sheet.title = 'Moved sheet'
        labels = ['付帯条件', '通貨', '費目表記', '課金単位', '会社名',
                  '輸送モード', '見積単価', 'サービス範囲', '適用条件']
        values = ['explicit additional terms', 'USD', 'Never-seen charge label', 'each',
                  'Synthetic Other Forwarder', 'sea', '=1+2' if formula else amount,
                  'explicit scope', '40FT']
        for index, (header, value) in enumerate(zip(labels, values), offset + 2):
            sheet.cell(7 + offset, index, header)
            if not (missing_company and header == '会社名'):
                sheet.cell(8 + offset, index, value)
            if duplicate_header:
                sheet.cell(20 + offset, index, header)
        book.save(path); book.close()

    def test_generated_header_shift_column_reorder_and_fresh_numeric_value(self):
        with tempfile.TemporaryDirectory() as temp:
            for offset in (0, 4):
                path = Path(temp) / ('unrelated-' + str(offset) + '.xlsx')
                self.workbook(path, offset=offset, amount=9321 + offset)
                doc = document(path); result = self.reader.read(path, doc)
                lines, issues = structure_facts(doc, 'real-generated-only', result)
                self.assertEqual(issues, [])
                self.assertEqual(len(lines), 1)
                self.assertEqual(lines[0].observation.unit_price, str(9321 + offset))
                self.assertEqual(lines[0].raw_value, str(9321 + offset))
                self.assertEqual(lines[0].observation.raw_label, 'Never-seen charge label')
                self.assertEqual(result.origin, 'REAL_SDK')
                self.assertEqual(result.sdk_identity['wheel_sha256'], WHEEL_SHA256)
                self.assertIsNotNone(memory_key(lines[0].observation, lines[0].source_ref, lines[0].input_kind))

    def test_formula_value_is_not_certified_or_recalculated(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'formula.xlsx'; self.workbook(path, formula=True)
            doc = document(path); result = self.reader.read(path, doc)
            lines, issues = structure_facts(doc, 'generated', result)
            self.assertFalse(lines)
            self.assertTrue(any(i.get('reason') == 'FORMULA_VALUE_UNCONFIRMED' for i in result.issues))

    def test_missing_company_stays_unknown_not_wildcard(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'missing.xlsx'; self.workbook(path, missing_company=True)
            doc = document(path); result = self.reader.read(path, doc)
            lines, _ = structure_facts(doc, 'generated', result)
            self.assertEqual(len(lines), 1)
            self.assertIsNone(lines[0].observation.forwarder_id)
            self.assertFalse(lines[0].observation.source_ready)
            self.assertIsNotNone(memory_key(lines[0].observation, lines[0].source_ref, lines[0].input_kind))  # fee-name-v2: optional company

    def test_duplicate_headers_leave_unresolved(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'duplicates.xlsx'; self.workbook(path, duplicate_header=True)
            doc = document(path); result = self.reader.read(path, doc)
            self.assertFalse(result.blocks)
            self.assertTrue(any(i.get('reason') == 'MULTIPLE_HEADER_ROWS' for i in result.issues))

    def test_sample_header_candidates_are_facts_not_company_role_decisions(self):
        path = Path(__file__).resolve().parents[1] / 'demo_materials/samples/01_A社_初回見積.xlsx'
        if not path.is_file():
            self.skipTest('fictional source sample not mounted')
        doc = document(path); result = self.reader.read(path, doc)
        candidates = result.sdk_identity['field_binding_candidates']
        self.assertTrue(candidates)
        source_ids = {b['source_id'] for b in result.blocks}
        for candidate in candidates:
            ref = candidate['source_ref']; unit = ref['evidence_unit']
            self.assertNotIn('field_name', unit)
            self.assertEqual(unit['scope'], 'DOCUMENT')
            self.assertEqual(ref['document_id'], doc['document_id'])
            self.assertEqual(ref['document_sha256'], doc['sha256'])
            self.assertNotIn(candidate['source_id'], source_ids)
            self.assertEqual(len(unit['applies_to_record_ids']), 5)
        lines, _ = structure_facts(doc, 'real-only', result)
        for line in lines:
            self.assertIsNone(line.observation.forwarder_id)
            self.assertIsNone(line.observation.context.transport_mode)
            self.assertIsNotNone(memory_key(line.observation, line.source_ref, line.input_kind))  # full F1 fee label, optional document roles

    def test_upload_changed_hash_rejected_before_sdk(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'changed.xlsx'; self.workbook(path)
            doc = document(path); doc['sha256'] = '0' * 64
            with self.assertRaisesRegex(ValueError, 'REAL_SDK_UPLOAD_HASH_MISMATCH'):
                self.reader.read(path, doc)

    def test_three_actual_samples_read_new_sdk_values_with_retained_raw_output(self):
        samples = Path(__file__).resolve().parents[1] / 'demo_materials' / 'samples'
        expected = [('01_A社_初回見積.xlsx', '配送手数料', '8000'),
                    ('02_B社_初回見積.pdf', 'DELIVERY CHARGE', '9000.00'),
                    ('03_C社_次回見積.xlsx', '配送手数料', '9500')]
        if not all((samples / name).is_file() for name, _, _ in expected):
            self.skipTest('three fictional raw sample originals not mounted')
        for name, target, amount in expected:
            with self.subTest(sample=name):
                path = samples / name; doc = document(path)
                before = doc['sha256']; result = self.reader.read(path, doc)
                lines, issues = structure_facts(doc, 'real-sdk-sample-test', result)
                self.assertEqual(len(lines), 5, issues)
                delivery = next(line for line in lines if line.observation.raw_label == target)
                self.assertEqual(delivery.observation.unit_price, amount)
                self.assertEqual(delivery.source_ref['fact_sources']['amount'], 'REAL_SDK')
                self.assertIsNone(delivery.ai_suggested_name)
                self.assertTrue(result.sdk_identity['public_outputs'])
                for call in result.sdk_identity['public_calls']:
                    raw = result.sdk_identity['public_output_json'][call['api']].encode('utf-8')
                    self.assertEqual(sha256(raw).hexdigest(), call['sha256'])
                    self.assertEqual(len(raw), call['size_bytes'])
                    self.assertEqual(json.loads(raw), result.sdk_identity['public_outputs'][call['api']])
                self.assertTrue(all(c['status'] == 'SUCCESS' for c in result.sdk_identity['public_calls']))
                self.assertEqual(sha256(path.read_bytes()).hexdigest(), before)
                if name.endswith('.pdf'):
                    self.assertEqual(result.sdk_identity['availability'], 'PARTIAL')
                    self.assertEqual(delivery.source_ref['locator_status'], 'PAGE_ONLY')
                    self.assertIsNone(delivery.observation.forwarder_id)
                    self.assertFalse(delivery.observation.context.complete())
                    self.assertTrue(any(i['code'] == 'SDK_LIMITATION' for i in result.issues))
