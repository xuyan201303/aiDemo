"""F2: current quote confirmation is independent of future learning safety.

All writes use the runner's disposable PostgreSQL and synthetic quotation
fixtures. AVAILABLE below is an in-memory document DTO only; immutable
document history is never updated to manufacture processing completion.
"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import os
import unittest
from uuid import uuid4

from openpyxl import load_workbook

from sanzo_core.postgres import PostgresRepository
from sanzo_demo.case_overview import summarize_case
from sanzo_demo.comparison import comparison_projection
from sanzo_demo.export import workbook_bytes
from sanzo_demo.review_state import review_facts
from sanzo_demo.row_confirmation import RowConfirmationCommand, RowConfirmationService, RowMember
from test_demo_condition_corrections import ConditionOverlayPostgresTests


DSN = os.environ.get("SANZO_TEST_DSN")


@unittest.skipUnless(DSN, "NOT_RUN: dedicated isolated SANZO_TEST_DSN required")
class F2CurrentReviewPostgresTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        PostgresRepository(DSN).migrate()

    def setUp(self):
        self.f = ConditionOverlayPostgresTests(
            "test_original_reading_and_source_survive_row_move_and_new_excel"
        )
        self.f.setUp()

    def confirm_row(self, item, fee_id):
        f = self.f
        snapshot = f.snapshot(item["case_id"])
        row = next(row for row in comparison_projection(
            snapshot["items"], snapshot["fees"], snapshot["documents"]
        )["rows"] if item["item_id"] in row["item_ids"])
        fee = next(fee for fee in snapshot["fees"] if fee["fee_id"] == fee_id)
        return RowConfirmationService(f.repo).confirm(f.actor, RowConfirmationCommand(
            item["case_id"], row["row_id"], snapshot["case"]["case_revision"],
            tuple(RowMember(member["item_id"], member["version"],
                            member["current_memory_version"])
                  for member in row["member_items"]),
            uuid4().hex, target_fee_id=fee_id, expected_fee_id=fee_id,
            expected_name_version=fee["name_version"],
        ))

    def test_confirmation_must_follow_condition_change_before_current_case_completes(self):
        f = self.f
        case_id = f.case("f2-order")
        item = f.line(case_id, "A", "SOURCE FEE", "1200")
        fee_id = f.classify(item, name="海上運賃")["fee_id"]
        f.conditions.correct(f.actor, f.command(item, value="40FT"))
        before = f.current(item)
        self.assertFalse(before["classification_confirmed_after_condition"])
        self.assertTrue(review_facts(before)["review_required"])
        self.assertIn("適用条件変更後の費目未確認", review_facts(before)["review_reasons"])

        saved = self.confirm_row(item, fee_id)
        self.assertFalse(saved["memory_reusable"])
        current = f.current(item)
        self.assertTrue(current["classification_confirmed_after_condition"])
        self.assertEqual(current["observation"]["context"]["equipment"], "20FT")
        self.assertEqual(current["effective_context"]["equipment"], "40FT")
        self.assertEqual(current["current_price"], "1200")
        self.assertEqual(current["currency"], "USD")
        facts = review_facts(current)
        self.assertTrue(facts["classification_confirmed"])
        self.assertFalse(facts["learning_eligible"])
        self.assertFalse(facts["review_required"])
        self.assertTrue(facts["learning_notices"])

        snapshot = deepcopy(f.snapshot(case_id))
        snapshot["documents"] = [
            {**document, "state": "AVAILABLE", "blocked_reason": None, "issues": []}
            for document in snapshot["documents"]
        ]
        self.assertEqual(summarize_case(snapshot)["status"], "確認済み")
        projected = comparison_projection(snapshot["items"], snapshot["fees"],
                                          snapshot["documents"])["rows"][0]
        self.assertEqual(projected["confirmation_status"], "✓ 確認済み")
        self.assertFalse(projected["cells"]["A"][0]["needs_review"])
        workbook = load_workbook(BytesIO(workbook_bytes(
            snapshot["case"], snapshot["items"], snapshot["documents"],
            fees=snapshot["fees"])), read_only=True)
        try:
            self.assertEqual(workbook["原価単価一覧表"]["A2"].value,
                             "保存済み案件スナップショット")
            self.assertEqual(workbook["未確認事項"].max_row, 2)
            self.assertIn("次回自動適用は未確認", workbook["確認情報"]["N2"].value)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
