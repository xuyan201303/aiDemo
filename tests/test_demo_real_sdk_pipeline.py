"""Actual fixed-wheel reads and disposable PostgreSQL; classification is offline.

These tests upload the approved original samples and run DocumentSdkAdapter.
Only the semantic classifier is a clearly named MECHANISM_TEST substitute.
They do not establish live OpenAI or browser demonstration success.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from sanzo_core.demo_store import DemoStore
from sanzo_core.domain import CorrectionCommand, Principal, SDKObservedLine, memory_key
from sanzo_core.postgres import PostgresRepository
from sanzo_core.service import ReviewService
from sanzo_demo.app import create_app
from sanzo_demo.document_sdk_adapter import DocumentSdkAdapter
from sanzo_demo.facts import structure_facts
from sanzo_demo.human_field_bindings import apply_human_bindings
from sanzo_demo.ingestion import IntegrationRuntime
from test_demo_ingestion_pipeline import MechanismExtractor
from test_demo_mock_isolation import mock_source

DSN = os.environ.get("SANZO_TEST_DSN")
SAMPLES = Path(__file__).resolve().parents[1] / "demo_materials/samples"
FIRST = "01_A社_初回見積.xlsx"
PDF = "02_B社_初回見積.pdf"
SECOND = "03_C社_次回見積.xlsx"
# Test operator choices from whole SDK source cells; never adapter parsing rules.
HUMAN_HEADER_CHOICES = {
    "company": "SANZO Demo Forwarder A（架空）",
    "transport_mode": "対象：海上輸送と付帯サービス。各行に記載した条件での単価です。",
}


@unittest.skipUnless(DSN, "NOT_RUN: isolated PostgreSQL SANZO_TEST_DSN required")
class RealSDKPipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        PostgresRepository(DSN).migrate()
        # Fail explicitly on a missing/wrong wheel or Python runtime; no Mock fallback.
        cls.reader = DocumentSdkAdapter()

    def setUp(self):
        self.folder = TemporaryDirectory(prefix="sanzo-real-sdk-pipeline-")
        self.addCleanup(self.folder.cleanup)
        self.storage = Path(self.folder.name)
        self.actor = Principal("real-sdk-test-" + uuid4().hex, "test-human-operator")
        self.repo = PostgresRepository(DSN)
        self.demo = DemoStore(self.repo)
        self.service = ReviewService(self.repo)
        self.classifier = MechanismExtractor({})
        self.runtime = IntegrationRuntime(self.repo, self.storage, reader=self.reader,
            extractor=self.classifier, reader_mode="REAL_SDK", allow_mechanism_tests=True)

    def new_case(self):
        case_id = str(uuid4())
        self.service.create_case(self.actor, case_id, "ACTUAL_SDK_OFFLINE_CLASSIFIER_TEST")
        return case_id

    def upload(self, filename, *, case_id=None, runtime=None):
        runtime = runtime or self.runtime
        case_id = case_id or self.new_case()
        raw = (SAMPLES / filename).read_bytes()
        with TestClient(create_app(DSN, self.storage, tenant_id=self.actor.tenant_id,
                                  integrations=runtime)) as client:
            response = client.post(f"/api/cases/{case_id}/documents",
                                   files={"file": (filename, raw)})
        self.assertEqual(response.status_code, 201, response.text)
        document = self.demo.get_document(self.actor, response.json()["document_id"])
        self.assertEqual(document["sha256"], sha256(raw).hexdigest())
        return document

    def human_bindings(self, document, *, runtime=None):
        review = (runtime or self.runtime).review_field_sources(self.actor, document["document_id"])
        self.assertEqual(review["origin"], "REAL_SDK")
        self.assertEqual(review["ai_requests"], 0)
        result = {}
        for role, complete_text in HUMAN_HEADER_CHOICES.items():
            candidates = [c for c in review["candidates"] if c["raw_text"] == complete_text]
            self.assertEqual(len(candidates), 1, (role, review))
            result[role] = candidates[0]["source_id"]
        return result

    def process(self, document, *, bindings=None, runtime=None):
        runtime = runtime or self.runtime
        task = runtime.submit(self.actor, document["document_id"], uuid4().hex,
                              field_bindings=bindings)
        return runtime.run(self.actor, task["task_id"])

    def items(self, document):
        return [i for i in self.demo.list_items(self.actor, document["case_id"])
                if i["observation"]["document_id"] == document["document_id"]]

    def confirm(self, item, name):
        current = next(i for i in self.demo.list_items(self.actor, item["case_id"])
                       if i["item_id"] == item["item_id"])
        return self.service.correct(self.actor, CorrectionCommand(
            item_id=current["item_id"], expected_case_revision=current["case_revision"],
            expected_item_version=current["version"],
            expected_memory_version=current["current_memory_version"],
            idempotency_key=uuid4().hex, new_fee_name=name))

    def slot_count(self, *, origin=None):
        with self.repo.transaction() as conn:
            query = "SELECT count(*) AS n FROM sanzo_core.memory_slots WHERE tenant_id=%s"
            values = (self.actor.tenant_id,)
            if origin:
                query += " AND origin=%s"; values += (origin,)
            return conn.execute(query, values).fetchone()["n"]

    def read_evidence(self, task):
        return next(e["metadata"]["read_evidence"] for e in task["events"]
                    if e["state"] == "EXTRACTING")

    def test_actual_three_files_target_reuses_only_real_confirmation_and_new_sdk_amount(self):
        first = self.upload(FIRST)
        task1 = self.process(first, bindings=self.human_bindings(first))
        self.assertEqual(task1["state"], "REVIEW_REQUIRED", task1)
        self.assertEqual(len(self.items(first)), 5)
        self.assertEqual(self.slot_count(), 0)
        second_company = self.upload(PDF, case_id=first["case_id"])
        pdf_task = self.process(second_company)
        self.assertEqual(pdf_task["state"], "REVIEW_REQUIRED", pdf_task)
        self.assertEqual(len(self.items(second_company)), 5)
        self.assertEqual(self.slot_count(), 0)
        self.assertEqual([c["purpose"] for c in self.classifier.calls],
                         ["fee_classification", "fee_classification"])
        for task in (task1, pdf_task):
            evidence = self.read_evidence(task)
            self.assertEqual(evidence["origin"], "REAL_SDK")
            self.assertEqual(evidence["sdk_identity"]["installed_file_check"], "VERIFIED")
            self.assertTrue(evidence["sdk_identity"]["public_outputs"])
            self.assertTrue(all(c["status"] == "SUCCESS" for c in evidence["sdk_identity"]["public_calls"]))
            self.assertEqual(task["result_json"]["ai_metadata"]["quotation_extraction_requests"], 0)
        delivery = next(i for i in self.items(first) if i["observation"]["raw_label"] == "配送手数料")
        self.assertEqual(delivery["current_price"], "8000")
        confirmed = self.confirm(delivery, "配送費")
        self.assertEqual(confirmed["memory_effect"], "ACTIVE")
        before = self.demo.case_snapshot(self.actor, first["case_id"])
        self.classifier.calls.clear()
        second = self.upload(SECOND)
        task2 = self.process(second, bindings=self.human_bindings(second))
        self.assertEqual(task2["state"], "REVIEW_REQUIRED", task2)
        target = next(i for i in self.items(second) if i["observation"]["raw_label"] == "配送手数料")
        self.assertEqual((target["current_fee_name"], target["current_price"], target["currency"]),
                         ("配送費", "9500", "JPY"))
        self.assertEqual(target["current_fee_id"], confirmed["fee_id"])
        self.assertEqual(target["classification_source"], "MEMORY")
        self.assertEqual(target["fact_sources"]["amount"], "REAL_SDK")
        amount = target["source_ref"]["field_sources"]["raw_value"]
        self.assertEqual(amount["original_value"], "9500")
        source = next(b for b in self.read_evidence(task2)["sources"]
                      if b["source_id"] == amount["source_block_id"])
        self.assertEqual(source["source_ref"]["sdk_value"], 9500)
        self.assertEqual(source["source_ref"]["sdk_physical_cell"]["raw_cached_value"], "9500")
        self.assertEqual(target["source_ref"]["document_sha256"], second["sha256"])
        for field in HUMAN_HEADER_CHOICES:
            self.assertEqual(target["source_ref"]["field_sources"][field]["evidence_unit"]["role_origin"], "HUMAN")
        self.assertEqual(len(self.classifier.calls), 1)
        self.assertEqual(len(self.classifier.calls[0]["blocks"]), 4)
        self.assertNotIn(target["item_id"], {i["item_id"] for i in self.classifier.calls[0]["blocks"]})
        self.assertTrue(all(set(i) == {"item_id", "raw_label", "context"}
                            for i in self.classifier.calls[0]["blocks"]))
        meta = task2["result_json"]["ai_metadata"]
        self.assertIn(target["item_id"], meta["classification_skipped_item_ids"])
        self.assertEqual(meta["memory_lookup_results"][target["item_id"]]["decision"], "MEMORY_APPLIED")
        self.assertEqual(meta["quotation_extraction_requests"], 0)
        self.assertEqual(self.demo.case_snapshot(self.actor, first["case_id"]), before)

    def test_all_five_real_confirmed_new_connection_http_process_needs_no_ai_or_key(self):
        first = self.upload(FIRST)
        task = self.process(first, bindings=self.human_bindings(first))
        self.assertEqual(task["state"], "REVIEW_REQUIRED", task)
        for item in self.items(first):
            label = item["observation"]["raw_label"]
            self.confirm(item, "配送費" if label == "配送手数料" else "HUMAN " + label)
        self.assertEqual(self.slot_count(origin="REAL_SDK"), 5)
        before = self.demo.case_snapshot(self.actor, first["case_id"])
        self.classifier.calls.clear()
        restarted = IntegrationRuntime(PostgresRepository(DSN), self.storage,
            reader=DocumentSdkAdapter(), extractor=None, reader_mode="REAL_SDK")
        second = self.upload(SECOND, runtime=restarted)
        bindings = self.human_bindings(second, runtime=restarted)
        with TestClient(create_app(DSN, self.storage, tenant_id=self.actor.tenant_id,
                                  integrations=restarted)) as client:
            response = client.post(f"/api/documents/{second['document_id']}/process",
                                   json={"idempotency_key": uuid4().hex, "field_bindings": bindings, "pause_ai": True})
            self.assertEqual(response.status_code, 202, response.text)
            task = client.get("/api/tasks/" + response.json()["task_id"]).json()
            exported = client.get(f"/api/cases/{second['case_id']}/export.xlsx")
            self.assertEqual(exported.status_code, 200)
            target = next(i for i in self.items(second) if i["observation"]["raw_label"] == "配送手数料")
            loc = target["source_ref"]["locator"]
            preview = client.get(f"/api/documents/{second['document_id']}/preview", params={
                "sheet": loc["sheet_name"], "cell": loc["A1_range"]})
            self.assertEqual(preview.status_code, 200)
            self.assertIn('data-cell="' + loc["A1_range"] + '"', preview.text)
            self.assertEqual(sha256(client.get(f"/api/documents/{second['document_id']}/original").content).hexdigest(), second["sha256"])
        # Original SDK binding limitations remain visible after human roles.
        self.assertIn(task["state"], {"AVAILABLE", "REVIEW_REQUIRED"}, task)
        self.assertTrue(all(i["classification_source"] == "MEMORY" for i in self.items(second)))
        self.assertEqual(self.classifier.calls, [])
        self.assertIsNone(restarted.extractor)
        meta = task["result_json"]["ai_metadata"]
        self.assertEqual((meta["quotation_extraction_requests"], meta["calls"],
                          meta["classification_requested_item_ids"]), (0, [], []))
        book = load_workbook(BytesIO(exported.content), read_only=True)
        try:
            values = list(book["明細と出典"].values)
            target_row = next(r for r in values[1:] if r[0] == target["item_id"])
            self.assertEqual(target_row[4:6], ("9500", "JPY"))
            main = list(book["原価単価一覧表"].values)
            self.assertTrue(any("配送費" in str(r[0]) and "9500 JPY" in str(r[1:]) for r in main))
        finally:
            book.close()
        self.assertEqual(self.demo.case_snapshot(self.actor, first["case_id"]), before)

    def test_paused_unknown_file_keeps_real_read_evidence_without_classification(self):
        document = self.upload(FIRST)
        bindings = self.human_bindings(document)
        with TestClient(create_app(DSN, self.storage, tenant_id=self.actor.tenant_id,
                                  integrations=self.runtime)) as client:
            response = client.post(f"/api/documents/{document['document_id']}/process",
                json={"idempotency_key": uuid4().hex, "field_bindings": bindings, "pause_ai": True})
            self.assertEqual(response.status_code, 202, response.text)
            task = client.get("/api/tasks/" + response.json()["task_id"]).json()
        self.assertEqual(task["state"], "BLOCKED", task)
        self.assertEqual(task["error_code"], "AI_PROCESSING_PAUSED")
        self.assertEqual(self.classifier.calls, [])
        self.assertEqual(self.slot_count(), 0)
        self.assertEqual(self.read_evidence(task)["origin"], "REAL_SDK")
        self.assertTrue(self.read_evidence(task)["sdk_identity"]["public_outputs"])
        self.assertEqual(self.items(document), [])

    def test_pdf_limited_evidence_allows_current_correction_without_fee_memory(self):
        document = self.upload(PDF)
        task = self.process(document)
        self.assertEqual(task["state"], "REVIEW_REQUIRED", task)
        items = self.items(document)
        self.assertEqual(len(items), 5)
        self.assertEqual(self.slot_count(), 0)
        evidence = self.read_evidence(task)
        self.assertTrue(evidence["reader_issues"])
        self.assertIn("OCR_NOT_EXECUTED", str(evidence["sdk_identity"]["limitations"]))
        self.assertIn("OCR_NOT_EXECUTED", str(task["result_json"]["issues"]))
        for item in items:
            self.assertIsNone(item["observation"]["context"]["transport_mode"])
            self.assertIsNone(item["current_fee_id"])
            self.assertEqual(item["source_ref"]["locator_status"], "PAGE_ONLY")
            self.assertEqual(set(item["source_ref"]["locator"]), {"page_number"})
        confirmed = self.confirm(items[0], "HUMAN FEE NAME")
        self.assertEqual(confirmed["memory_effect"], "NOT_ELIGIBLE_CONTEXT_OR_SOURCE")
        self.assertEqual(self.slot_count(), 0)
        with TestClient(create_app(DSN, self.storage, tenant_id=self.actor.tenant_id,
                                  integrations=self.runtime)) as client:
            preview = client.get(f"/api/documents/{document['document_id']}/preview", params={"page": 1})
            self.assertEqual(preview.status_code, 200)
            book = load_workbook(BytesIO(client.get(f"/api/cases/{document['case_id']}/export.xlsx").content), read_only=True)
        try:
            notes = "\n".join(str(c or "") for row in book["未確認事項"].values for c in row)
            self.assertIn("OCR_NOT_EXECUTED", notes)
            self.assertIn("EXTRACTION_UNRESOLVED", notes)
            self.assertIn("適用条件未確認", notes)
        finally:
            book.close()

    def test_invalid_human_source_binding_retains_completed_real_read_failure_evidence(self):
        document = self.upload(FIRST)
        task = self.process(document, bindings={"company": "UNLISTED_HUMAN_SOURCE_CANDIDATE"})
        self.assertEqual(task["state"], "FAILED", task)
        self.assertEqual(self.classifier.calls, [])
        self.assertEqual(self.items(document), [])
        self.assertEqual(self.slot_count(), 0)
        failure = task["events"][-1]
        self.assertEqual(failure["state"], "FAILED")
        metadata = failure["metadata"]
        self.assertTrue(metadata["read_completed"])
        self.assertTrue(metadata["sdk_read_completed"])
        self.assertFalse(metadata["mock_read_completed"])
        evidence = metadata["read_evidence"]
        self.assertEqual(evidence["origin"], "REAL_SDK")
        self.assertEqual(evidence["document_sha256"], document["sha256"])
        identity = evidence["sdk_identity"]
        self.assertEqual(identity["real_sdk_execution"], "EXECUTED")
        self.assertEqual(identity["installed_file_check"], "VERIFIED")
        self.assertTrue(identity["public_outputs"])
        self.assertGreater(len(identity["public_calls"]), 0)
        self.assertTrue(all(call["status"] == "SUCCESS" for call in identity["public_calls"]))
        self.assertEqual(len([block for block in evidence["sources"]
            if block["source_ref"]["evidence_unit"]["field_name"] == "raw_label"]), 5)
        # A fresh repository must still expose the original public SDK output;
        # binding validation failure cannot rewrite an actual read as NOT_RUN.
        restarted = IntegrationRuntime(PostgresRepository(DSN), self.storage,
            reader=DocumentSdkAdapter(), extractor=None, reader_mode="REAL_SDK")
        saved = restarted.store.get_task(self.actor, task["task_id"])
        self.assertEqual(saved["events"][-1]["metadata"]["read_evidence"], evidence)

    def test_existing_mock_memory_same_business_context_never_applies_to_real_read(self):
        incoming = self.upload(FIRST)
        reading = self.reader.read(self.storage / incoming["storage_key"], incoming)
        bound = apply_human_bindings(reading, incoming, self.actor, self.human_bindings(incoming))
        lines, _ = structure_facts(incoming, self.actor.tenant_id, bound)
        line = next(l for l in lines if l.observation.raw_label == "配送手数料")
        # Only this old-memory seed is synthetic. No Mock reader/JSON feeds the
        # real incoming task, which still reads uploaded bytes through the wheel.
        old_case = self.new_case(); old_document = str(uuid4()); old_source = uuid4().hex
        doc = self.demo.register_document(self.actor, old_case, document_id=old_document,
            original_filename="EXPLICIT_OLD_MOCK_MECHANISM_SEED.pdf", media_type="application/pdf",
            size_bytes=1, sha256="a" * 64, storage_key=old_document + ".pdf")
        obs = replace(line.observation, item_id=str(uuid4()), case_id=old_case,
                      document_id=old_document, source_id=old_source)
        legacy = SDKObservedLine(obs, mock_source(old_document, old_source),
                                 raw_value=line.raw_value, input_kind="MOCK_SDK")
        with self.repo.transaction() as conn:
            self.demo._insert_line(conn, self.actor, doc, legacy)
            self.repo.advance_case(conn, self.actor.tenant_id, old_case)
        old_item = self.demo.list_items(self.actor, old_case)[0]
        self.confirm(old_item, "MOCK ONLY NAME")
        self.assertEqual(self.slot_count(origin="MOCK_SDK"), 1)
        self.assertNotEqual(memory_key(legacy.observation, legacy.source_ref, legacy.input_kind),
                            memory_key(line.observation, line.source_ref, line.input_kind))
        task = self.process(incoming, bindings=self.human_bindings(incoming))
        self.assertEqual(task["state"], "REVIEW_REQUIRED", task)
        target = next(i for i in self.items(incoming) if i["observation"]["raw_label"] == "配送手数料")
        self.assertIsNone(target["current_fee_id"])
        self.assertEqual(target["classification_source"], "AI")
        self.assertIn(target["item_id"], {i["item_id"] for i in self.classifier.calls[0]["blocks"]})
        self.assertEqual(self.slot_count(origin="REAL_SDK"), 0)
        self.assertEqual(self.slot_count(origin="MOCK_SDK"), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
