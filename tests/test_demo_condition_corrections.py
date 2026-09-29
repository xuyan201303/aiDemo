"""Condition overlay probes use only synthetic lines and an isolated PostgreSQL DSN.

No SDK, OpenAI, uploaded customer file, or live Demo database is touched.
"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import os
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook

from sanzo_core.condition_corrections import ConditionCorrectionCommand, ConditionCorrectionService
from sanzo_core.demo_store import DemoStore
from sanzo_core.domain import (ConflictError, Context, CorrectionCommand, Observation,
                               Principal, RevokeCommand, SDKObservedLine, SyntheticLine)
from sanzo_core.postgres import PostgresRepository
from sanzo_core.service import ReviewService
from sanzo_demo.case_overview import summarize_case
from sanzo_demo.comparison import comparison_projection
from sanzo_demo.export import workbook_bytes
from sanzo_demo.review_state import review_facts
from sanzo_demo.row_confirmation import (RowConfirmationCommand, RowConfirmationService,
                                         RowMember, SingleMemberCorrectionCommand,
                                         SingleMemberCorrectionService)

DSN = os.environ.get("SANZO_TEST_DSN")


class ConditionCommandValidationTests(unittest.TestCase):
    def command(self, **changes):
        values = dict(case_id="case", row_id="a" * 64, item_id="item",
                      expected_case_revision=1, expected_item_version=1,
                      expected_memory_version=0, idempotency_key="one",
                      field="equipment", value="40FT")
        values.update(changes)
        return ConditionCorrectionCommand(**values)

    def test_only_bounded_equipment_or_billing_unit_values(self):
        self.command().validate()
        self.command(field="billing_unit", value="B/L").validate()
        for changes in ({"field": "transport_mode"}, {"field": "source_ref"},
                        {"value": " "}, {"value": "x" * 1025},
                        {"expected_item_version": True}, {"row_id": "bad"},
                        {"expected_memory_evidence": []}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                self.command(**changes).validate()


@unittest.skipUnless(DSN, "NOT_RUN: dedicated isolated SANZO_TEST_DSN required")
class ConditionOverlayPostgresTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        PostgresRepository(DSN).migrate()

    def setUp(self):
        self.tenant = "condition-injected-" + uuid4().hex
        self.actor = Principal(self.tenant, "human-condition-reviewer")
        self.repo = PostgresRepository(DSN)
        self.store = DemoStore(self.repo)
        self.review = ReviewService(self.repo)
        self.conditions = ConditionCorrectionService(self.repo)

    def case(self, suffix):
        case_id = "case-" + suffix + "-" + uuid4().hex
        self.review.create_case(self.actor, case_id, "INJECTED_TEST_CASE")
        return case_id

    def line(self, case_id, company, label, price, *, equipment="20FT",
             billing_unit="CONTAINER"):
        item_id = "item-" + uuid4().hex
        document_id = "document-" + uuid4().hex
        source_id = "source-" + uuid4().hex
        observation = Observation(self.tenant, item_id, case_id, document_id,
            source_id, company, label, price, "USD",
            Context("SEA", equipment, billing_unit, "ORIGIN", "STANDARD"), True)
        source = {"synthetic": True, "document_id": document_id,
                  "document_revision": 1, "document_sha256": "a" * 64,
                  "source_block_ids": [source_id], "kind": "EXCEL_RANGE",
                  "locator_status": "VERIFIED",
                  "locator": {"sheet_name": "INJECTED", "A1_range": "B14"}}
        self.store.register_document(self.actor, case_id, document_id=document_id,
            original_filename=document_id + ".xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=1, sha256="a" * 64, storage_key=document_id + ".xlsx")
        return self.store.add_synthetic_lines(self.actor, document_id,
            [SyntheticLine(observation, source)])[0]

    def classify(self, item, *, name=None, fee_id=None):
        current = self.current(item)
        return self.review.correct(self.actor, CorrectionCommand(
            item_id=current["item_id"], expected_case_revision=current["case_revision"],
            expected_item_version=current["version"],
            expected_memory_version=current["current_memory_version"],
            idempotency_key=uuid4().hex,
            new_fee_name=name, target_fee_id=fee_id))

    def current(self, item):
        return next(i for i in self.store.list_items(self.actor, item["case_id"])
                    if i["item_id"] == item["item_id"])

    def snapshot(self, case_id):
        return self.store.case_snapshot(self.actor, case_id)

    def command(self, item, *, field="equipment", value="40FT", key=None):
        snap = self.snapshot(item["case_id"])
        row = next(row for row in comparison_projection(snap["items"], snap["fees"], snap["documents"])["rows"]
                   if item["item_id"] in row["item_ids"])
        member = next(i for i in row["member_items"] if i["item_id"] == item["item_id"])
        resolution = member.get("memory_resolution") or {}
        return ConditionCorrectionCommand(
            case_id=item["case_id"], row_id=row["row_id"], item_id=item["item_id"],
            expected_case_revision=snap["case"]["case_revision"],
            expected_item_version=member["version"],
            expected_memory_version=member["current_memory_version"],
            expected_memory_evidence=(resolution.get("evidence_versions")
                if resolution.get("kind") == "LEGACY_PROJECTION" else None),
            idempotency_key=key or uuid4().hex, field=field, value=value)

    def assert_current_conflict_remains_pending(self, case_id, item, reason):
        current = self.current(item)
        facts = review_facts(current)
        self.assertIn(reason, facts["review_reasons"])
        self.assertTrue(facts["review_required"])
        snapshot = deepcopy(self.snapshot(case_id))
        # Isolate current review from ingestion: synthetic test documents do
        # not pass through a real SDK and cannot be marked AVAILABLE in SQL.
        snapshot["documents"] = [
            {**document, "state": "AVAILABLE", "blocked_reason": None, "issues": []}
            for document in snapshot["documents"]
        ]
        self.assertEqual(summarize_case(snapshot)["status"], "確認必要")
        self.assertTrue(any(row["item_id"] == item["item_id"] and
                            review_facts(row)["review_required"]
                            for row in snapshot["items"]))
        book = load_workbook(BytesIO(workbook_bytes(
            snapshot["case"], snapshot["items"], snapshot["documents"],
            fees=snapshot["fees"])), read_only=True)
        try:
            self.assertIn("DRAFT", book["原価単価一覧表"]["A2"].value)
            issues = list(book["未確認事項"].values)
            self.assertTrue(any(row[0] == item["item_id"] and reason in (row[2] or "")
                                for row in issues[2:]))
        finally:
            book.close()

    def test_prior_revocation_is_not_downgraded_by_later_condition_guard(self):
        case_id = self.case("guard-after-revoke")
        original = self.line(case_id, "A", "REVOKED SAME SOURCE", "100")
        self.classify(original, name="撤回対象費目")
        before = self.current(original)
        self.review.revoke_memory(self.actor, RevokeCommand(
            item_id=before["item_id"], expected_case_revision=before["case_revision"],
            expected_item_version=before["version"],
            expected_memory_version=before["current_memory_version"],
            idempotency_key=uuid4().hex, reason="人工確認を撤回"))
        self.assertIn("確認記録が撤回されています", review_facts(self.current(original))["review_reasons"])

        later = self.line(self.case("later-correction"), "A", "REVOKED SAME SOURCE", "200")
        self.conditions.correct(self.actor, self.command(later, value="40FT"))
        # The condition event must still protect future reuse, but it must not
        # erase the older explicit revocation from the current case/XLSX.
        self.assert_current_conflict_remains_pending(case_id, original,
                                                      "確認記録が撤回されています")
        # A later ordinary fee confirmation cannot erase a still-unresolved
        # revocation merely because the condition guard also applies.
        self.classify(original, fee_id=self.current(original)["current_fee_id"])
        self.assert_current_conflict_remains_pending(case_id, original,
                                                      "確認記録が撤回されています")

    def test_prior_shared_key_conflict_is_not_downgraded_by_later_condition_guard(self):
        case_id = self.case("guard-after-shared-conflict")
        original = self.line(case_id, "A", "SHARED SAME SOURCE", "100")
        fee_id = self.classify(original, name="元の費目")["fee_id"]
        other = self.line(self.case("other-fee"), "A", "UNRELATED FEE", "10")
        other_fee_id = self.classify(other, name="別の費目")["fee_id"]
        contested = self.line(self.case("same-key-different-fee"), "A", "SHARED SAME SOURCE", "200")
        snapshot = self.snapshot(contested["case_id"])
        row = next(row for row in comparison_projection(
            snapshot["items"], snapshot["fees"], snapshot["documents"])["rows"]
            if contested["item_id"] in row["item_ids"])
        member = next(member for member in row["member_items"]
                      if member["item_id"] == contested["item_id"])
        name_version = next(fee["name_version"] for fee in snapshot["fees"]
                            if fee["fee_id"] == other_fee_id)
        corrected = SingleMemberCorrectionService(self.repo).correct(self.actor,
            SingleMemberCorrectionCommand(
                case_id=contested["case_id"], row_id=row["row_id"],
                item_id=contested["item_id"],
                expected_case_revision=snapshot["case"]["case_revision"],
                expected_item_version=member["version"],
                expected_memory_version=member["current_memory_version"],
                idempotency_key=uuid4().hex, target_fee_id=other_fee_id,
                expected_fee_id=other_fee_id, expected_name_version=name_version))
        self.assertEqual(corrected["member"]["memory_action"], "CONFLICTED_SHARED_KEY")
        self.assertNotEqual(fee_id, other_fee_id)
        self.assertIn("確認記録が競合しています", review_facts(self.current(original))["review_reasons"])

        later = self.line(self.case("later-shared-condition"), "A", "SHARED SAME SOURCE", "300")
        self.conditions.correct(self.actor, self.command(later, value="40FT"))
        self.assert_current_conflict_remains_pending(case_id, original,
                                                      "確認記録が競合しています")

    def test_original_reading_and_source_survive_row_move_and_new_excel(self):
        case_id = self.case("move")
        twenty = self.line(case_id, "A", "OCEAN FREIGHT 20", "1200")
        fee_id = self.classify(twenty, name="海上運賃")["fee_id"]
        forty = self.line(case_id, "B", "OCEAN FREIGHT 40", "1900", equipment="40FT")
        self.classify(forty, fee_id=fee_id)
        wrong = self.line(case_id, "C", "WRONG READ", "1180")
        self.classify(wrong, fee_id=fee_id)
        before = deepcopy(self.current(wrong))
        other_before = [deepcopy(self.current(x)) for x in (twenty, forty)]
        old_excel = workbook_bytes(self.snapshot(case_id)["case"], self.snapshot(case_id)["items"],
                                   self.snapshot(case_id)["documents"], fees=self.snapshot(case_id)["fees"])
        result = self.conditions.correct(self.actor, self.command(wrong))
        self.assertEqual(result["original_value"], "20FT")
        self.assertEqual(result["effective_value"], "40FT")
        self.assertFalse(result["future_reuse_eligible"])
        after = self.current(wrong)
        self.assertEqual(after["observation"], before["observation"])
        self.assertEqual(after["source_ref"], before["source_ref"])
        self.assertEqual(after["current_price"], before["current_price"])
        self.assertEqual(after["currency"], before["currency"])
        self.assertEqual(after["current_fee_name"], before["current_fee_name"])
        self.assertEqual(after["effective_context"]["equipment"], "40FT")
        self.assertEqual(after["observation"]["context"]["equipment"], "20FT")
        for old, original in zip(other_before, (twenty, forty)):
            fresh = self.current(original)
            for field in ("observation", "source_ref", "current_fee_id", "current_fee_name",
                          "current_fee_name_version", "current_price", "currency", "version",
                          "effective_context", "condition_overrides", "current_memory_version"):
                self.assertEqual(fresh[field], old[field], field)
        snap = self.snapshot(case_id)
        rows = comparison_projection(snap["items"], snap["fees"], snap["documents"])["rows"]
        by_equipment = {row["equipment"]: row for row in rows}
        self.assertEqual(set(by_equipment), {"20FT", "40FT"})
        self.assertEqual(by_equipment["20FT"]["item_ids"], [twenty["item_id"]])
        self.assertEqual(set(by_equipment["40FT"]["item_ids"]), {forty["item_id"], wrong["item_id"]})
        new_excel = workbook_bytes(snap["case"], snap["items"], snap["documents"], fees=snap["fees"])
        book = load_workbook(BytesIO(new_excel), read_only=True)
        try:
            self.assertEqual({book["原価単価一覧表"].cell(i, 1).value for i in (4, 5)},
                             {row["display_name"] for row in rows})
            detail = list(book["明細と出典"].values)
            self.assertTrue(any(wrong["item_id"] in row and "40FT" in row for row in detail))
        finally:
            book.close()
        self.assertNotEqual(old_excel, new_excel)
        with self.repo.transaction() as conn:
            raw = conn.execute("""
                SELECT observation,source_ref,current_price_text,currency
                FROM sanzo_core.items WHERE tenant_id=%s AND item_id=%s
            """, (self.tenant, wrong["item_id"])).fetchone()
            self.assertEqual(raw["observation"]["context"]["equipment"], "20FT")
            self.assertEqual(raw["source_ref"], before["source_ref"])
            history = conn.execute("""
                SELECT actor_id,field_name,original_value,before_effective_value,
                       human_confirmed_value,item_version,case_revision
                FROM sanzo_core.condition_correction_events
                WHERE tenant_id=%s AND item_id=%s
            """, (self.tenant, wrong["item_id"])).fetchall()
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["actor_id"], self.actor.actor_id)
        self.assertEqual((history[0]["original_value"], history[0]["before_effective_value"],
                          history[0]["human_confirmed_value"]), ("20FT", "20FT", "40FT"))

    def test_original_memory_key_fails_closed_and_other_key_stays_active(self):
        case_id = self.case("memory")
        wrong_seed = self.line(case_id, "C", "O/F 20", "100")
        wrong_fee = self.classify(wrong_seed, name="海上運賃")["fee_id"]
        good_seed = self.line(case_id, "A", "OCEAN 20", "120")
        self.classify(good_seed, fee_id=wrong_fee)
        wrong_next = self.line(case_id, "C", "O/F 20", "105")
        self.assertEqual(wrong_next["saved_status"], "MEMORY_APPLIED")
        result = self.conditions.correct(self.actor, self.command(wrong_next))
        self.assertEqual(result["memory_action"], "ORIGINAL_SIGNATURE_CONFLICTED")
        self.assertEqual(self.current(wrong_next)["current_memory_status"], "CONFLICTED")
        self.assertFalse(self.current(wrong_next)["future_reuse_eligible"])
        # Even an ordinary later confirmation cannot certify that the reader's
        # original 20FT condition is correct. The response, stored slot and
        # effective read must now all agree that reuse remains blocked.
        current_seed = self.current(wrong_seed)
        reactivated = self.review.correct(self.actor, CorrectionCommand(
            item_id=current_seed["item_id"],
            expected_case_revision=current_seed["case_revision"],
            expected_item_version=current_seed["version"],
            expected_memory_version=current_seed["current_memory_version"],
            idempotency_key=uuid4().hex, target_fee_id=wrong_fee,
            replace_existing_mapping=True))
        self.assertEqual(reactivated["memory_effect"], "CONFLICTED")
        self.assertEqual(reactivated["memory_action"],
                         "CONFLICTED_CONDITION_CORRECTION_GUARD")
        with self.repo.transaction() as conn:
            from sanzo_core.domain import canonical_json, memory_key
            observation = Observation.from_json(canonical_json(current_seed["observation"]))
            key = memory_key(observation, current_seed["source_ref"], current_seed["input_kind"])
            self.assertEqual(self.repo.slot(conn, self.tenant, key)["status"], "CONFLICTED")
            self.assertEqual(self.repo.lookup_memory(conn, observation,
                current_seed["source_ref"], current_seed["input_kind"])["status"], "CONFLICTED")
        after_case = self.case("next")
        wrong_later = self.line(after_case, "C", "O/F 20", "205")
        good_later = self.line(after_case, "A", "OCEAN 20", "220")
        self.assertIsNone(wrong_later["current_fee_id"])
        self.assertEqual(wrong_later["current_memory_status"], "CONFLICTED")
        self.assertEqual(wrong_later["current_price"], "205")
        self.assertEqual(good_later["current_fee_id"], wrong_fee)
        self.assertEqual(good_later["saved_status"], "MEMORY_APPLIED")
        self.assertEqual(good_later["current_price"], "220")

    def test_guarded_original_key_without_slot_cannot_be_published_as_active(self):
        seed = self.line(self.case("unlearned-guard"), "C", "SAME MISREAD FEE", "100")
        self.assertIsNone(self.current(seed)["current_memory_status"])
        self.conditions.correct(self.actor, self.command(seed, value="40FT"))
        later = self.line(self.case("first-confirmation"), "C", "SAME MISREAD FEE", "200")
        snap = self.snapshot(later["case_id"])
        row = comparison_projection(snap["items"], snap["fees"], snap["documents"])["rows"][0]
        member = row["member_items"][0]
        command = RowConfirmationCommand(
            case_id=later["case_id"], row_id=row["row_id"],
            expected_case_revision=snap["case"]["case_revision"],
            members=(RowMember(member["item_id"], member["version"],
                               member["current_memory_version"]),),
            idempotency_key=uuid4().hex, new_fee_name="確認した費目")
        result = RowConfirmationService(self.repo).confirm(self.actor, command)
        self.assertFalse(result["memory_reusable"])
        self.assertEqual(result["members"][0]["memory_effect"], "CONFLICTED")
        self.assertEqual(result["members"][0]["memory_action"],
                         "CONFLICTED_CONDITION_CORRECTION_GUARD")
        self.assertEqual(self.current(later)["current_fee_id"], result["fee_id"])
        self.assertEqual(self.current(later)["current_memory_status"], "CONFLICTED")
        with self.repo.transaction() as conn:
            from sanzo_core.domain import canonical_json, memory_key
            observation = Observation.from_json(canonical_json(later["observation"]))
            key = memory_key(observation, later["source_ref"], later["input_kind"])
            self.assertEqual(self.repo.slot(conn, self.tenant, key)["status"], "CONFLICTED")
            self.assertTrue(self.repo.condition_correction_guarded(conn, self.tenant, key))
        future = self.line(self.case("future-still-guarded"), "C", "SAME MISREAD FEE", "300")
        self.assertEqual((future["saved_status"], future["current_fee_id"],
                          future["current_price"]), ("REVIEW_REQUIRED", None, "300"))
        # A different, unguarded signature still learns normally.
        ordinary = self.line(self.case("ordinary-rule"), "A", "SAFE OTHER FEE", "410")
        confirmed = self.classify(ordinary, name="別費目")
        self.assertEqual(confirmed["memory_effect"], "ACTIVE")
        reused = self.line(self.case("ordinary-next"), "A", "SAFE OTHER FEE", "420")
        self.assertEqual((reused["saved_status"], reused["current_fee_id"],
                          reused["current_price"]), ("MEMORY_APPLIED", confirmed["fee_id"], "420"))

    def test_guarded_key_blocks_injected_legacy_projection_and_new_active_write(self):
        # This is an explicitly injected old-policy fixture in disposable PG.
        # No SDK reader or customer file is used to construct these source facts.
        label = "INJECTED LEGACY FEE"

        def injected_line(case_id, price):
            item_id, document_id = "item-" + uuid4().hex, "document-" + uuid4().hex
            record_id = "INJECTED_RECORD_" + uuid4().hex
            source_id = "INJECTED_LABEL_" + uuid4().hex
            obs = Observation(self.tenant, item_id, case_id, document_id, source_id,
                "A", label, price, "USD",
                Context("SEA", "20FT", "CONTAINER", "ORIGIN", "STANDARD"), True)
            unit = {"kind": "EXCEL_CELL", "field_name": "raw_label",
                    "record_id": record_id, "scope": "LINE"}
            source = {"origin": "SDK_OBSERVED", "extraction_origin": "REAL_SDK",
                "document_id": document_id, "document_sha256": "a" * 64,
                "source_block_ids": [source_id], "kind": "EXCEL_RANGE",
                "locator_status": "VERIFIED",
                "locator": {"sheet_name": "INJECTED_TEST_CASE", "A1_range": "B14"},
                "text_verification": "VERIFIED", "evidence_unit": unit,
                "field_sources": {"raw_label": {"value": label, "original_value": label,
                    "source_block_id": source_id,
                    "record_key": ["EXCEL_RANGE", "INJECTED_TEST_CASE", record_id],
                    "evidence_unit": unit}}}
            for name, value in (("equipment", "20FT"), ("billing_unit", "CONTAINER"),
                                ("charge_scope", "ORIGIN"), ("extra_terms", "STANDARD")):
                field_id = "INJECTED_" + name + "_" + uuid4().hex
                source["source_block_ids"].append(field_id)
                source["field_sources"][name] = {"value": value, "original_value": value,
                    "source_block_id": field_id,
                    "record_key": ["EXCEL_RANGE", "INJECTED_TEST_CASE", record_id],
                    "evidence_unit": {"kind": "EXCEL_CELL", "field_name": name,
                        "record_id": record_id, "scope": "LINE"}}
            self.store.register_document(self.actor, case_id, document_id=document_id,
                original_filename="INJECTED_TEST_CASE.xlsx",
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                size_bytes=1, sha256="a" * 64, storage_key=document_id + ".xlsx")
            line = SDKObservedLine(obs, source, raw_value=price)
            with self.repo.transaction() as conn:
                document = self.store._document(conn, self.actor, document_id)
                self.repo.case(conn, self.tenant, case_id, lock=True)
                self.store._insert_line(conn, self.actor, document, line)
                self.repo.advance_case(conn, self.tenant, case_id)
                return self.store._item(conn, self.actor, item_id)

        misread = injected_line(self.case("legacy-guard"), "100")
        self.conditions.correct(self.actor, self.command(misread, value="40FT"))
        old_key = lambda observation, source=None, input_kind=None: observation.memory_key()
        with (patch("sanzo_core.fee_policy.uses_fee_policy", return_value=False),
              patch("sanzo_core.postgres.memory_key", old_key),
              patch("sanzo_core.service.memory_key", old_key),
              patch("sanzo_core.demo_store.memory_key", old_key)):
            historical = injected_line(self.case("legacy-history"), "200")
            legacy = self.classify(historical, name="旧方式の確認名")
        self.assertEqual(legacy["memory_effect"], "ACTIVE")
        next_item = injected_line(self.case("legacy-next"), "300")
        with self.repo.transaction() as conn:
            from sanzo_core.domain import canonical_json, memory_key
            observation = Observation.from_json(canonical_json(next_item["observation"]))
            key = memory_key(observation, next_item["source_ref"], next_item["input_kind"])
            self.assertTrue(self.repo.condition_correction_guarded(conn, self.tenant, key))
            self.assertIsNone(self.repo.slot(conn, self.tenant, key))
            self.assertIsNone(self.repo.lookup_memory(conn, observation,
                next_item["source_ref"], next_item["input_kind"]))
            with patch.object(self.repo, "condition_correction_guarded", return_value=False):
                projection = self.repo.lookup_memory(conn, observation,
                    next_item["source_ref"], next_item["input_kind"])
            self.assertEqual((projection["resolution"], projection["status"]),
                             ("LEGACY_PROJECTION", "ACTIVE"))
            self.assertEqual(self.repo.slot(conn, self.tenant, old_key(observation))["status"], "ACTIVE")
        result = self.classify(next_item, name="今回の確認名")
        self.assertEqual((result["memory_effect"], result["memory_action"]),
                         ("CONFLICTED", "CONFLICTED_CONDITION_CORRECTION_GUARD"))
        self.assertEqual(self.current(next_item)["current_memory_status"], "CONFLICTED")
        future = injected_line(self.case("legacy-future"), "400")
        self.assertEqual((future["saved_status"], future["current_fee_id"],
                          future["current_price"]), ("REVIEW_REQUIRED", None, "400"))

    def test_billing_unit_replay_stale_revision_and_rollback(self):
        case_id = self.case("billing")
        item = self.line(case_id, "A", "HANDLING", "300", equipment="20FT",
                         billing_unit="CONTAINER")
        fee_id = self.classify(item, name="取扱料")["fee_id"]
        peer = self.line(case_id, "B", "HANDLING", "350", equipment="20FT",
                         billing_unit="CONTAINER")
        self.classify(peer, fee_id=fee_id)
        command = self.command(item, field="billing_unit", value="B/L")
        first = self.conditions.correct(self.actor, command)
        self.assertEqual(first, self.conditions.correct(self.actor, command))
        self.assertEqual(self.current(item)["effective_context"]["billing_unit"], "B/L")
        snap = self.snapshot(case_id)
        rows = comparison_projection(snap["items"], snap["fees"], snap["documents"])["rows"]
        self.assertEqual({row["distinguishing_conditions"].get("billing_unit") for row in rows},
                         {"CONTAINER", "B/L"})
        self.assertEqual(next(row for row in rows if item["item_id"] in row["item_ids"])
                         ["distinguishing_conditions"]["billing_unit"], "B/L")
        with self.assertRaises(ConflictError):
            self.conditions.correct(self.actor, self.command(item, field="equipment", value="40FT",
                                                      key=command.idempotency_key))
        with self.assertRaises(ConflictError):
            self.conditions.correct(self.actor, command.__class__(**{
                **command.__dict__, "idempotency_key": uuid4().hex}))
        self.assertEqual(self.current(item)["current_fee_id"], fee_id)
        before = deepcopy(self.current(item))
        fault = ConditionCorrectionService(self.repo,
            fault_hook=lambda stage: (_ for _ in ()).throw(RuntimeError("INJECTED_FAILURE"))
            if stage == "after_condition_overlay" else None)
        with self.assertRaisesRegex(RuntimeError, "INJECTED_FAILURE"):
            fault.correct(self.actor, self.command(item, field="equipment", value="40FT"))
        self.assertEqual(self.current(item), before)
        with self.repo.transaction() as conn:
            count = conn.execute("""
                SELECT count(*) AS n FROM sanzo_core.condition_correction_events
                WHERE tenant_id=%s AND item_id=%s
            """, (self.tenant, item["item_id"])).fetchone()["n"]
        self.assertEqual(count, 1)

    def test_explicit_rename_after_condition_correction_uses_stable_fee_identity(self):
        case_id = self.case("rename-after-condition")
        wrong = self.line(case_id, "A", "O/F 20", "100")
        fee_id = self.classify(wrong, name="海上運賃")["fee_id"]
        forty = self.line(case_id, "B", "O/F 40", "200", equipment="40FT")
        self.classify(forty, fee_id=fee_id)
        true_twenty = self.line(case_id, "C", "OCEAN 20", "110")
        self.classify(true_twenty, fee_id=fee_id)
        self.conditions.correct(self.actor, self.command(wrong))
        before = self.snapshot(case_id)
        old_adopted = {i["item_id"]: (i["current_fee_name"],
            i["current_fee_name_version"], i["observation"], i["source_ref"],
            i["current_price"], i["currency"]) for i in before["items"]}
        old_excel = workbook_bytes(before["case"], before["items"], before["documents"], fees=before["fees"])
        old_book = load_workbook(BytesIO(old_excel), read_only=True)
        try:
            old_labels = [old_book["原価単価一覧表"].cell(row, 1).value for row in (4, 5)]
        finally:
            old_book.close()

        self.assertEqual(set(old_labels), {"海上運賃（20FT）", "海上運賃（40FT）"})
        projection = comparison_projection(before["items"], before["fees"], before["documents"])
        forty_row = next(row for row in projection["rows"] if row["equipment"] == "40FT")
        self.assertTrue(forty_row["rename_allowed"])
        members = tuple(RowMember(item["item_id"], item["version"],
            item["current_memory_version"]) for item in forty_row["member_items"])
        renamed = RowConfirmationService(self.repo).confirm(self.actor, RowConfirmationCommand(
            case_id=case_id, row_id=forty_row["row_id"],
            expected_case_revision=before["case"]["case_revision"],
            members=members, idempotency_key=uuid4().hex,
            intent="RENAME_FEE", expected_fee_id=fee_id,
            expected_name_version=forty_row["rename_name_version"],
            new_fee_name="海上輸送費"))
        self.assertEqual(renamed["fee_name"], "海上輸送費")
        self.assertFalse(renamed["memory_reusable"])
        after = self.snapshot(case_id)
        labels = {row["display_name"] for row in comparison_projection(
            after["items"], after["fees"], after["documents"])["rows"]}
        self.assertEqual(labels, {"海上輸送費（20FT）", "海上輸送費（40FT）"})
        for item in after["items"]:
            self.assertEqual((item["current_fee_name"], item["current_fee_name_version"],
                item["observation"], item["source_ref"], item["current_price"],
                item["currency"]), old_adopted[item["item_id"]])
        self.assertEqual(next(i for i in after["items"] if i["item_id"] == wrong["item_id"])
                         ["current_memory_status"], "CONFLICTED")
        old_book = load_workbook(BytesIO(old_excel), read_only=True)
        try:
            self.assertEqual([old_book["原価単価一覧表"].cell(row, 1).value for row in (4, 5)], old_labels)
        finally:
            old_book.close()

    def test_r2_guard_then_new_shared_key_conflict_remains_current_blocker(self):
        original = self.line(self.case("r2-original"), "A", "SAME SOURCE", "100")
        fee_a = self.classify(original, name="費用甲")["fee_id"]
        self.conditions.correct(self.actor, self.command(original, value="40FT"))
        other = self.line(self.case("r2-other"), "A", "OTHER SOURCE", "10")
        fee_b = self.classify(other, name="費用乙")["fee_id"]
        second = self.line(self.case("r2-second"), "A", "SAME SOURCE", "200")
        self.classify(second, fee_id=fee_a)
        third = self.line(self.case("r2-third"), "A", "SAME SOURCE", "300")
        self.classify(third, fee_id=fee_a)
        snap = self.snapshot(third["case_id"])
        row = next(row for row in comparison_projection(snap["items"], snap["fees"], snap["documents"])["rows"]
                   if third["item_id"] in row["item_ids"])
        version = next(fee["name_version"] for fee in snap["fees"] if fee["fee_id"] == fee_b)
        result = SingleMemberCorrectionService(self.repo).correct(self.actor,
            SingleMemberCorrectionCommand(third["case_id"], row["row_id"], third["item_id"],
                snap["case"]["case_revision"], self.current(third)["version"],
                self.current(third)["current_memory_version"], uuid4().hex,
                target_fee_id=fee_b, expected_fee_id=fee_b, expected_name_version=version))
        self.assertEqual(result["member"]["memory_action"], "CONFLICTED_SHARED_KEY")
        self.assert_current_conflict_remains_pending(second["case_id"], second,
                                                      "確認記録が競合しています")
        self.assert_current_conflict_remains_pending(third["case_id"], third,
                                                      "確認記録が競合しています")
        conflict = self.snapshot(third["case_id"])
        conflict_row = comparison_projection(conflict["items"], conflict["fees"],
                                             conflict["documents"])["rows"][0]
        self.assertEqual(conflict_row["confirmation_status"], "⚠ 確認必要")
        self.assertEqual(next(iter(conflict_row["cells"].values()))[0]["status"],
                         "⚠ 確認必要")
        future = self.line(self.case("r2-future"), "A", "SAME SOURCE", "400")
        self.assertIsNone(self.current(future)["current_fee_id"])

    def test_r2_condition_change_does_not_leave_row_confirmed(self):
        item = self.line(self.case("r2-row"), "A", "OCEAN R2", "100")
        fee_id = self.classify(item, name="海上運賃")["fee_id"]
        before = self.snapshot(item["case_id"])
        before_row = comparison_projection(before["items"], before["fees"], before["documents"])["rows"][0]
        self.assertEqual(before_row["confirmation_status"], "✓ 確認済み")
        self.conditions.correct(self.actor, self.command(item, value="40FT"))
        pending = self.snapshot(item["case_id"])
        pending_row = comparison_projection(pending["items"], pending["fees"], pending["documents"])["rows"][0]
        self.assertTrue(pending_row["member_items"][0]["review_facts"]["review_required"])
        self.assertEqual(pending_row["confirmation_status"], "⚠ 確認必要")
        self.assertEqual(next(iter(pending_row["cells"].values()))[0]["status"], "⚠ 確認必要")
        pending_book = load_workbook(BytesIO(workbook_bytes(
            pending["case"], pending["items"], pending["documents"], fees=pending["fees"])),
            read_only=True)
        try:
            audit = pending_book["確認情報"]
            self.assertEqual((audit[2][1].value, audit[2][6].value),
                             ("⚠ 確認必要", "未確認"))
        finally:
            pending_book.close()
        self.classify(item, fee_id=fee_id)
        done = self.snapshot(item["case_id"])
        done_row = comparison_projection(done["items"], done["fees"], done["documents"])["rows"][0]
        self.assertEqual(done_row["confirmation_status"], "✓ 確認済み")
        done_book = load_workbook(BytesIO(workbook_bytes(
            done["case"], done["items"], done["documents"], fees=done["fees"])),
            read_only=True)
        try:
            audit = done_book["確認情報"]
            self.assertEqual((audit[2][1].value, audit[2][6].value),
                             ("✓ 確認済み", "確認済み"))
        finally:
            done_book.close()

    def test_r2_condition_guard_then_revocation_remains_current_blocker(self):
        item = self.line(self.case("r2-guard-revoke"), "A", "REVOKE AFTER GUARD", "100")
        self.classify(item, name="撤回対象")
        self.conditions.correct(self.actor, self.command(item, value="40FT"))
        current = self.current(item)
        result = self.review.revoke_memory(self.actor, RevokeCommand(
            item_id=current["item_id"], expected_case_revision=current["case_revision"],
            expected_item_version=current["version"],
            expected_memory_version=current["current_memory_version"],
            idempotency_key=uuid4().hex, reason="人工確認を撤回"))
        self.assertEqual(result["memory_effect"], "REVOKED")
        self.assert_current_conflict_remains_pending(item["case_id"], item,
                                                      "確認記録が撤回されています")
        future = self.line(self.case("r2-after-revoke"), "A", "REVOKE AFTER GUARD", "200")
        self.assertIsNone(self.current(future)["current_fee_id"])


if __name__ == "__main__":
    unittest.main()
