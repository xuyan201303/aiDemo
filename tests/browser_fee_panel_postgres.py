"""Real-browser, real-row API and real PostgreSQL saves in the disposable test DB.
Only test-server localhost port 18765; never production/Demo cases or real AI.
"""
import argparse,json,traceback
from pathlib import Path
from io import BytesIO
from hashlib import sha256
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
BASE='http://127.0.0.1:18765'

def run(out):
    out.mkdir(parents=True,exist_ok=False);result={'scope':'DEDICATED_POSTGRES_SYNTHETIC_BROWSER_SAVES','checks':[],'requests':[],'ai_calls':0}
    def dump(n,d):(out/(n+'.json')).write_text(json.dumps(d,ensure_ascii=False,indent=2))
    def check(n,v=True):
        assert v,n
        result['checks'].append(n);print('PASS '+n,flush=True)
    try:
        with sync_playwright() as pw:
            browser=pw.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
            page=browser.new_page(viewport={'width':1366,'height':768});errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            def get(path):
                r=page.request.get(BASE+path);assert r.ok;return r.json()
            def open_case(cid):
                page.goto(BASE,wait_until='networkidle');page.locator('#case-select').select_option(cid)
                page.locator('.row-confirm-button').first.click();expect(page.locator('#edit-dialog')).to_be_visible()
            def guard(r):
                assert r.request.url.startswith(BASE+'/')
                if r.request.method=='POST':
                    assert '/comparison/' in r.request.url and r.request.url.endswith('/confirmation')
                    result['requests'].append(r.request.post_data_json)
                r.continue_()
            page.route('**/*',guard)
            audit=get('/__test/audit');dump('audit_before',audit)
            check('no memory before confirmation',not audit['memory_versions'])
            before=get('/api/cases/first');dump('first_before',before);open_case('first')
            page.locator('#edit-fee-name').fill('配送費');page.locator('.member-preview').first.click()
            expect(page.frame_locator('#preview-frame').locator('.source-cell')).to_have_count(3)
            page.screenshot(path=str(out/'browser_test_original.png'));page.locator('#preview-dialog .close-dialog').click();expect(page.locator('#edit-fee-name')).to_have_value('配送費')
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/confirmation')) as pending:
                page.locator('#edit-form').evaluate('(f)=>{f.requestSubmit();f.requestSubmit();}')
            response=pending.value;check('whole row HTTP200',response.status==200);saved=response.json();dump('save_result',saved)
            expect(page.locator('#edit-dialog')).not_to_be_visible();expect(page.locator('.row-status')).to_have_text('✓ 確認済み');expect(page.locator('#notice')).to_contain_text('名称は保存されましたが、一部の学習結果に確認が必要')
            check('double submit one HTTP request',len(result['requests'])==1)
            check('no prices in save payload',not {'unit_price','currency'} & result['requests'][0].keys() and len(result['requests'][0]['members'])==2)
            after=get('/api/cases/first');dump('first_after',after)
            for old,new in zip(before['items'],after['items']):
                check('all immutable facts '+old['item_id'],all(old[k]==new[k] for k in ['observation','raw_value','source_ref','current_price','currency','ai_suggested_name']))
                check('every displayed member confirmed '+old['item_id'],new['current_fee_id']==saved['fee_id'] and new['saved_status']=='HUMAN_CONFIRMED')
            audit=get('/__test/audit');dump('audit_after_first',audit)
            check('eligible A only memory',len(audit['memory_versions'])==1 and audit['memory_versions'][0]['raw_label']=='配送原文A')
            check('both member history',len(audit['correction_events'])==2)
            page.locator('#comparison-table').screenshot(path=str(out/'browser_after_save.png'))
            # Ordinary main-dialog rename must supersede the existing A mapping.
            # B remains ineligible and the mixed result must not claim full reuse.
            first_saved=saved;first_after=after;first_history=audit
            open_case('first');expect(page.locator('#edit-details')).not_to_have_attribute('open','')
            expect(page.locator('#replace-mapping')).not_to_be_checked()
            page.locator('#edit-fee-name').fill('配送費_R1_TEST')
            page.screenshot(path=str(out/'browser_ordinary_rename.png'))
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/confirmation')) as pending:
                page.locator('#save-correction').click()
            check('ordinary rename HTTP200',pending.value.status==200);saved=pending.value.json();dump('ordinary_rename_result',saved)
            expect(page.locator('#edit-dialog')).not_to_be_visible()
            expect(page.locator('#notice')).to_contain_text('名称は保存されましたが、一部の学習結果に確認が必要')
            check('ordinary rename no advanced replacement',result['requests'][-1]['replace_existing_mapping'] is False)
            check('A superseded ACTIVE B remains ineligible',[m['memory_effect'] for m in saved['members']]==['ACTIVE','NOT_ELIGIBLE_CONTEXT_OR_SOURCE'] and saved['members'][0]['memory_action']=='ACTIVE_SUPERSEDED' and saved['memory_reusable'] is False)
            after=get('/api/cases/first');dump('first_after_ordinary_rename',after)
            check('ordinary rename one shared new fee',saved['fee_id']!=first_saved['fee_id'] and all(i['current_fee_id']==saved['fee_id'] for i in after['items']))
            check('ordinary rename keeps all original facts',all(all(old[k]==new[k] for k in ['observation','source_ref','raw_value','current_price','currency']) for old,new in zip(first_after['items'],after['items'])))
            audit=get('/__test/audit');dump('audit_after_ordinary_rename',audit)
            check('old memory version retained and current ACTIVE v2',len(audit['memory_versions'])==2 and first_history['memory_versions'][0] in audit['memory_versions'] and len(audit['memory_slots'])==1 and audit['memory_slots'][0]['status']=='ACTIVE' and audit['memory_slots'][0]['version']==2 and audit['memory_slots'][0]['fee_id']==saved['fee_id'])
            check('ineligible B still no memory',all(v['raw_label']=='配送原文A' for v in audit['memory_versions']))
            page.screenshot(path=str(out/'browser_ordinary_rename_saved.png'),full_page=True)
            r=page.request.get(BASE+'/api/cases/first/export.xlsx');check('export HTTP200',r.ok)
            book=load_workbook(BytesIO(r.body()),read_only=True);sheet=book['原価単価一覧表'];check('Excel name price exact',[sheet.cell(4,c).value for c in [1,2,3]]==['配送費_R1_TEST（TRUCK）','8000 JPY','9000 JPY']);dump('export',{'sha256':sha256(r.body()).hexdigest(),'row':[sheet.cell(4,c).value for c in [1,2,3]],'sheets':book.sheetnames});book.close()
            check('create next after human memory',page.request.post(BASE+'/__test/next').ok)
            new=get('/api/cases/next');dump('next',new);i=new['items'][0]
            check('next actual memory fee and new9500',i['current_fee_id']==saved['fee_id'] and i['current_price']=='9500' and i['saved_status']=='MEMORY_APPLIED')
            old=get('/api/cases/first');check('old8000 unchanged after new case',old['items'][0]['current_price']=='8000' and old['items']==after['items'])
            check('next uses new human name without additional manual edit',i['current_fee_name']=='配送費_R1_TEST')
            page.goto(BASE,wait_until='networkidle');page.locator('#case-select').select_option('next')
            expect(page.locator('.row-confirm-button')).to_have_count(1)
            expect(page.locator('.price-cell')).to_have_count(1)
            expect(page.locator('.price-cell')).to_contain_text('9,500')
            page.screenshot(path=str(out/'browser_next_memory_reuse.png'),full_page=True)
            page.locator('#comparison-table').screenshot(path=str(out/'browser_next_memory_table.png'))
            # The existing one-member next case is fully eligible. Reconfirmation
            # reuses the same name and must give the positive next-file notice.
            page.locator('.row-confirm-button').click();expect(page.locator('#edit-fee-name')).to_have_value('配送費_R1_TEST')
            expect(page.locator('#replace-mapping')).not_to_be_checked()
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/confirmation')) as pending:
                page.locator('#save-correction').click()
            check('all eligible confirmation HTTP200',pending.value.status==200);all_saved=pending.value.json();dump('all_eligible_save_result',all_saved)
            check('all eligible truthful reusable result',all_saved['memory_reusable'] and all(m['memory_effect']=='ACTIVE' for m in all_saved['members']))
            expect(page.locator('#edit-dialog')).not_to_be_visible()
            expect(page.locator('#notice')).to_contain_text('同じ原文・条件の次回以降の読み取りに反映されます。')
            check('all eligible save preserves current9500',get('/api/cases/next')['items'][0]['current_price']=='9500')
            page.screenshot(path=str(out/'browser_all_eligible_saved.png'),full_page=True)
            # Inject a genuine failure on the second member inside the existing transaction.
            open_case('rollback');baseline=get('/__test/audit');check('arm bounded transaction fault',page.request.post(BASE+'/__test/fault').ok)
            page.locator('#edit-fee-name').fill('Rollback confirmation')
            with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/confirmation')) as pending:page.locator('#save-correction').click()
            check('injected second member returned500',pending.value.status==500);expect(page.locator('#edit-error')).to_contain_text('保存結果を確認できません')
            check('transaction failure changes no table',get('/__test/audit')==baseline)
            expect(page.locator('#edit-dialog')).to_be_visible();expect(page.locator('#edit-fee-name')).to_have_value('Rollback confirmation')
            page.screenshot(path=str(out/'browser_rollback.png'))
            failed=result['requests'][-1].copy()
            page.locator('#save-correction').click();expect(page.locator('#edit-dialog')).not_to_be_visible()
            check('retry same exact rollback request',result['requests'][-1]==failed)
            # Lose a successful HTTP response. The actual DB commits once, then replay returns it.
            open_case('lost');page.locator('#edit-fee-name').fill('Response confirmed')
            captured=[]
            def lose_once(r):
                response=r.fetch();assert response.status==200;captured.append({'request':r.request.post_data_json,'response':response.json()});r.abort('connectionreset')
            page.route('**/api/cases/lost/comparison/*/confirmation',lose_once,times=1)
            page.locator('#save-correction').click();expect(page.locator('#edit-error')).to_contain_text('保存結果を確認できません')
            committed=get('/__test/audit');dump('lost_committed',captured)
            expect(page.locator('#edit-fee-name')).to_have_value('Response confirmed');page.locator('#save-correction').click();expect(page.locator('#edit-dialog')).not_to_be_visible()
            check('lost response retry same idempotency key',result['requests'][-1]==captured[0]['request'])
            check('lost response creates no duplicate history/memory',get('/__test/audit')==committed)
            # A tab with an older row must receive409 and preserve its draft.
            open_case('many');expect(page.locator('.row-member')).to_have_count(4);page.locator('#edit-fee-name').fill('All displayed members')
            snapshot=get('/api/cases/many');row=snapshot['comparison']['rows'][0]
            body={'expected_case_revision':snapshot['case']['case_revision'],'members':[{'item_id':i['item_id'],'expected_item_version':i['version'],'expected_memory_version':i['current_memory_version']} for i in row['member_items']], 'new_fee_name':'Other tab','idempotency_key':'synthetic-other-tab'}
            check('concurrent other tab actual row save',page.request.post(BASE+'/api/cases/many/comparison/'+row['row_id']+'/confirmation',data=body).ok)
            page.locator('#save-correction').click();expect(page.locator('#edit-error')).to_contain_text('他の更新が反映されています');expect(page.locator('#edit-fee-name')).to_have_value('All displayed members')
            page.screenshot(path=str(out/'browser_version_conflict.png'))
            check('stale request does not partially overwrite',all(i['current_fee_name']=='Other tab' for i in get('/api/cases/many')['items']))
            page.locator('#edit-details > summary').click();page.locator('#refresh-item').click();expect(page.locator('.row-member')).to_have_count(4);page.locator('#save-correction').click();expect(page.locator('#edit-dialog')).not_to_be_visible()
            many=get('/api/cases/many');check('three companies duplicate A all four saved',len(many['items'])==4 and all(i['current_fee_name']=='All displayed members' for i in many['items']))
            dump('final_audit',get('/__test/audit'));check('no browser JS error',not errors);result['status']='PASS';browser.close()
    except Exception:result.update(status='FAIL',error=traceback.format_exc());print(result['error'])
    finally:dump('result',result)
    return 0 if result['status']=='PASS' else 1
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--evidence',type=Path,required=True);raise SystemExit(run(p.parse_args().evidence))
