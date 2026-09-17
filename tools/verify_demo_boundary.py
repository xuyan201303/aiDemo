#!/usr/bin/env python3
"""Browser boundary check ONLY; never substitutes for real SDK/AI two-round acceptance.

Uploads the first two user-provided fictional originals to this local project.
Requires SDK and AI unbound; does not open/parse/upload the third workbook.
"""
import argparse
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import sys
import traceback

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

PROJECT = Path(__file__).resolve().parents[1]
BASE = "http://127.0.0.1:8765"


def canonical_hash(value):
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence-dir", required=True, type=Path)
    parser.add_argument("--chrome", type=Path, default=Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"))
    parser.add_argument("--reader-mode", choices=["REAL_SDK", "MOCK_SDK"], default="REAL_SDK")
    args = parser.parse_args()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    evidence = args.evidence_dir / ("browser_boundary_" + stamp)
    evidence.mkdir(parents=True, exist_ok=False)
    result = {"status": "FAIL", "scope": "LOCAL_BROWSER_UPLOAD_PREVIEW_BLOCKED_BOUNDARY", "reader_mode": args.reader_mode,
              "sdk_calls": 0, "openai_calls": 0, "round_1": "NOT_RUN", "round_2": "NOT_RUN",
              "third_sample": "NOT_OPENED_NOT_PARSED_NOT_UPLOADED", "checks": [], "screenshots": []}
    requests, errors = [], []

    def check(name, condition=True):
        if not condition:
            raise AssertionError(name)
        result["checks"].append(name)
        print("PASS: " + name, flush=True)

    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(executable_path=str(args.chrome), headless=True)
            context = browser.new_context(viewport={"width": 1480, "height": 1100}, accept_downloads=True)
            context.route("**/*", lambda route: route.continue_() if route.request.url.startswith(BASE + "/") else route.abort())
            page = context.new_page()
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.on("request", lambda req: requests.append({"method": req.method, "url": req.url}))
            status = context.request.get(BASE + "/api/status").json()
            check("Live bindings absent; no real calls allowed", not status["sdk"]["ready"] and not status["ai"]["ready"])
            check("Explicit expected reader mode", status["reader_mode"] == args.reader_mode)
            if args.reader_mode == "MOCK_SDK":
                check("Mock raw reader assembled without SDK", status["reader"]["ready"] is True)
            check("Manual debug disabled", status["debug_manual"] is False)
            old = context.request.get(BASE + "/api/cases").json()
            old_item_hashes = {case["case_id"]: canonical_hash(context.request.get(BASE + "/api/cases/" + case["case_id"]).json()["items"]) for case in old}
            page.goto(BASE)
            if args.reader_mode == "MOCK_SDK":
                expect(page.locator("#connection-panel")).to_contain_text("デモ用解析データ／SDK未接続")
            page.locator("#new-case-button").click()
            page.locator("#case-name").fill("接続境界確認・" + args.reader_mode + "・AI未接続 " + stamp)
            page.locator("#new-case-form button[type=submit]").click()
            expect(page.locator("#new-case-form")).to_be_hidden()
            case_id = page.locator("#case-select").input_value()
            check("Distinct boundary case created", bool(case_id))
            result["case_id"] = case_id
            samples = [PROJECT / "demo_materials/samples" / name for name in (
                "01_A社_初回見積.xlsx", "02_B社_初回見積.pdf")]
            result["uploaded_samples"] = [{"filename": p.name, "sha256": sha256(p.read_bytes()).hexdigest()} for p in samples]
            page.locator("#files").set_input_files([str(p) for p in samples])
            page.locator("#upload-button").click()
            expect(page.locator(".document-card")).to_have_count(2)
            expect(page.locator("#upload-button")).to_be_enabled()
            check("Two supplied files uploaded in browser")
            check("Internal manual entry not offered", page.locator(".debug-action").count() == 0)
            before = context.request.get(BASE + "/api/cases/" + case_id).json()
            check("Upload has no manufactured items", before["items"] == [] and before["tasks"] == [])
            result["documents"] = [{"document_id": d["document_id"], "sha256": d["sha256"], "original_filename": d["original_filename"]} for d in before["documents"]]
            for index in range(2):
                card = page.locator(".document-card").nth(index)
                card.get_by_role("button", name="原文をプレビュー", exact=False).click()
                expect(page.locator("#preview-dialog")).to_be_visible()
                if args.reader_mode == "MOCK_SDK":
                    expect(page.locator("#preview-note")).to_contain_text("デモ用解析データ／SDK未接続")
                frame = page.frame_locator("#preview-frame")
                doc = before["documents"][index]
                if doc["original_filename"].endswith(".pdf"):
                    img = frame.locator("#pdf-page-image")
                    expect(img).to_be_visible()
                    expect(img).to_have_js_property("complete", True)
                    check("PDF original page rendered", img.evaluate("e => e.naturalWidth > 0 && e.naturalHeight > 0"))
                else:
                    expect(frame.locator("table").first).to_be_visible()
                    check("Excel original cells visible without SDK claim")
                screenshot = evidence / ("preview_" + str(index + 1) + ".png")
                page.screenshot(path=str(screenshot), full_page=True)
                result["screenshots"].append(str(screenshot))
                page.locator("#preview-dialog .close-dialog").click()
            for doc in before["documents"]:
                selector = '.document-card[data-document-id="' + doc["document_id"] + '"]'
                page.locator(selector + " .process-button").click()
                expect(page.locator(selector + " .document-state")).to_have_attribute("data-status", "BLOCKED")
                expect(page.locator(selector + " .process-button")).to_be_visible()
            after = context.request.get(BASE + "/api/cases/" + case_id).json()
            check("Both attempts persisted as BLOCKED", len(after["tasks"]) == 2 and all(t["status"] == "BLOCKED" for t in after["tasks"]))
            counter = "real_sdk_calls" if args.reader_mode == "MOCK_SDK" else "sdk_calls"
            check("No SDK or AI calls occurred", all(t["events"][-1]["metadata"][counter] == 0 and t["events"][-1]["metadata"]["ai_calls"] == 0 for t in after["tasks"]))
            check("No partial quotation was saved", after["items"] == [])
            expected_states = ["QUEUED", "READING", "EXTRACTING", "BLOCKED"] if args.reader_mode == "MOCK_SDK" else ["QUEUED", "BLOCKED"]
            check("Actual read and blocked events match mode", all([e["state"] for e in t["events"]] == expected_states for t in after["tasks"]))
            if args.reader_mode == "MOCK_SDK":
                for task in after["tasks"]:
                    event = next(e for e in task["events"] if e["state"] == "EXTRACTING")
                    proof = event["metadata"]["read_evidence"]
                    check("Fixture integrity separate from SDK execution", proof["origin"] == "MOCK_SDK" and proof["sdk_identity"]["fixture_integrity"] == "VERIFIED" and event["metadata"]["real_sdk_calls"] == 0)
                    check("Missing AI authorization is explicit", task["error_code"] == "OPENAI_AUTHORIZATION_NOT_CONFIRMED")

            page.reload()
            page.locator("#case-select").select_option(case_id)
            expect(page.locator('.document-state[data-status="BLOCKED"]')).to_have_count(2)
            check("Reload retains task states")
            # Local download endpoint, no externally uploaded output.
            expect(page.locator("#export-button")).to_be_enabled()
            with page.expect_download() as download_info:
                page.locator("#export-button").click()
            download = download_info.value
            exported_bytes = Path(download.path()).read_bytes()
            check("Excel download available from browser", bool(exported_bytes))
            book = load_workbook(BytesIO(exported_bytes), read_only=True, data_only=False)
            try:
                rows = [[cell.value for cell in row] for row in book["未確認事項"]]
                values = [str(v) for row in rows for v in row]
                check("Export keeps both unprocessed filenames", all(p.name in values for p in samples))
                check("Export reports BLOCKED", "BLOCKED" in values)
                if args.reader_mode == "MOCK_SDK":
                    all_values = [str(cell.value) for sheet in book for row in sheet for cell in row]
                    check("Export retains mock source notice", any("デモ用解析データ／SDK未接続" in value for value in all_values))
                check("Export has no invented detail row", book["明細と出典"].max_row == 1)
                result["export"] = {"sha256": sha256(exported_bytes).hexdigest(), "sheets": book.sheetnames, "issues": rows, "detail_row_count": 0}
            finally:
                book.close()
            screenshot = evidence / "blocked_tasks.png"
            page.screenshot(path=str(screenshot), full_page=True)
            result["screenshots"].append(str(screenshot))
            result["tasks"] = after["tasks"]
            current = {case_id: canonical_hash(context.request.get(BASE + "/api/cases/" + case_id).json()["items"]) for case_id in old_item_hashes}
            check("Existing case item snapshots unchanged", current == old_item_hashes)
            result["old_case_item_hashes"] = old_item_hashes
            check("Browser has no JavaScript exceptions", not errors)
            check("No manual input requests", not any("/manual-lines" in r["url"] for r in requests))
            browser.close()
            result["status"] = "PASS_BOUNDARY_ONLY"
    except Exception:
        result["error"] = traceback.format_exc()
    finally:
        result["browser_errors"] = errors
        result["requests"] = requests
        result["argv"] = [sys.executable, *sys.argv]
        result["exit_code"] = 0 if result["status"] == "PASS_BOUNDARY_ONLY" else 1
        (evidence / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps({"status": result["status"], "exit_code": result["exit_code"], "evidence": str(evidence)}, ensure_ascii=False))
    return result["exit_code"]


if __name__ == "__main__":
    raise SystemExit(main())
