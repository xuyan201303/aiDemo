"""Approved live browser rehearsal. Mock reader only; no fake AI or direct line entry."""
from pathlib import Path
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import argparse, json, time, traceback
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
BASE="http://127.0.0.1:8765"
NAMES=["01_A社_初回見積.xlsx","02_B社_初回見積.pdf","03_A社_次回見積.xlsx"]
def run(out, resume_empty_case=None, resume_first_evidence=None):
    out.mkdir(parents=True,exist_ok=False)
    result={"scope":"MOCK_SDK_MEMORY_FIRST_UNKNOWN_ONLY_AI","real_sdk":"NOT_RUN","checks":[],"rounds":{}}
    def save(name,data):
        (out/(name+".json")).write_text(json.dumps(data,ensure_ascii=False,indent=2))
    def check(name,value):
        if not value:raise AssertionError(name)
        result["checks"].append(name);print("PASS "+name,flush=True)
    try:
        with sync_playwright() as pw:
            browser=pw.chromium.launch(executable_path="/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",headless=True)
            page=browser.new_page(viewport={"width":1500,"height":1150},accept_downloads=True)
            page_errors=[]
            page.on("pageerror",lambda e:page_errors.append(str(e)))
            def get(path):
                r=page.request.get(BASE+path);assert r.ok;return r.json()
            status=get("/api/status")
            save("status_before",status)
            budget=status["ai"]["budget"]
            if resume_first_evidence:
                budget=json.loads((resume_first_evidence/"status_before.json").read_text())["ai"]["budget"]
                check("Resume preserves fixed grant",budget["authorization_id"]==status["ai"]["budget"]["authorization_id"] and budget["max_requests"]==status["ai"]["budget"]["max_requests"])
                result["resumed_evidence"]=str(resume_first_evidence)
            check("Approved real AI + Mock reader ready",status["reader_mode"]=="MOCK_SDK" and status["reader"]["ready"] and status["ai"]["ready"])
            check("Fixed model/budget preserved",json.loads((ROOT/"demo_materials/config/ai-authorization.json").read_text())["model"]=="gpt-5.6-terra" and budget["authorization_id"]=="sanzo-three-file-mock-sdk-v1" and budget["max_requests"]==json.loads((ROOT/"demo_materials/config/ai-authorization.json").read_text())["max_requests"] and budget["max_requests"]-budget["requests_used"]>=3)
            page.goto(BASE,wait_until="networkidle")
            check("Current customer code loaded",page.locator("#amount-only").count()==0 and "Demoモード" in page.locator("#connection-panel").inner_text())
            stamp=datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
            def new_case(round_name):
                page.locator("#new-case-button").click()
                page.locator("#case-name").fill("No.202608310011 Memory First "+stamp+" "+round_name)
                with page.expect_response(lambda r:r.request.method=="POST" and r.url==BASE+"/api/cases") as response:
                    page.locator("#new-case-form button[type=submit]").click()
                case=response.value.json()["case_id"]
                expect(page.locator("#case-select")).to_have_value(case)
                return case
            def batch(case,names,tag):
                expect(page.locator("#files")).to_be_enabled(timeout=10000)
                page.locator("#files").set_input_files([str(ROOT/"demo_materials/samples"/n) for n in names])
                expect(page.locator("#pending-files li")).to_have_count(len(names))
                page.screenshot(path=str(out/("browser_"+tag+"_selected.png")),full_page=True)
                page.locator("#upload-button").click()
                deadline=time.monotonic()+300
                while time.monotonic()<deadline:
                    page.wait_for_timeout(1000)
                    snapshot=get("/api/cases/"+case)
                    tasks=snapshot["tasks"]
                    terminal=[t for t in tasks if t["status"] in {"AVAILABLE","REVIEW_REQUIRED","FAILED","BLOCKED","INTERRUPTED"}]
                    if any(t["status"] in {"FAILED","BLOCKED","INTERRUPTED"} for t in terminal):
                        save(tag+"_failed",snapshot)
                        page.screenshot(path=str(out/("browser_"+tag+"_failed.png")),full_page=True)
                        raise RuntimeError("Live batch stopped; preserved failed task, no retry")
                    if len(terminal)==len(names):
                        expect(page.locator("#files")).to_be_enabled(timeout=10000)
                        save(tag,snapshot)
                        return snapshot
                raise TimeoutError("Batch not terminal; no resubmit")
            if resume_first_evidence:
                prior_result=json.loads((resume_first_evidence/"progress.json").read_text())
                first=prior_result["rounds"]["first_case_id"]
                prior=get("/api/cases/"+first)
                check("Resume completed first uploads only",len(prior["documents"])==2 and len(prior["items"])==10 and all(t["status"] in {"AVAILABLE","REVIEW_REQUIRED"} for t in prior["tasks"]))
                check("Resume samples bound to actual files",{d["sha256"] for d in prior["documents"]}=={sha256((ROOT/"demo_materials/samples"/n).read_bytes()).hexdigest() for n in NAMES[:2]})
                page.locator("#case-select").select_option(first)
            elif resume_empty_case:
                prior=get("/api/cases/"+resume_empty_case)
                check("Resume only the empty pre-upload case",not prior["documents"] and not prior["items"] and not prior["tasks"])
                page.locator("#case-select").select_option(resume_empty_case)
                first=resume_empty_case
            else:
                first=new_case("第1輪")
            result["rounds"]["first_case_id"]=first;save("progress",result)
            def routing(snapshot, tag):
                records=[]; calls=[]
                by_id={i["item_id"]:i for i in snapshot["items"]}
                for task in snapshot["tasks"]:
                    meta=task["result_json"]["ai_metadata"]
                    check(tag+" deterministic facts/no quotation extraction",meta["quotation_extraction_requests"]==0 and meta["fact_materialization"]=="complete-reader-fields-v1")
                    requested=set(meta["classification_requested_item_ids"])
                    check(tag+" unknown-only classification",all(meta["memory_lookup_results"][i]["decision"]!="MEMORY_APPLIED" for i in requested))
                    check(tag+" no extraction calls",all(c["purpose"]=="fee_classification" and c["execution_kind"]=="OPENAI_API" and c["http_status"]==200 for c in meta["calls"]))
                    calls.extend(meta["calls"])
                    evidence={b["source_id"]:b for b in task["result_json"]["read_evidence"]["sources"]}
                    for item_id in task["result_json"]["item_ids"]:
                        item=by_id[item_id]; source=item["source_ref"]; fields=source["field_sources"]
                        check(tag+" fact provenance "+item_id,all(v=="MOCK_SDK" for v in item["fact_sources"].values()) and all(v["original_value"]==evidence[v["source_block_id"]]["raw_text"] for v in fields.values()))
                        records.append({"item_id":item_id,"raw_label":item["observation"]["raw_label"],
                            "current_price":item["current_price"],"currency":item["currency"],
                            "classification_source":item["classification_source"],"fact_sources":item["fact_sources"],
                            "classification_requested":item_id in requested,
                            "lookup":meta["memory_lookup_results"][item_id],
                            "source_ref":source,"committed_match":next(m for m in task["result_json"]["classification_results"] if m["item_id"]==item_id)})
                evidence={"quotation_extraction_requests":0,"classification_requests":len(calls),"calls":calls,"items":records}
                save(tag+"_routing",evidence);result["rounds"][tag+"_calls"]=len(calls)
                return evidence
            before=get("/api/cases/"+first) if resume_first_evidence else batch(first,NAMES[:2],"first_before_confirmation")
            if resume_first_evidence:save("first_before_confirmation",before)
            routing(before,"first")
            check("First round retains ten original items",len(before["items"])==10)
            check("Two uploaded documents, two forwarder columns",len(before["documents"])==2 and len(before["comparison"]["companies"])==2)
            result["first_pair_count"]=len(before["comparison"]["rows"])
            five_rows=len(before["comparison"]["rows"])==5 and all(sum(map(len,r["cells"].values()))==2 for r in before["comparison"]["rows"])
            result["five_paired_rows"]="PASS" if five_rows else "NOT_MET"
            if not five_rows:
                conflicts=[i for i in before["items"] if i.get("current_memory_status") in {"CONFLICTED","REVOKED"}]
                check("Existing memory conflict remains unresolved without AI override",bool(conflicts) and all(not i["current_fee_id"] and not i["ai_suggested_name"] for i in conflicts))
                save("preserved_memory_conflicts",[{"item_id":i["item_id"],"raw_label":i["observation"]["raw_label"],"memory_status":i["current_memory_status"],"memory_version":i["current_memory_version"]} for i in conflicts])
            b=[i for i in before["items"] if "B（架空）" in i["observation"]["forwarder_id"]]
            check("B proposals stay unconfirmed and conditions unknown",len(b)==5 and all(i.get("ai_suggested_name") and not i["current_fee_id"] and not i["observation"]["context"]["ready"] for i in b))
            for i in b:
                check("B price cell has no fee controls "+i["item_id"],page.locator('[data-item-id="'+i["item_id"]+'"] button').count()==0)
            table=page.locator("#comparison-table").inner_text()
            check("Customer table has no implementation/source clutter",not any(x in table for x in ["合成・","Synthetic ","Composite ","Merged ","MOCK_SDK","SDK未接続",".xlsx",".pdf","transport_mode"]))
            page.locator("#comparison-table").screenshot(path=str(out/"browser_first_table_before.png"))
            delivery=next(i for i in before["items"] if i["observation"]["raw_label"]=="配送手数料")
            target_id=delivery["item_id"]
            delivery_row=next(r for r in before["comparison"]["rows"] if target_id in r["item_ids"])
            page.locator('[data-row-id="'+delivery_row["row_id"]+'"] .row-confirm-button').click()
            check("Price and currency readonly",page.locator("#edit-price,#edit-currency,#amount-only").count()==0)
            page.locator(".member-preview").first.click()
            page.frame_locator("#preview-frame").locator("body").wait_for()
            page.wait_for_timeout(700)
            check("Excel preview uses first uploaded original",delivery["observation"]["document_id"] in page.locator("#preview-frame").get_attribute("src"))
            page.screenshot(path=str(out/"browser_first_excel_preview.png"),full_page=True)
            page.locator("#preview-dialog .close-dialog").click()
            target=next((f for f in before["fees"] if (f.get("display_name") or f.get("name"))=="配送費"),None)
            if target:
                page.locator("#edit-fee-name").fill(target["display_name"])
                page.locator('[data-fee-id="'+target["fee_id"]+'"]').click()
            else:
                page.locator("#edit-fee-name").fill("配送費")
            page.screenshot(path=str(out/"browser_confirmation.png"),full_page=True)
            with page.expect_response(lambda r:r.request.method=="POST" and r.url.endswith("/confirmation")) as correction:
                page.locator("#save-correction").click()
            response=correction.value
            save("human_confirmation",{"http_status":response.status,"request":response.request.post_data_json,"response":response.json()})
            check("Name confirmation request contains no price/currency",not {"unit_price","currency"} & set(response.request.post_data_json))
            check("Confirmation saved",response.ok)
            expect(page.locator("#edit-dialog")).not_to_be_visible()
            after=get("/api/cases/"+first);save("first_after_confirmation",after)
            confirmed=next(i for i in after["items"] if i["item_id"]==target_id)
            check("Confirmed customer name and source price preserved",confirmed["current_fee_name"]=="配送費" and confirmed["current_price"]==delivery["current_price"]=="8000" and confirmed["currency"]==delivery["currency"])
            check("Raw label and source preserved",confirmed["observation"]==delivery["observation"] and confirmed["source_ref"]==delivery["source_ref"])
            check("Reusable memory saved",next(m for m in response.json()["members"] if m["item_id"]==target_id)["memory_effect"]=="ACTIVE" and confirmed["current_memory_version"]>0)
            check("Other current-case items unchanged",all(i==next(j for j in before["items"] if j["item_id"]==i["item_id"]) or all(i[k]==next(j for j in before["items"] if j["item_id"]==i["item_id"])[k] for k in ["observation","source_ref","current_price","currency","current_fee_id","ai_suggested_name"]) for i in after["items"] if i["item_id"] not in delivery_row["item_ids"]))
            page.locator("#comparison-table").screenshot(path=str(out/"browser_first_table_after.png"))
            # Open B's uploaded PDF itself, not a fixture-generated screenshot.
            pdf=next(d for d in after["documents"] if d["original_filename"].endswith(".pdf"))
            page.locator('[data-document-id="'+pdf["document_id"]+'"] .document-actions button').first.click()
            page.frame_locator("#preview-frame").locator("body").wait_for();page.wait_for_timeout(700)
            check("PDF preview uses actual uploaded PDF",pdf["document_id"] in page.locator("#preview-frame").get_attribute("src"))
            page.screenshot(path=str(out/"browser_pdf_preview.png"),full_page=True)
            page.locator("#preview-dialog .close-dialog").click()
            def export(snapshot,tag):
                with page.expect_download() as download:page.locator("#export-button").click()
                path=out/(tag+".xlsx");download.value.save_as(path)
                book=load_workbook(path,data_only=False)
                sheet=book["原価単価一覧表"];projection=snapshot["comparison"]
                check(tag+" shared row count",sheet.max_row==len(projection["rows"])+3)
                rows=[]
                for number,row in enumerate(projection["rows"],4):
                    check(tag+" shared row name "+str(number),sheet.cell(number,1).value==row["display_name"])
                    values=[]
                    for col,company in enumerate(projection["companies"],2):
                        expected="\\n".join(c["price"]+" "+c["currency"] for c in row["cells"][company]) or "—"
                        check(tag+" exact current price/currency "+str(number)+company,sheet.cell(number,col).value==expected)
                        values.append(expected)
                    rows.append({"name":row["display_name"],"values":values})
                alltext=" ".join(str(c.value or "") for cells in sheet for c in cells)
                check(tag+" clean main sheet",not any(s in alltext for s in ["MOCK_SDK","SDK未接続","AI_REQUIRED","MEMORY_APPLIED","REVIEW_REQUIRED","原文："]))
                save(tag+"_verification",{"sha256":sha256(path.read_bytes()).hexdigest(),"rows":rows,"sheets":book.sheetnames})
                book.close()
            export(after,"first_export")
            second=new_case("第2輪")
            result["rounds"]["second_case_id"]=second;save("progress",result)
            fresh=batch(second,[NAMES[2]],"second")
            routing(fresh,"second")
            next_delivery=next(i for i in fresh["items"] if i["observation"]["raw_label"]=="配送手数料")
            check("Second round reuses confirmed fee identity",next_delivery["current_fee_id"]==confirmed["current_fee_id"] and next_delivery["saved_status"]=="MEMORY_APPLIED" and next_delivery["current_fee_name"]=="配送費")
            check("Second price comes from new file",next_delivery["current_price"]=="9500" and next_delivery["currency"]=="JPY")
            old=get("/api/cases/"+first)
            old_delivery=next(i for i in old["items"] if i["item_id"]==target_id)
            check("Old first-case price remains8000",old_delivery["current_price"]=="8000")
            metadata=fresh["tasks"][0]["result_json"]["ai_metadata"]
            check("Matched item skips name classification",next_delivery["item_id"] in metadata["classification_skipped_item_ids"] and next_delivery["classification_source"]=="MEMORY")
            check("9500 comes from this reader amount field",next_delivery["source_ref"]["field_sources"]["raw_value"]["original_value"]=="9500" and next_delivery["fact_sources"]["amount"]=="MOCK_SDK")
            target_lookup=metadata["memory_lookup_results"][next_delivery["item_id"]]
            check("Second lookup uses first round human memory version",target_lookup["memory_version"]==confirmed["current_memory_version"])
            result["rounds"]["second_delivery_classification_calls"]=0
            page.locator("#comparison-table").screenshot(path=str(out/"browser_second_table.png"))
            next_row=next(r for r in fresh["comparison"]["rows"] if next_delivery["item_id"] in r["item_ids"])
            page.locator('[data-row-id="'+next_row["row_id"]+'"] .row-confirm-button').click()
            page.locator(".member-preview").first.click()
            page.frame_locator("#preview-frame").locator("body").wait_for()
            check("Second preview uses new document",next_delivery["observation"]["document_id"] in page.locator("#preview-frame").get_attribute("src"))
            page.screenshot(path=str(out/"browser_second_preview.png"),full_page=True)
            page.locator("#preview-dialog .close-dialog").click();page.locator("#edit-dialog .close-dialog").click()
            export(fresh,"second_export")
            final_status=get("/api/status");save("status_after",final_status)
            result["budget_before"]=budget;result["budget_after"]=final_status["ai"]["budget"]
            check("Actual budget delta equals classification requests",final_status["ai"]["budget"]["requests_used"]-budget["requests_used"]==result["rounds"]["first_calls"]+result["rounds"]["second_calls"])
            check("No browser script errors",not page_errors)
            result["status"]="SUCCESS"
            browser.close()
    except Exception:
        result["status"]="PARTIAL";result["error"]=traceback.format_exc()
        print(result["error"],flush=True)
    finally:
        save("result",result)
    return 0 if result["status"]=="SUCCESS" else 1
if __name__=="__main__":
    p=argparse.ArgumentParser();p.add_argument("--evidence",type=Path,required=True);p.add_argument("--execute-approved-ai",action="store_true");p.add_argument("--resume-empty-case")
    p.add_argument("--resume-first-evidence",type=Path)
    args=p.parse_args()
    if not args.execute_approved_ai:p.error("Explicit --execute-approved-ai is required; consumes shared approved budget.")
    raise SystemExit(run(args.evidence,args.resume_empty_case,args.resume_first_evidence))
