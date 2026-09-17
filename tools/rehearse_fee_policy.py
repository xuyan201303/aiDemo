"""Explicit approved browser rehearsal: real fixed-wheel SDK, existing AI and DB.

This is a test driver, not an input reader. Only the existing name confirmation
UI writes human names. No company or transport preflight is submitted. No amounts, labels or memory
answers are submitted to processing. Never retries processing or renews quota.
"""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import argparse
import json
import re
import time
import traceback

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
BASE = 'http://127.0.0.1:8765'
NAMES = ('01_A社_初回見積.xlsx', '02_B社_初回見積.pdf', '03_A社_次回見積.xlsx')
WHEEL_SHA = 'd791e0602ef138f50a065cd4cc5d3dca801226822b74328230f2faa2c0c75666'
TERMINAL = {'AVAILABLE', 'REVIEW_REQUIRED', 'FAILED', 'BLOCKED', 'INTERRUPTED'}


def run(out, resume_first_evidence=None, stop_after_first=False):
    out.mkdir(parents=True, exist_ok=False)
    result = {'scope': 'SIMPLIFY_CORE_LEARNING_FLOW_R1', 'status': 'RUNNING',
              'real_sdk': 'NOT_RUN', 'real_ai': 'NOT_RUN', 'two_rounds': 'NOT_RUN',
              'checks': [], 'rounds': {}, 'sdk_preflights': [], 'human_role_confirmations': [],
              'human_name_confirmations': [], 'unresolved': [], 'process_requests': []}
    page = browser = None

    def save(name, value):
        path = out / (name + '.json'); path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')

    def checkpoint():
        save('progress', result)

    def check(name, passed):
        if not passed:
            raise AssertionError(name)
        result['checks'].append(name)
        print('PASS ' + name, flush=True)

    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome', headless=True)
            page = browser.new_page(viewport={'width': 1366, 'height': 768}, accept_downloads=True)
            page_errors = []
            page.on('pageerror', lambda error: page_errors.append(str(error)))

            def get(path):
                response = page.request.get(BASE + path)
                if not response.ok:
                    raise RuntimeError('Read-only status request failed: ' + path + ' HTTP ' + str(response.status))
                return response.json()

            def capture_response(response):
                # Do not collect headers, credentials, or unrelated requests.
                if response.request.method == 'POST' and re.fullmatch(re.escape(BASE) + r'/api/documents/[^/]+/process', response.url):
                    result['process_requests'].append({'url':response.url,'http_status':response.status,'body':response.request.post_data_json})
                    checkpoint()
                if response.request.method == 'GET' and re.fullmatch(re.escape(BASE) + r'/api/documents/[^/]+/field-sources', response.url):
                    try:
                        value = response.json()
                        result['sdk_preflights'].append({'url': response.url, 'http_status': response.status, 'result': value})
                        checkpoint()
                    except Exception:
                        result['sdk_preflights'].append({'url': response.url, 'http_status': response.status, 'body_unavailable': True})

            page.on('response', capture_response)
            status = get('/api/status'); save('status_before', status)
            auth = json.loads((ROOT / 'demo_materials/config/ai-authorization.json').read_text())
            check('Real reader and approved AI runtime are assembled', status['reader_mode'] == 'REAL_SDK' and status['reader']['ready'] and status['sdk']['ready'] and status['ai']['ready'])
            budget = status['ai']['budget']
            check('Fixed approved model and shared cumulative grant unchanged', auth['model'] == 'gpt-5.6-terra' and auth['max_requests'] == 100 and budget['authorization_id'] == 'sanzo-three-file-mock-sdk-v1' and budget['max_requests'] == 100)
            check('Existing quota has room for at most three unknown-only calls', budget['requests_used'] <= 97)
            result['budget_before'] = budget
            if resume_first_evidence:
                previous = json.loads((resume_first_evidence / 'progress.json').read_text())
                check('Resume has first case and no already-created second case', bool(previous.get('rounds', {}).get('first_case_id')) and not previous.get('rounds', {}).get('second_case_id'))
                previous_budget = json.loads((resume_first_evidence / 'status_before.json').read_text())['ai']['budget']
                check('Resume preserves budget grant and never resets usage', previous_budget['authorization_id'] == budget['authorization_id'] and previous_budget['max_requests'] == budget['max_requests'] and previous_budget['requests_used'] <= budget['requests_used'])
                result['resumed_evidence'] = str(resume_first_evidence)
                result['budget_before'] = previous_budget
                result['sdk_preflights'] = previous.get('sdk_preflights', [])
            page.goto(BASE, wait_until='networkidle')
            check('Customer page remains price-readonly', page.locator('#edit-price,#edit-currency,#amount-only').count() == 0)
            stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')

            def new_case(label):
                page.locator('#new-case-button').click()
                page.locator('#case-name').fill('No.202608310011 SIMPLIFY R1 ' + stamp + ' ' + label)
                page.screenshot(path=str(out / ('browser_' + label + '_new_case.png')), full_page=True)
                with page.expect_response(lambda r: r.request.method == 'POST' and r.url == BASE + '/api/cases') as pending:
                    page.locator('#new-case-form button[type=submit]').click()
                response = pending.value
                check('New isolated rehearsal case created ' + label, response.ok)
                case_id = response.json()['case_id']
                expect(page.locator('#case-select')).to_have_value(case_id)
                return case_id

            def batch(case_id, names, tag):
                expect(page.locator('#files')).to_be_enabled(timeout=15000)
                page.locator('#files').set_input_files([str(ROOT / 'demo_materials/samples' / name) for name in names])
                expect(page.locator('#pending-files li')).to_have_count(len(names))
                page.screenshot(path=str(out / ('browser_' + tag + '_selected.png')), full_page=True)
                check(tag + ' no field review DOM or loaded script', page.locator('#field-review-dialog,script[src*=field_review]').count() == 0)
                # Exactly one user action for the selected original files.
                page.locator('#upload-button').click()
                page.screenshot(path=str(out / ('browser_' + tag + '_direct_processing.png')), full_page=True)
                deadline = time.monotonic() + 360
                while time.monotonic() < deadline:
                    page.wait_for_timeout(500)
                    check_no_role = page.locator('#field-review-dialog').count() == 0
                    if not check_no_role or result['sdk_preflights']:
                        raise AssertionError('Unexpected company/transport preflight')
                    snapshot = get('/api/cases/' + case_id)
                    if any(task['status'] in {'FAILED', 'BLOCKED', 'INTERRUPTED'} for task in snapshot['tasks']):
                        save(tag + '_failed', snapshot)
                        page.screenshot(path=str(out / ('browser_' + tag + '_failed.png')), full_page=True)
                        raise RuntimeError('Processing failed: task evidence preserved; no retry or subsequent AI submission')
                    done = [task for task in snapshot['tasks'] if task['status'] in TERMINAL]
                    if len(done) == len(names):
                        expect(page.locator('#files')).to_be_enabled(timeout=15000)
                        save(tag, snapshot)
                        return snapshot
                raise TimeoutError('Batch did not finish; preserved tasks without re-submission')

            def routing(snapshot, tag):
                calls, items, sdk_calls = [], [], []
                by_id = {item['item_id']: item for item in snapshot['items']}
                for task in snapshot['tasks']:
                    payload = task['result_json']; evidence = payload['read_evidence']; identity = evidence['sdk_identity']
                    check(tag + ' actual fixed SDK provenance ' + task['task_id'], evidence['origin'] == 'REAL_SDK' and identity['wheel_sha256'] == WHEEL_SHA and identity['installed_file_check'] == 'VERIFIED' and not identity.get('fixture_sha256'))
                    sdk_calls.extend(identity['public_calls'])
                    for index, call in enumerate(identity['public_calls'], 1):
                        check(tag + ' SDK public call succeeded ' + str(index), call['status'] == 'SUCCESS')
                        raw = identity['public_output_json'][call['api']].encode('utf-8')
                        check(tag + ' exact public output SHA ' + str(index), sha256(raw).hexdigest() == call['sha256'])
                        path = out / 'sdk_outputs' / tag / task['document_id'] / (str(index).zfill(2) + '.json')
                        path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(raw)
                    meta = payload['ai_metadata']; requested = set(meta['classification_requested_item_ids'])
                    check(tag + ' quotation extraction AI is zero', meta['quotation_extraction_requests'] == 0 and meta['fact_materialization'] == 'complete-reader-fields-v1')
                    check(tag + ' only unknown names sent to AI', all(meta['memory_lookup_results'][item_id]['decision'] != 'MEMORY_APPLIED' for item_id in requested))
                    check(tag + ' real authorized classification calls only', all(c['purpose'] == 'fee_classification' and c['execution_kind'] == 'OPENAI_API' and c['http_status'] == 200 and c['requested_model'] == 'gpt-5.6-terra' and c['automatic_retries'] == 0 for c in meta['calls']))
                    calls.extend(meta['calls'])
                    source_blocks = {block['source_id']: block for block in evidence['sources']}
                    for item_id in payload['item_ids']:
                        item = by_id[item_id]; fields = item['source_ref']['field_sources']
                        check(tag + ' immutable real SDK facts ' + item_id, item['input_kind'] == 'SDK_OBSERVED' and all(value == 'REAL_SDK' for value in item['fact_sources'].values()) and all(v['original_value'] == source_blocks[v['source_block_id']]['raw_text'] for v in fields.values()))
                        items.append({'item_id': item_id, 'raw_label': item['observation']['raw_label'], 'price': item['current_price'], 'currency': item['currency'], 'classification_source': item['classification_source'], 'classification_requested': item_id in requested, 'lookup': meta['memory_lookup_results'][item_id], 'source_ref': item['source_ref'], 'committed_match': next(m for m in payload['classification_results'] if m['item_id'] == item_id)})
                summary = {'classification_requests': len(calls), 'quotation_extraction_requests': 0, 'calls': calls, 'sdk_public_calls': sdk_calls, 'items': items}
                save(tag + '_routing', summary)
                result['rounds'][tag + '_calls'] = len(calls)
                result['real_sdk'] = 'PASS'
                if calls:
                    result['real_ai'] = 'EXECUTED'
                checkpoint()
                return summary

            if resume_first_evidence:
                first_id = previous['rounds']['first_case_id']
                before = get('/api/cases/' + first_id)
                check('Resume uses completed originals and tasks only', len(before['documents']) == 2 and len(before['items']) == 10 and all(t['status'] in {'AVAILABLE', 'REVIEW_REQUIRED'} for t in before['tasks']))
                check('Resume original hashes unchanged', {d['sha256'] for d in before['documents']} == {sha256((ROOT / 'demo_materials/samples' / name).read_bytes()).hexdigest() for name in NAMES[:2]})
                page.locator('#case-select').select_option(first_id)
                save('first_before_confirmation', before)
            else:
                first_id = new_case('第1輪')
                result['rounds']['first_case_id'] = first_id; checkpoint()
                before = batch(first_id, NAMES[:2], 'first_before_confirmation')
            result['rounds']['first_case_id'] = first_id; checkpoint()
            first_route = routing(before, 'first')
            check('Both originals produce all ten quotation items', len(before['items']) == 10 and len(before['documents']) == 2)
            a_doc = next(d for d in before['documents'] if d['original_filename'] == NAMES[0])
            b_doc = next(d for d in before['documents'] if d['original_filename'] == NAMES[1])
            a_items = [i for i in before['items'] if i['observation']['document_id'] == a_doc['document_id']]
            b_items = [i for i in before['items'] if i['observation']['document_id'] == b_doc['document_id']]
            check('No document roles were invented or requested', len(a_items) == len(b_items) == 5 and all(i['observation']['forwarder_id'] is None and i['observation']['context']['transport_mode'] is None for i in a_items+b_items))
            if resume_first_evidence:
                original = json.loads((resume_first_evidence / 'first_before_confirmation.json').read_text())
                b_before = [i for i in original['items'] if i['observation']['document_id'] == b_doc['document_id']]
                check('Preserved pre-confirmation B snapshot has no learned memory', all(i['current_memory_version'] == 0 and i['current_fee_id'] is None for i in b_before))
                check('Resumed B confirmations persisted across restart', all(i['current_memory_version'] > 0 and i['saved_status'] == 'HUMAN_CONFIRMED' for i in b_items))
            else:
                check('B AI suggestions have no confirmation memory', all(i['current_memory_version'] == 0 and i['current_fee_id'] is None for i in b_items))
            result['unresolved'].append({'fields':['company','transport_mode'], 'state':'OPTIONAL_METADATA_UNKNOWN', 'not_a_learning_prerequisite':True})
            page.screenshot(path=str(out / 'browser_first_ai_unknown_state.png'), full_page=True)
            page.locator('#comparison-table').screenshot(path=str(out / 'browser_first_table_before.png'))
            delivery = next(i for i in a_items if i['observation']['raw_label'] == '配送手数料')
            check('First delivery amount is SDK8000', delivery['current_price'] == '8000' and delivery['source_ref']['field_sources']['raw_value']['original_value'] == '8000')

            def preview_item(snapshot, item, tag):
                row = next(row for row in snapshot['comparison']['rows'] if item['item_id'] in row['item_ids'])
                page.locator('[data-row-id="' + row['row_id'] + '"] .row-confirm-button').click()
                draft = page.locator('#edit-fee-name').input_value()
                page.locator('[data-member-id="' + item['item_id'] + '"] .member-preview').click()
                page.frame_locator('#preview-frame').locator('body').wait_for()
                url = page.locator('#preview-frame').get_attribute('src')
                check(tag + ' preview points to current original', item['observation']['document_id'] in url)
                source = item['source_ref']; locator = source['locator']
                if source['kind'] == 'EXCEL_RANGE':
                    check(tag + ' source cell remains in preview URL', locator['A1_range'] in url)
                    expect(page.frame_locator('#preview-frame').locator('[data-cell="' + locator['A1_range'] + '"]')).to_contain_text(item['observation']['raw_label'])
                else:
                    check(tag + ' PDF remains page-only without invented bbox', source['locator_status'] == 'PAGE_ONLY' and set(locator) == {'page_number'})
                page.screenshot(path=str(out / ('browser_' + tag + '_preview.png')), full_page=True)
                page.locator('#preview-dialog .close-dialog').click()
                expect(page.locator('#edit-fee-name')).to_have_value(draft)
                page.locator('#edit-dialog .close-dialog').click()

            preview_item(before, delivery, 'first_delivery')
            preview_item(before, b_items[0], 'first_pdf')
            # The operator confirms each displayed row using existing UI. These
            # expected name pairings are rehearsal choices, never SDK/AI inputs.
            counterparts = {"OCEAN FREIGHT 20'": "O/F 20'", "OCEAN FREIGHT 40'": "O/F 40'",
                "DOCUMENTATION FEE":"DOC FEE", "CUSTOMS CLEARANCE":"通関手数料", "DELIVERY CHARGE":"配送手数料"}
            ordered = [delivery] + [item for item in a_items if item['item_id'] != delivery['item_id']] + b_items
            for number, original in enumerate(ordered, 1):
                snapshot = get('/api/cases/' + first_id)
                current = next(i for i in snapshot['items'] if i['item_id'] == original['item_id'])
                if original['observation']['document_id'] == b_doc['document_id']:
                    peer = next(i for i in snapshot['items'] if i['observation']['document_id'] == a_doc['document_id'] and i['observation']['raw_label'] == counterparts[current['observation']['raw_label']])
                    desired = peer['current_fee_name']
                else:
                    desired = '配送費' if original['item_id'] == delivery['item_id'] else (current.get('current_fee_name') or current.get('ai_suggested_name'))
                check('Observed candidate exists before human confirmation ' + str(number), isinstance(desired, str) and bool(desired.strip()))
                if current.get('saved_status') == 'HUMAN_CONFIRMED' and current.get('current_fee_name') == desired:
                    result['human_name_confirmations'].append({'item_id': current['item_id'], 'name': desired, 'preserved_prior_confirmation': True, 'memory_version': current['current_memory_version']})
                    continue
                row = next(r for r in snapshot['comparison']['rows'] if current['item_id'] in r['item_ids'])
                page.locator('[data-row-id="' + row['row_id'] + '"] .row-confirm-button').click()
                page.locator('#edit-fee-name').fill(desired)
                candidates = page.locator('#fee-candidates [data-fee-id]')
                if candidates.count():
                    exact = candidates.filter(has_text=desired)
                    if exact.count() == 1:
                        exact.first.click()
                check('Normal confirmation form has no editable price', page.locator('#edit-dialog input[name=unit_price],#edit-price,#edit-currency').count() == 0)
                if number == 1:
                    page.screenshot(path=str(out / 'browser_human_delivery_confirmation.png'), full_page=True)
                    page.locator('#edit-dialog').screenshot(path=str(out / 'browser_human_delivery_dialog.png'))
                with page.expect_response(lambda r: r.request.method == 'POST' and r.url.endswith('/confirmation')) as pending:
                    page.locator('#save-correction').click()
                response = pending.value
                record = {'item_id': current['item_id'], 'name': desired, 'http_status': response.status, 'request': response.request.post_data_json, 'response': response.json()}
                result['human_name_confirmations'].append(record); checkpoint()
                check('Human name transaction committed ' + str(number), response.ok)
                check('No price facts submitted for confirmation ' + str(number), not {'unit_price', 'currency', 'amount'} & set(record['request']))
                expect(page.locator('#edit-dialog')).not_to_be_visible()
                saved = get('/api/cases/' + first_id)
                after_item = next(i for i in saved['items'] if i['item_id'] == current['item_id'])
                check('Confirmation preserves original/current prices and sources ' + str(number), all(after_item[k] == current[k] for k in ('observation', 'source_ref', 'current_price', 'currency')))
                check('Grounded member creates active fee-name memory ' + str(number), next(m for m in record['response']['members'] if m['item_id'] == current['item_id'])['memory_effect'] == 'ACTIVE' and after_item['current_memory_version'] > 0)
            first_saved = get('/api/cases/' + first_id); save('first_after_confirmation', first_saved)
            confirmed = next(i for i in first_saved['items'] if i['item_id'] == delivery['item_id'])
            check('All five A facts now individually confirmed', all(i.get('current_fee_id') and i.get('current_memory_version', 0) > 0 for i in first_saved['items'] if i['observation']['document_id'] == a_doc['document_id']))
            check('All five B names human-confirmed and learned despite unknown roles', all(i.get('current_memory_version', 0) > 0 and i['current_memory_status']=='ACTIVE' and i['saved_status']=='HUMAN_CONFIRMED' for i in first_saved['items'] if i['observation']['document_id'] == b_doc['document_id']))
            check('Five paired comparison rows use filenames for unknown companies', len(first_saved['comparison']['rows'])==5 and all(len(r['member_items'])==2 for r in first_saved['comparison']['rows']) and set(first_saved['comparison']['companies'])==set(NAMES[:2]))
            page.screenshot(path=str(out / 'browser_first_saved_learning.png'), full_page=True)
            page.locator('#comparison-table').screenshot(path=str(out / 'browser_first_B_pdf_paired.png'))
            check('Delivery final human name saved', confirmed['current_fee_name'] == '配送費' and confirmed['current_price'] == '8000')
            page.locator('#comparison-table').screenshot(path=str(out / 'browser_first_table_after.png'))

            def export(snapshot, tag):
                page.screenshot(path=str(out / ('browser_' + tag + '_download.png')), full_page=True)
                with page.expect_download() as pending:
                    page.locator('#export-button').click()
                download = pending.value
                raw = Path(download.path()).read_bytes()
                book = load_workbook(BytesIO(raw), data_only=False)
                try:
                    sheet = book['原価単価一覧表']; projection = snapshot['comparison']
                    check(tag + ' shared projection row count', sheet.max_row == len(projection['rows']) + 3)
                    rows = []
                    for row_number, row in enumerate(projection['rows'], 4):
                        check(tag + ' row name ' + str(row_number), sheet.cell(row_number, 1).value == row['display_name'])
                        values = []
                        for column, company in enumerate(projection['companies'], 2):
                            expected = '\n'.join(cell['price'] + ' ' + cell['currency'] for cell in row['cells_by_forwarder'][company]) or '—'
                            check(tag + ' current source price ' + str(row_number) + '/' + str(column), sheet.cell(row_number, column).value == expected)
                            values.append(expected)
                        rows.append({'name': row['display_name'], 'prices': values})
                    main_text = '\n'.join(str(cell.value or '') for row in sheet for cell in row)
                    check(tag + ' clean price sheet', not any(term in main_text for term in ('MOCK_SDK', 'SDK未接続', 'AI_REQUIRED', 'MEMORY_APPLIED', '原文：')))
                    issue_sheet = [[str(cell.value or '') for cell in row] for row in book['未確認事項']]
                    save(tag + '_verification', {'sha256': sha256(raw).hexdigest(), 'size_bytes': len(raw), 'sheetnames': book.sheetnames, 'rows': rows, 'issues_sheet': issue_sheet, 'binary_in_bundle': False})
                finally:
                    book.close(); download.delete()

            export(first_saved, 'first_export')
            if stop_after_first:
                stage_status = get('/api/status'); save('status_after_first', stage_status)
                result['budget_after'] = stage_status['ai']['budget']
                check('First phase cumulative delta equals recorded calls',
                      result['budget_after']['requests_used'] - result['budget_before']['requests_used'] == first_route['classification_requests'])
                result['stage'] = 'FIRST_COMPLETE'
                result['status'] = 'PARTIAL'
                result['two_rounds'] = 'SECOND_NOT_RUN'
                result['next_step'] = 'Owner-managed service restart, then --resume-first-evidence this evidence directory; no repeat uploads or AI for first case'
                checkpoint(); browser.close()
                return 0
            if resume_first_evidence:
                recorded = json.loads((resume_first_evidence / 'first_after_confirmation.json').read_text())
                recorded_by_id = {item['item_id']: item for item in recorded['items']}
                check('All first-case saved facts and fee identities survived service restart',
                      len(recorded_by_id) == len(first_saved['items']) and all(item['item_id'] in recorded_by_id and all(
                          item[key] == recorded_by_id[item['item_id']][key]
                          for key in ('observation', 'source_ref', 'current_price', 'currency', 'current_fee_id', 'current_memory_version'))
                          for item in first_saved['items']))
                result['restart_persistence'] = 'PASS'
            second_start = get('/api/status')['ai']['budget']; save('second_budget_before', second_start)
            second_id = new_case('第2輪'); result['rounds']['second_case_id'] = second_id; checkpoint()
            second = batch(second_id, [NAMES[2]], 'second')
            second_route = routing(second, 'second')
            fresh = next(i for i in second['items'] if i['observation']['raw_label'] == '配送手数料')
            check('Second file yields all five actual items', len(second['items']) == 5 and len(second['documents']) == 1)
            check('Second delivery uses first human fee identity', fresh['current_fee_id'] == confirmed['current_fee_id'] and fresh['current_fee_name'] == '配送費' and fresh['classification_source'] == 'MEMORY')
            check('Second delivery9500 is new SDK fact not learned price', fresh['current_price'] == '9500' and fresh['currency'] == 'JPY' and fresh['source_ref']['field_sources']['raw_value']['original_value'] == '9500' and fresh['fact_sources']['amount'] == 'REAL_SDK')
            target = next(i for i in second_route['items'] if i['item_id'] == fresh['item_id'])
            check('Target item never sent for AI classification', not target['classification_requested'] and target['lookup']['decision'] == 'MEMORY_APPLIED' and target['lookup']['memory_version'] == confirmed['current_memory_version'])
            check('All five confirmed contexts reuse memory with no document AI', len(second_route['items']) == 5 and all(not i['classification_requested'] and i['lookup']['decision'] == 'MEMORY_APPLIED' for i in second_route['items']) and second_route['classification_requests'] == 0)
            result['rounds']['second_delivery_classification_calls'] = 0
            result['rounds']['second_quotation_extraction_calls'] = 0
            old = get('/api/cases/' + first_id); save('first_after_second_round', old)
            original_by_id = {i['item_id']: i for i in first_saved['items']}
            check('First case remains unchanged after second-file processing', all(all(item[key] == original_by_id[item['item_id']][key] for key in ('observation', 'source_ref', 'current_price', 'currency', 'current_fee_id')) for item in old['items']))
            check('First delivery still8000', next(i for i in old['items'] if i['item_id'] == delivery['item_id'])['current_price'] == '8000')
            page.locator('#comparison-table').screenshot(path=str(out / 'browser_second_table.png'))
            page.screenshot(path=str(out / 'browser_second_memory_9500.png'), full_page=True)
            preview_item(second, fresh, 'second_delivery')
            export(second, 'second_export')
            page.locator('#case-select').select_option(first_id)
            expect(page.locator('#comparison-table tbody tr')).to_have_count(5)
            page.locator('#comparison-table').screenshot(path=str(out / 'browser_first_still_8000.png'))
            page.locator('#case-select').select_option(second_id)
            final = get('/api/status'); save('status_after', final)
            result['budget_after'] = final['ai']['budget']
            check('Whole second process consumed zero cumulative requests', final['ai']['budget']['requests_used'] == second_start['requests_used'])
            check('Recorded real calls equal cumulative request delta', final['ai']['budget']['requests_used'] - result['budget_before']['requests_used'] == first_route['classification_requests'] + second_route['classification_requests'])
            check('No model or quota switch', final['ai']['budget']['authorization_id'] == budget['authorization_id'] and final['ai']['budget']['max_requests'] == 100)
            check('Normal uploads made zero company/transport preflights', result['sdk_preflights'] == [] and result['human_role_confirmations'] == [])
            check('Browser has no uncaught script errors', not page_errors)
            result['real_sdk'] = 'PASS'; result['two_rounds'] = 'PASS'; result['status'] = 'SUCCESS'
            checkpoint()
            browser.close()
    except Exception:
        result['status'] = 'PARTIAL'; result['error'] = traceback.format_exc()
        if page is not None:
            try:
                page.screenshot(path=str(out / 'browser_stopped_on_failure.png'), full_page=True)
            except Exception:
                pass
        print(result['error'], flush=True)
    finally:
        save('result', result)
    return 0 if result['status'] == 'SUCCESS' else 1


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--evidence', type=Path, required=True)
    parser.add_argument('--execute-approved-ai', action='store_true')
    parser.add_argument('--resume-first-evidence', type=Path)
    parser.add_argument('--stop-after-first', action='store_true', help='Save FIRST_COMPLETE/PARTIAL evidence and stop before second upload, enabling an explicit service restart.')
    args = parser.parse_args()
    if not args.execute_approved_ai:
        parser.error('Explicit --execute-approved-ai required; uses existing cumulative grant, with no retries.')
    if args.stop_after_first and args.resume_first_evidence:
        parser.error('--stop-after-first and --resume-first-evidence cannot be combined')
    raise SystemExit(run(args.evidence, args.resume_first_evidence, args.stop_after_first))
