"""Browser row confirmation of existing real Demo outputs. No AI or upload."""
import argparse,json,traceback
from pathlib import Path
from hashlib import sha256
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
BASE='http://127.0.0.1:8765'
def run(out,first_id,second_id,do_save):
    out.mkdir(parents=True,exist_ok=False)
    result={'scope':'REAL_POSTGRES_SAVED_MOCK_OPENAI_OUTPUT_ROW_UI','real_ai_calls_this_check':0,'real_sdk':'NOT_RUN','checks':[],'saves':[]}
    def save(n,d):(out/(n+'.json')).write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    def check(n,v):
        if not v:raise AssertionError(n)
        result['checks'].append(n);print('PASS '+n,flush=True)
    try:
        with sync_playwright() as pw:
            browser=pw.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
            page=browser.new_page(viewport={'width':1500,'height':1100},accept_downloads=True)
            errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            def guard(r):
                assert r.request.url.startswith(BASE+'/')
                assert r.request.method=='GET' or (do_save and r.request.method=='POST' and '/comparison/' in r.request.url and r.request.url.endswith('/confirmation'))
                r.continue_()
            page.route('**/*',guard)
            def get(path):
                r=page.request.get(BASE+path);assert r.ok;return r.json()
            budget=get('/api/status')['ai']['budget'];save('budget_before',budget)
            first=get('/api/cases/'+first_id);second=get('/api/cases/'+second_id)
            save('first_before',first);save('second_before',second)
            page.goto(BASE,wait_until='networkidle')
            def select(cid):
                page.locator('#case-select').select_option(cid)
                expect(page.locator('#files')).to_be_enabled()
                expect(page.locator('.compare-table tbody tr')).to_have_count(5)
            select(first_id)
            check('ten original items',len(first['items'])==10)
            for row in page.locator('.compare-table tbody tr').all():
                check('one left status and control',row.locator('td').first.locator('.row-status').count()==1 and row.locator('td').first.locator('.row-confirm-button').count()==1)
                for cell in row.locator('.price-cell').all():
                    check('price-only cell',cell.locator('button,input,.row-status,.quote-status').count()==0 and not any(x in cell.inner_text() for x in ['AI仮判定','確認済み','確認必要','確認・修正']))
            page.locator('#comparison-table').screenshot(path=str(out/'browser_first_before.png'))
            chosen=[('ocean20',next(r for r in first['comparison']['rows'] if r['differentiating_conditions'].get('equipment')=='20FT')),
                    ('delivery',next(r for r in first['comparison']['rows'] if any(i['observation']['raw_label']=='配送手数料' for i in r['member_items'])))]
            for tag,row in chosen:
                page.locator('[data-row-id="'+row['row_id']+'"] .row-confirm-button').click()
                expect(page.locator('#edit-dialog')).to_be_visible();expect(page.locator('.row-member')).to_have_count(2)
                check(tag+' no price inputs',page.locator('#edit-dialog').locator('#edit-price,#edit-currency,#amount-only,input[name=unit_price],input[name=currency]').count()==0)
                for item in row['member_items']:
                    card=page.locator('[data-member-id="'+item['item_id']+'"]')
                    expect(card).to_contain_text(item['observation']['raw_label'])
                    doc=next(d for d in first['documents'] if d['document_id']==item['observation']['document_id'])
                    expect(card).to_contain_text(doc['original_filename']);page.locator('#edit-details').evaluate('(e)=>{e.open=true;}')
                    detail=page.locator('[data-detail-member-id="'+item['item_id']+'"]')
                    check(tag+' eligibility '+item['item_id'],('自動学習の記録は作りません' in detail.inner_text()) == (not item['learning_eligible']))
                    card.locator('.member-preview').click();page.frame_locator('#preview-frame').locator('body').wait_for();page.wait_for_timeout(400)
                    check(tag+' original preview '+item['item_id'],doc['document_id'] in page.locator('#preview-frame').get_attribute('src'))
                    page.screenshot(path=str(out/('browser_'+tag+'_preview_'+('A' if 'A（架空）' in item['forwarder'] else 'B')+'.png')),full_page=True)
                    page.locator('#preview-dialog .close-dialog').click()
                page.locator('#edit-dialog').screenshot(path=str(out/('browser_'+tag+'_dialog.png')))
                if do_save:
                    current=get('/api/cases/'+first_id);fee=next((f for f in current['fees'] if f['display_name']==row['label']),None)
                    if fee:
                        page.locator('#edit-fee-name').fill(fee['display_name'])
                        page.locator('[data-fee-id="'+fee['fee_id']+'"]').click()
                    else:
                        page.locator('#edit-fee-name').fill(row['label'])
                    with page.expect_response(lambda r:r.request.method=='POST' and r.url.endswith('/confirmation')) as pending:page.locator('#save-correction').click()
                    response=pending.value;record={'tag':tag,'http_status':response.status,'request':response.request.post_data_json,'response':response.json()}
                    result['saves'].append(record);save('save_'+tag,record);check(tag+' HTTP200',response.status==200)
                    body=record['request'];check(tag+' name-only whole-row payload',len(body['members'])==2 and not {'unit_price','currency'}&body.keys() and all(not {'unit_price','currency'}&m.keys() for m in body['members']))
                    check(tag+' eligible memory only',sorted(m['memory_effect'] for m in record['response']['members'])==['ACTIVE','NOT_ELIGIBLE_CONTEXT_OR_SOURCE'])
                    expect(page.locator('#edit-dialog')).not_to_be_visible()
                    expect(page.locator('[data-row-id="'+row['row_id']+'"] .row-status')).to_have_text('✓ 確認済み')
                else:page.locator('#edit-dialog .close-dialog').click()
            after=get('/api/cases/'+first_id);save('first_after',after)
            for old in first['items']:
                new=next(i for i in after['items'] if i['item_id']==old['item_id'])
                check('immutable original facts '+old['item_id'],all(new[k]==old[k] for k in ['observation','source_ref','current_price','currency','ai_suggested_name','raw_value']))
            def export(snapshot,tag):
                with page.expect_download() as pending:page.locator('#export-button').click()
                path=out/(tag+'_export.xlsx');pending.value.save_as(path);book=load_workbook(path);sheet=book['原価単価一覧表'];projection=snapshot['comparison'];rows=[]
                for n,row in enumerate(projection['rows'],4):
                    check(tag+' row name '+str(n),sheet.cell(n,1).value==row['display_name']);values=[]
                    for col,company in enumerate(projection['companies'],2):
                        value='\n'.join(c['price']+' '+c['currency'] for c in row['cells_by_forwarder'][company]) or '—'
                        check(tag+' price-only export '+str(n)+company,sheet.cell(n,col).value==value);values.append(value)
                    rows.append({'name':row['display_name'],'values':values})
                check(tag+' confirmation audit included','確認情報' in book.sheetnames and book['確認情報'].max_row==len(snapshot['items'])+1)
                check(tag+' raw items retained',book['明細と出典'].max_row==len(snapshot['items'])+1)
                save(tag+'_export_check',{'rows':rows,'sha256':sha256(path.read_bytes()).hexdigest(),'sheets':book.sheetnames});book.close()
            page.locator('#comparison-table').screenshot(path=str(out/'browser_first_final.png'));page.screenshot(path=str(out/'browser_first_page.png'),full_page=True);export(after,'first')
            select(second_id);fresh=get('/api/cases/'+second_id);save('second_after',fresh)
            check('second case saved facts unchanged',all(all(i[k]==next(j for j in second['items'] if j['item_id']==i['item_id'])[k] for k in ['observation','source_ref','current_price','currency','current_fee_id','version']) for i in fresh['items']))
            delivery=next(i for i in fresh['items'] if i['observation']['raw_label']=='配送手数料');old=next(i for i in after['items'] if i['observation']['raw_label']=='配送手数料')
            check('second9500 first8000 same confirmed fee',delivery['current_price']=='9500' and old['current_price']=='8000' and delivery['current_fee_id']==old['current_fee_id'] and delivery['current_fee_name']=='配送費')
            check('second actual saved memory reuse',delivery['saved_status']=='MEMORY_APPLIED' and delivery['initial_classification']['memory_version']>=1)
            page.locator('#comparison-table').screenshot(path=str(out/'browser_second_final.png'));export(fresh,'second')
            final=get('/api/status')['ai']['budget'];save('budget_after',final)
            check('no new AI calls or budget reset',final['requests_used']==budget['requests_used'] and final['authorization_id']==budget['authorization_id'] and final['max_requests']==budget['max_requests'])
            check('no browser errors',not errors);result['status']='PASS';browser.close()
    except Exception:
        result['status']='FAIL';result['error']=traceback.format_exc();print(result['error'],flush=True)
    finally:save('result',result)
    return 0 if result['status']=='PASS' else 1
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--evidence',type=Path,required=True);p.add_argument('--first-case',required=True);p.add_argument('--second-case',required=True);p.add_argument('--save',action='store_true');args=p.parse_args()
    raise SystemExit(run(args.evidence,args.first_case,args.second_case,args.save))
