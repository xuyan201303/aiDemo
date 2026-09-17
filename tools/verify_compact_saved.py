"""Read-only browser verification of saved Demo facts with local static UI overlay.
The running immutable web image is not restarted and its keys are never read.
No POST, new AI, upload, or Demo database writes. Sources/projection served from
this checkout are explicitly recorded, separate from the live container image.
"""
import argparse,json,traceback
from hashlib import sha256
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
from sanzo_demo.comparison import comparison_projection
ROOT=Path(__file__).resolve().parents[1];BASE='http://127.0.0.1:8765'

def run(out,first_id,second_id,live_ui=False):
    out.mkdir(exist_ok=False,parents=True)
    result={'scope':'READ_ONLY_SAVED_DEMO_LIVE_UI' if live_ui else 'READ_ONLY_SAVED_DEMO_LOCAL_UI_OVERLAY','ai_calls':0,'post_requests':0,'checks':[],
        'static_sha256':{n:sha256((ROOT/'backend/sanzo_demo/static'/n).read_bytes()).hexdigest() for n in ['index.html','app.js','style.css','workflow.js']},
        'projection_sha256':sha256((ROOT/'backend/sanzo_demo/comparison.py').read_bytes()).hexdigest()}
    def dump(n,v):(out/(n+'.json')).write_text(json.dumps(v,ensure_ascii=False,indent=2))
    def check(n,v=True):
        assert v,n
        result['checks'].append(n);print('PASS '+n,flush=True)
    try:
        with sync_playwright() as pw:
            browser=pw.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
            page=browser.new_page(viewport={'width':1366,'height':768});errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            def get(path):
                r=page.request.get(BASE+path);assert r.ok;return r.json()
            if live_ui:
                for n in result['static_sha256']:
                    response=page.request.get(BASE+('/' if n=='index.html' else '/static/'+n))
                    check('live serves checkout '+n,response.ok and sha256(response.body()).hexdigest()==result['static_sha256'][n])
            baseline={i:get('/api/cases/'+i) for i in [first_id,second_id]};budget=get('/api/status')['ai']['budget']
            dump('saved_before',baseline);dump('budget_before',budget)
            def guard(r):
                assert r.request.url.startswith(BASE+'/') and r.request.method=='GET'
                path=r.request.url.removeprefix(BASE)
                if live_ui:return r.continue_()
                if path=='/':return r.fulfill(path=str(ROOT/'backend/sanzo_demo/static/index.html'))
                if path.startswith('/static/'):
                    name=path.removeprefix('/static/');assert name in result['static_sha256']
                    return r.fulfill(path=str(ROOT/'backend/sanzo_demo/static'/name))
                if path in ['/api/cases/'+first_id,'/api/cases/'+second_id]:
                    response=r.fetch();data=response.json();data['comparison']=comparison_projection(data['items'],data['fees'])
                    return r.fulfill(response=response,json=data)
                r.continue_()
            page.route('**/*',guard);page.goto(BASE,wait_until='networkidle');page.locator('#case-select').select_option(first_id)
            expect(page.locator('.compare-table tbody tr')).to_have_count(5)
            first=baseline[first_id];projection=comparison_projection(first['items'],first['fees']);row=next(r for r in projection['rows'] if any(i['observation']['raw_label']=='配送手数料' for i in r['member_items']))
            page.locator('[data-row-id="'+row['row_id']+'"] .row-confirm-button').click()
            expect(page.locator('.row-member')).to_have_count(2);expect(page.locator('#edit-fee-name')).to_have_value('配送費')
            box=page.locator('#edit-dialog').bounding_box();check('normal panel fits1366x768',0<=box['y'] and box['y']+box['height']<=768)
            check('one name input visible',page.locator('#edit-dialog input:visible').count()==1)
            page.screenshot(path=str(out/'browser_normal_full.png'));page.locator('#edit-dialog').screenshot(path=str(out/'browser_normal_panel.png'))
            check('condition notice hidden by default',not page.locator('#edit-warning').is_visible())
            check('default single name and save',page.locator('#edit-dialog input:visible').count()==1 and page.locator('#edit-dialog select:visible,#edit-dialog input[type=checkbox]:visible').count()==0)
            for control in ['#edit-warning','#edit-dimensions','#refresh-item','#replace-mapping','#revoke-button','#edit-status','#edit-unified-name']:
                expect(page.locator(control)).not_to_be_visible()
            heights=page.locator('.row-member').evaluate_all('(es)=>es.map(e=>e.getBoundingClientRect().height)')
            result['card_heights_px']=heights
            check('typical card heights90to120px',all(90<=h<=120 for h in heights))
            page.locator('#edit-details > summary').click();expect(page.locator('#edit-warning')).to_be_visible()
            page.locator('#edit-details > summary').click()
            for item in row['member_items']:
                card=page.locator('[data-member-id="'+item['item_id']+'"]');expect(card).to_contain_text(item['observation']['raw_label'])
                doc_id=item['observation']['document_id'];doc=next(d for d in first['documents'] if d['document_id']==doc_id)
                expect(card).to_contain_text(doc['original_filename']);card.locator('.member-preview').click()
                page.frame_locator('#preview-frame').locator('body').wait_for();page.wait_for_timeout(200)
                check('original file and source '+doc_id,doc_id in page.locator('#preview-frame').get_attribute('src'))
                page.screenshot(path=str(out/('browser_preview_'+('excel' if doc['original_filename'].endswith('.xlsx') else 'pdf')+'.png')))
                page.locator('#preview-dialog .close-dialog').click();expect(page.locator('#edit-fee-name')).to_have_value('配送費')
            page.locator('#edit-fee-name').fill('保存しない入力途中');page.locator('.member-preview').first.click();page.locator('#preview-dialog .close-dialog').click()
            expect(page.locator('#edit-fee-name')).to_have_value('保存しない入力途中');check('actual original preview retains edited draft')
            page.once('dialog',lambda d:d.accept());page.locator('#edit-dialog .close-dialog').click()
            page.locator('#comparison-table').screenshot(path=str(out/'browser_comparison_saved.png'))
            for cell in page.locator('.price-cell').all():check('read-only price cell',cell.locator('button,input,.row-status').count()==0)
            # Existing English confirmed name is preserved; no translation.
            english=next(r for r in projection['rows'] if any('Ocean Freight' in (i.get('current_fee_name') or '') for i in r['member_items']))
            page.locator('[data-row-id="'+english['row_id']+'"] .row-confirm-button').click()
            expect(page.locator('#edit-fee-name')).to_have_value(english['label']);check('confirmed English name unchanged')
            page.locator('#edit-dialog').screenshot(path=str(out/'browser_ocean40_panel.png'))
            page.locator('#edit-dialog .close-dialog').click()
            for cid in [first_id,second_id]:
                snap=baseline[cid];p=comparison_projection(snap['items'],snap['fees'])
                response=page.request.get(BASE+'/api/cases/'+cid+'/export.xlsx');assert response.ok
                book=load_workbook(BytesIO(response.body()),read_only=True);sheet=book['原価単価一覧表'];rows=[]
                for n,r in enumerate(p['rows'],4):
                    values=[r['display_name']]+['\n'.join(c['price']+' '+c['currency'] for c in r['cells_by_forwarder'][company]) or '—' for company in p['companies']]
                    check('export values '+cid+' '+str(n),[sheet.cell(n,col).value for col in range(1,len(values)+1)]==values);rows.append(values)
                dump('export_'+cid,{'sha256':sha256(response.body()).hexdigest(),'rows':rows,'sheets':book.sheetnames});book.close()
            second=baseline[second_id];old=next(i for i in first['items'] if i['observation']['raw_label']=='配送手数料');new=next(i for i in second['items'] if i['observation']['raw_label']=='配送手数料')
            check('saved9500new8000oldsameconfirmedfee',new['current_price']=='9500' and old['current_price']=='8000' and new['current_fee_id']==old['current_fee_id'] and new['saved_status']=='MEMORY_APPLIED')
            page.locator('#case-select').select_option(second_id);expect(page.locator('.compare-table tbody tr')).to_have_count(5)
            page.locator('#comparison-table').screenshot(path=str(out/'browser_second_saved.png'))
            after={i:get('/api/cases/'+i) for i in baseline};final_budget=get('/api/status')['ai']['budget'];dump('saved_after',after);dump('budget_after',final_budget)
            check('saved cases complete snapshots unchanged',after==baseline);check('budget unchanged',budget==final_budget);check('no browser errors',not errors)
            result['status']='PASS';browser.close()
    except Exception:result.update(status='FAIL',error=traceback.format_exc());print(result['error'])
    finally:dump('result',result)
    return 0 if result['status']=='PASS' else 1
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--evidence',type=Path,required=True);p.add_argument('--first-case',required=True);p.add_argument('--second-case',required=True);p.add_argument('--live-ui',action='store_true');a=p.parse_args();raise SystemExit(run(a.evidence,a.first_case,a.second_case,a.live_ui))
